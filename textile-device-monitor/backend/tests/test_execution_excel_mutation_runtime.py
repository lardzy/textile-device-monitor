from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.execution import (
    AuthContext,
    copy_run_mutation,
    preflight_run_mutation,
    publish_run_mutation,
    run_mutation_detail,
    verify_run_mutation,
    write_run_mutation,
)
from app.config import settings
from app.database import Base
from app.execution.catalog import (
    bind_user_role,
    create_workflow,
    ensure_default_rbac,
    publish_workflow,
)
from app.execution.engine import (
    NodeExecutionContext,
    claim_human_task,
    claim_next_node,
    create_run,
    execute_claimed_node,
    set_run_control_status,
    submit_human_task,
)
from app.execution.errors import ExecutionApiError
from app.execution.excel_runtime import register_excel_executors
from app.execution.models import (
    ExecutionArtifact,
    ExecutionArtifactRelation,
    ExecutionAuditLog,
    ExecutionCategory,
    ExecutionFileMutation,
    ExecutionHumanTask,
    ExecutionNodeRun,
    ExecutionPublishReceipt,
    ExecutionUser,
)
from app.execution.mutations import FileMutationService
from app.execution.mutation_runtime import register_mutation_executors
from app.execution.persistence import ensure_storage_roots
from app.execution.registry import node_registry
from app.execution.schemas import (
    MutationPreflightRequest,
    MutationPublishRequest,
    MutationStageRequest,
    MutationVerifyRequest,
    MutationWriteRequest,
)
from app.execution.security import hash_password


def _definition() -> dict:
    node_specs = [
        ("start", "core.start", {}),
        ("classify", "excel.classify", {}),
        ("extract", "excel.extract_summary", {}),
        ("copy", "workbook.copy", {"staging_root_id": "execution_staging"}),
        ("write", "workbook.write_cells", {}),
        ("verify", "workbook.verify", {}),
        ("confirm", "human.confirm", {"title": "确认发布"}),
        (
            "publish",
            "artifact.publish",
            {
                "publish_root_id": "execution_publish",
                "confirmation_node_id": "confirm",
            },
        ),
        ("end", "core.end", {}),
    ]
    mappings = {
        "classify": {"source": "$.inputs.source"},
        "extract": {"source": "$.inputs.source"},
        "copy": {
            "source": "$.inputs.source",
            "mutation_id": "$.inputs.mutation_id",
        },
        "write": {
            "mutation_id": "$.inputs.mutation_id",
            "working_copy": "$.nodes.copy.output.working_copy",
            "writes": "$.inputs.writes",
        },
        "verify": {
            "mutation_id": "$.inputs.mutation_id",
            "working_copy": "$.nodes.copy.output.working_copy",
            "writes": "$.inputs.writes",
            "target": "$.inputs.target",
        },
        "confirm": {
            "approval_context":
                "$.nodes.verify.output.approval_context",
        },
        "publish": {
            "mutation_id": "$.inputs.mutation_id",
            "working_copy": "$.nodes.copy.output.working_copy",
            "target": "$.inputs.target",
        },
        "end": {"published": "$.nodes.publish.output"},
    }
    return {
        "schema_version": "1.0",
        "metadata": {"name": "Excel 受控写入测试"},
        "input_schema": {
            "type": "object",
            "properties": {
                "inspection_number": {"type": "string"},
                "source": {"type": "object"},
                "mutation_id": {"type": "string"},
                "writes": {"type": "array"},
                "target": {"type": "object"},
            },
            "required": [
                "inspection_number",
                "source",
                "mutation_id",
                "writes",
                "target",
            ],
        },
        "global_schema": {"type": "object", "properties": {}},
        "root_slots": [
            {
                "name": "source",
                "root_id": "special_wool_records",
                "access": "read",
            },
            {
                "name": "staging",
                "root_id": "execution_staging",
                "access": "write",
            },
            {
                "name": "publish",
                "root_id": "execution_publish",
                "access": "publish",
            },
        ],
        "credential_slots": [],
        "nodes": [
            {
                "id": node_id,
                "type": node_type,
                "type_version": 1,
                "name": node_id,
                "config": config,
                "input_mapping": mappings.get(node_id, {}),
                "ui": {},
            }
            for node_id, node_type, config in node_specs
        ],
        "edges": [
            {
                "id": f"edge-{index}",
                "source": node_specs[index][0],
                "target": node_specs[index + 1][0],
            }
            for index in range(len(node_specs) - 1)
        ],
    }


class ExecutionExcelMutationRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        root = Path(self.tempdir.name)
        self.source = root / "source"
        self.runtime = root / "runtime"
        self.publish = root / "publish"
        self.records = self.source / "2026-特种毛"
        for path in (
            self.records,
            self.source / "2026-麻棉",
            self.source / "2026-电镜",
            self.runtime,
            self.publish,
        ):
            path.mkdir(parents=True)
        self.source_workbook = self.records / "26X0001.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "原始记录"
        sheet["A1"] = "棉和再生纤维素纤维物理法"
        sheet["B2"] = "26X0001"
        sheet["C3"] = 12
        workbook.save(self.source_workbook)
        workbook.close()

        self.original_settings = (
            settings.EXECUTION_SOURCE_ROOT,
            settings.EXECUTION_RUNTIME_ROOT,
            settings.EXECUTION_PUBLISH_ROOT,
        )
        settings.EXECUTION_SOURCE_ROOT = str(self.source)
        settings.EXECUTION_RUNTIME_ROOT = str(self.runtime)
        settings.EXECUTION_PUBLISH_ROOT = str(self.publish)

        self.engine = create_engine("sqlite:///:memory:")
        self.Session = sessionmaker(bind=self.engine, autoflush=False)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        ensure_default_rbac(self.db)
        self.user = ExecutionUser(
            username="excel-admin",
            display_name="Excel 管理员",
            password_hash=hash_password("excel-test-password"),
            role="admin",
        )
        category = ExecutionCategory(key="excel-test", name="Excel 测试")
        self.db.add_all([self.user, category])
        self.db.flush()
        bind_user_role(self.db, self.user, "admin", created_by_id=self.user.id)
        self.workflow = create_workflow(
            self.db,
            actor=self.user,
            slug="excel-controlled-write",
            category_id=category.id,
            name="Excel 受控写入",
            description=None,
            definition=_definition(),
            capabilities={"read": True, "write": True},
            is_enabled=True,
        )
        publish_workflow(
            self.db,
            workflow_id=self.workflow.id,
            expected_revision=1,
            actor=self.user,
            release_note="test",
        )
        ensure_storage_roots(self.db)
        self.run, _ = create_run(
            self.db,
            workflow=self.workflow,
            actor=self.user,
            inspection_number="26X0001",
            input_data={
                "source": {
                    "root_id": "special_wool_records",
                    "relative_path": "26X0001.xlsx",
                },
                "mutation_id": "mutation-worker-e2e",
                "writes": [
                    {
                        "sheet": "原始记录",
                        "cell": "C3",
                        "value": 77,
                    }
                ],
                "target": {
                    "root_id": "execution_publish",
                    "relative_path": "results/worker-e2e.xlsx",
                },
            },
            global_data={},
            idempotency_key="excel-mutation-run-1",
        )
        self.db.commit()
        register_excel_executors()
        register_mutation_executors()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        (
            settings.EXECUTION_SOURCE_ROOT,
            settings.EXECUTION_RUNTIME_ROOT,
            settings.EXECUTION_PUBLISH_ROOT,
        ) = self.original_settings
        self.tempdir.cleanup()

    def _context(
        self,
        node_id: str,
        *,
        input_data: dict,
        config: dict | None = None,
    ) -> NodeExecutionContext:
        node_run = (
            self.db.query(ExecutionNodeRun)
            .filter(
                ExecutionNodeRun.run_id == self.run.id,
                ExecutionNodeRun.node_id == node_id,
            )
            .one()
        )
        node = next(
            item
            for item in self.run.definition_snapshot["nodes"]
            if item["id"] == node_id
        )
        node = {**node, "config": config if config is not None else node["config"]}
        return NodeExecutionContext(
            db=self.db,
            run=self.run,
            node_run=node_run,
            node=node,
            input_data=input_data,
            worker_id="test-worker",
            lease_token="test-lease",
        )

    def _execute(self, node_type: str, context: NodeExecutionContext) -> dict:
        executor = node_registry.executor(node_type, 1)
        self.assertIsNotNone(executor)
        return executor(context)

    def _set_direct_publish_state(
        self,
        *,
        approval_context: dict,
        approved: bool,
    ) -> tuple[ExecutionNodeRun, ExecutionHumanTask]:
        confirm = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="confirm")
            .one()
        )
        confirm.status = "succeeded"
        confirm.input_data = {"approval_context": approval_context}
        confirm.output_data = {"approved": approved}
        task = (
            self.db.query(ExecutionHumanTask)
            .filter_by(node_run_id=confirm.id)
            .one_or_none()
        )
        if task is None:
            task = ExecutionHumanTask(
                run_id=self.run.id,
                node_run_id=confirm.id,
                title="确认发布",
                form_schema={},
                status="completed",
                revision=1,
                assigned_user_id=self.user.id,
            )
            self.db.add(task)
        task.status = "completed"
        task.claimed_by_id = self.user.id
        task.completed_by_id = self.user.id
        task.result_data = {"approved": approved}
        publish = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="publish")
            .one()
        )
        publish.status = "running"
        publish.lease_owner = "test-worker"
        publish.lease_token = "test-lease"
        self.db.flush()
        return confirm, task

    def _run_worker_node(self) -> None:
        node = claim_next_node(
            self.db,
            worker_id="excel-e2e-worker",
            lease_seconds=30,
        )
        self.assertIsNotNone(node)
        node_id = node.id
        lease_token = node.lease_token
        self.db.commit()
        execute_claimed_node(self.db, node_id, lease_token)
        self.db.commit()

    def _run_until_node_status(
        self,
        run_id: str,
        node_id: str,
        statuses: set[str],
    ) -> ExecutionNodeRun:
        for _ in range(20):
            node = (
                self.db.query(ExecutionNodeRun)
                .filter_by(run_id=run_id, node_id=node_id)
                .one()
            )
            if node.status in statuses:
                return node
            self._run_worker_node()
        self.fail(
            f"node {node_id} did not reach {sorted(statuses)}"
        )

    def _prepare_public_publish(
        self,
    ) -> tuple[AuthContext, dict, dict]:
        auth = AuthContext(session=None, user=self.user)
        mutation_id = self.run.input_data["mutation_id"]
        source = self.run.input_data["source"]
        writes = self.run.input_data["writes"]
        target = self.run.input_data["target"]

        self._run_until_node_status(self.run.id, "copy", {"ready"})
        preflight_run_mutation(
            self.run.id,
            MutationPreflightRequest(
                mutation_id=mutation_id,
                source=source,
            ),
            auth=auth,
            db=self.db,
        )
        copy_run_mutation(
            self.run.id,
            MutationStageRequest(node_id="copy"),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )

        self._run_until_node_status(self.run.id, "write", {"ready"})
        write_run_mutation(
            self.run.id,
            MutationWriteRequest(node_id="write", writes=writes),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )

        self._run_until_node_status(self.run.id, "verify", {"ready"})
        verified = verify_run_mutation(
            self.run.id,
            MutationVerifyRequest(node_id="verify", target=target),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        approval_context = verified["verification"]["approval_context"]

        self._run_until_node_status(
            self.run.id,
            "confirm",
            {"waiting_human"},
        )
        task = self.db.query(ExecutionHumanTask).one()
        task = claim_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            actor=self.user,
        )
        self.db.commit()
        submit_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"approved": True},
            actor=self.user,
        )
        self.db.commit()
        publish_node = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="publish")
            .one()
        )
        self.assertEqual(publish_node.status, "ready")
        return auth, target, approval_context

    def test_classification_and_summary_use_safe_artifact_reference(self):
        source = {
            "root_id": "special_wool_records",
            "relative_path": "26X0001.xlsx",
        }
        classified = self._execute(
            "excel.classify",
            self._context(
                "classify",
                input_data={"source": source},
                config={
                    "types": [
                        {
                            "name": "棉再物理法",
                            "features": [
                                {
                                    "sheet": "原始记录",
                                    "cell": "A1",
                                    "contains": "再生纤维素",
                                }
                            ],
                        }
                    ]
                },
            ),
        )
        self.assertEqual(classified["detected_type"], "棉再物理法")
        self.assertEqual(classified["sheet_names"], ["原始记录"])

        summary = self._execute(
            "excel.extract_summary",
            self._context(
                "extract",
                input_data={"source": source},
                config={
                    "fields": [
                        {
                            "name": "inspection_number",
                            "sheet": "原始记录",
                            "cell": "B2",
                            "required": True,
                        },
                        {
                            "name": "count",
                            "sheet": "原始记录",
                            "cell": "C3",
                        },
                    ]
                },
            ),
        )
        self.assertEqual(
            summary["values"],
            {"inspection_number": "26X0001", "count": 12},
        )

    def test_worker_human_confirmation_and_publish_full_chain(self):
        for _ in range(7):
            self._run_worker_node()

        task = self.db.query(ExecutionHumanTask).one()
        self.assertEqual(task.status, "open")
        self.assertIn("approval_context", task.node_run.input_data)
        task = claim_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            actor=self.user,
        )
        self.db.commit()
        self.assertIs(
            task.form_schema["properties"]["approved"]["const"],
            True,
        )
        with self.assertRaises(ExecutionApiError) as rejected:
            submit_human_task(
                self.db,
                task_id=task.id,
                expected_revision=task.revision,
                data={"approved": False},
                actor=self.user,
            )
        self.assertEqual(
            rejected.exception.code,
            "human_task_input_invalid",
        )
        self.db.rollback()
        task = self.db.get(ExecutionHumanTask, task.id)
        submit_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"approved": True},
            actor=self.user,
        )
        self.db.commit()

        self._run_worker_node()
        self._run_worker_node()
        self.db.refresh(self.run)
        self.assertEqual(self.run.status, "completed")
        published_path = self.publish / "results" / "worker-e2e.xlsx"
        self.assertTrue(published_path.exists())
        book = load_workbook(published_path, read_only=True)
        try:
            self.assertEqual(book["原始记录"]["C3"].value, 77)
        finally:
            book.close()

    def test_copy_write_verify_publish_preserves_source_and_is_idempotent(self):
        source = {
            "root_id": "special_wool_records",
            "relative_path": "26X0001.xlsx",
        }
        mutation_id = "mutation-runtime-0001"
        copy_result = self._execute(
            "workbook.copy",
            self._context(
                "copy",
                input_data={"source": source, "mutation_id": mutation_id},
            ),
        )
        writes = [{"sheet": "原始记录", "cell": "C3", "value": 42}]
        common_input = {
            "mutation_id": mutation_id,
            "working_copy": copy_result["working_copy"],
            "writes": writes,
            "target": {
                "root_id": "execution_publish",
                "relative_path": "results/26X0001.xlsx",
            },
        }
        self._execute(
            "workbook.write_cells",
            self._context("write", input_data=common_input),
        )
        mutation = (
            self.db.query(ExecutionFileMutation)
            .filter_by(mutation_id=mutation_id)
            .one()
        )
        self.assertEqual(mutation.status, "written")

        confirm = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="confirm")
            .one()
        )
        confirm.status = "succeeded"
        publish_context = self._context(
            "publish",
            input_data={
                **common_input,
                "target": {
                    "root_id": "execution_publish",
                    "relative_path": "results/26X0001.xlsx",
                },
            },
        )
        with self.assertRaises(ExecutionApiError) as blocked:
            self._execute("artifact.publish", publish_context)
        self.assertEqual(blocked.exception.code, "mutation_not_verified")
        self.assertFalse((self.publish / "results" / "26X0001.xlsx").exists())

        verified = self._execute(
            "workbook.verify",
            self._context("verify", input_data=common_input),
        )
        self.assertTrue(verified["verified"])
        self.assertEqual(mutation.status, "verified")

        confirm, task = self._set_direct_publish_state(
            approval_context=verified["approval_context"],
            approved=False,
        )
        with self.assertRaises(ExecutionApiError) as confirmation_blocked:
            self._execute("artifact.publish", publish_context)
        self.assertEqual(
            confirmation_blocked.exception.code,
            "publish_confirmation_required",
        )
        confirm.output_data = {"approved": True}
        task.result_data = {"approved": True}
        unauthorized = ExecutionUser(
            username="unauthorized-approver",
            display_name="无权批准人",
            password_hash=hash_password("unauthorized-password"),
            role="user",
        )
        self.db.add(unauthorized)
        self.db.flush()
        task.claimed_by_id = unauthorized.id
        task.completed_by_id = unauthorized.id
        with self.assertRaises(ExecutionApiError) as unauthorized_publish:
            self._execute("artifact.publish", publish_context)
        self.assertEqual(
            unauthorized_publish.exception.code,
            "publish_approver_not_authorized",
        )
        task.claimed_by_id = self.user.id
        task.completed_by_id = self.user.id
        published = self._execute("artifact.publish", publish_context)
        published_again = self._execute("artifact.publish", publish_context)
        self.assertFalse(published["reused"])
        self.assertTrue(published_again["reused"])
        self.db.commit()

        source_book = load_workbook(self.source_workbook, read_only=True)
        published_book = load_workbook(
            self.publish / "results" / "26X0001.xlsx",
            read_only=True,
        )
        try:
            self.assertEqual(source_book["原始记录"]["C3"].value, 12)
            self.assertEqual(published_book["原始记录"]["C3"].value, 42)
        finally:
            source_book.close()
            published_book.close()

        self.assertEqual(mutation.status, "published")
        self.assertEqual(self.db.query(ExecutionPublishReceipt).count(), 1)
        self.assertGreaterEqual(self.db.query(ExecutionArtifact).count(), 4)
        self.assertGreaterEqual(self.db.query(ExecutionArtifactRelation).count(), 3)

    def test_verify_rejects_writes_that_differ_from_persisted_change_plan(self):
        source = {
            "root_id": "special_wool_records",
            "relative_path": "26X0001.xlsx",
        }
        mutation_id = "mutation-runtime-plan-binding"
        copied = self._execute(
            "workbook.copy",
            self._context(
                "copy",
                input_data={"source": source, "mutation_id": mutation_id},
            ),
        )
        original_writes = [{"sheet": "原始记录", "cell": "C3", "value": 42}]
        self._execute(
            "workbook.write_cells",
            self._context(
                "write",
                input_data={
                    "mutation_id": mutation_id,
                    "working_copy": copied["working_copy"],
                    "writes": original_writes,
                },
            ),
        )

        with self.assertRaises(ExecutionApiError) as blocked:
            self._execute(
                "workbook.verify",
                self._context(
                    "verify",
                    input_data={
                        "mutation_id": mutation_id,
                        "working_copy": copied["working_copy"],
                        "writes": [
                            {
                                "sheet": "原始记录",
                                "cell": "C3",
                                "value": 999,
                            }
                        ],
                    },
                ),
            )
        self.assertEqual(
            blocked.exception.code,
            "mutation_change_plan_mismatch",
        )

    def test_designer_test_mode_never_writes_publish_target(self):
        self.run.mode = "test"
        source = {
            "root_id": "special_wool_records",
            "relative_path": "26X0001.xlsx",
        }
        mutation_id = "mutation-runtime-dry-run"
        copied = self._execute(
            "workbook.copy",
            self._context(
                "copy",
                input_data={"source": source, "mutation_id": mutation_id},
            ),
        )
        target = {
            "root_id": "execution_publish",
            "relative_path": "results/never-published.xlsx",
        }
        common = {
            "mutation_id": mutation_id,
            "working_copy": copied["working_copy"],
            "writes": [{"sheet": "原始记录", "cell": "C3", "value": 88}],
            "target": target,
        }
        self._execute(
            "workbook.write_cells",
            self._context("write", input_data=common),
        )
        verified = self._execute(
            "workbook.verify",
            self._context("verify", input_data=common),
        )
        self._set_direct_publish_state(
            approval_context=verified["approval_context"],
            approved=True,
        )

        result = self._execute(
            "artifact.publish",
            self._context("publish", input_data=common),
        )
        self.assertTrue(result["dry_run"])
        self.assertFalse(
            (self.publish / "results" / "never-published.xlsx").exists()
        )
        mutation = (
            self.db.query(ExecutionFileMutation)
            .filter_by(mutation_id=mutation_id)
            .one()
        )
        self.assertEqual(mutation.status, "verified")
        self.assertEqual(self.db.query(ExecutionPublishReceipt).count(), 0)

    def test_copy_failure_keeps_preflight_plan_and_error_for_node_audit(self):
        mutation_id = "mutation-copy-failure"

        class FailingCopyService:
            def prepare_working_copy(_self, *args, **kwargs):
                planned = (
                    self.db.query(ExecutionFileMutation)
                    .filter_by(mutation_id=mutation_id)
                    .one()
                )
                self.assertEqual(planned.status, "planned")
                self.assertIsNotNone(planned.source_artifact_id)
                self.assertIn("source", planned.change_plan)
                raise OSError("simulated-copy-failure")

        with patch(
            "app.execution.mutation_runtime._mutation_service",
            return_value=FailingCopyService(),
        ):
            with self.assertRaises(ExecutionApiError) as captured:
                self._execute(
                    "workbook.copy",
                    self._context(
                        "copy",
                        input_data={
                            "source": {
                                "root_id": "special_wool_records",
                                "relative_path": "26X0001.xlsx",
                            },
                            "mutation_id": mutation_id,
                        },
                    ),
                )

        self.assertEqual(captured.exception.code, "file_mutation_failed")
        mutation = (
            self.db.query(ExecutionFileMutation)
            .filter_by(mutation_id=mutation_id)
            .one()
        )
        self.assertEqual(mutation.status, "failed")
        self.assertEqual(mutation.error_code, "file_mutation_failed")
        self.assertIn("working_copy", mutation.change_plan)
        self.assertIsNone(mutation.working_artifact_id)

    def test_public_preflight_rejects_skipped_copy_node(self):
        copy_node = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="copy")
            .one()
        )
        copy_node.status = "skipped"
        self.run.status = "running"
        self.db.commit()

        with self.assertRaises(ExecutionApiError) as blocked:
            preflight_run_mutation(
                self.run.id,
                MutationPreflightRequest(
                    mutation_id=self.run.input_data["mutation_id"],
                    node_id="copy",
                    source=self.run.input_data["source"],
                ),
                auth=AuthContext(session=None, user=self.user),
                db=self.db,
            )
        self.assertEqual(blocked.exception.code, "mutation_node_not_reached")
        self.assertEqual(
            self.db.query(ExecutionFileMutation).count(),
            0,
        )

    def test_publish_cancel_interleaving_records_side_effect_once(self):
        auth, target, approval_context = self._prepare_public_publish()
        mutation_id = self.run.input_data["mutation_id"]
        original_publish = FileMutationService.publish
        cancel_observed = False

        def cancel_after_fence(service, *args, **kwargs):
            nonlocal cancel_observed
            fenced = (
                self.db.query(ExecutionFileMutation)
                .filter_by(
                    run_id=self.run.id,
                    mutation_id=mutation_id,
                )
                .one()
            )
            self.assertEqual(fenced.status, "publishing")
            self.assertIsNotNone(fenced.publish_fence_token)
            settling_run = set_run_control_status(
                self.db,
                run_id=self.run.id,
                action="cancel",
                actor=self.user,
            )
            self.assertEqual(settling_run.status, "cancel_pending")
            self.db.commit()
            cancel_observed = True
            return original_publish(service, *args, **kwargs)

        with patch.object(
            FileMutationService,
            "publish",
            new=cancel_after_fence,
        ):
            result = publish_run_mutation(
                self.run.id,
                MutationPublishRequest(
                    node_id="publish",
                    target=target,
                    approval_context=approval_context,
                ),
                mutation_id=mutation_id,
                auth=auth,
                db=self.db,
            )

        self.assertTrue(cancel_observed)
        self.assertFalse(result["reused"])
        self.db.refresh(self.run)
        self.assertEqual(self.run.status, "cancelled")
        self.assertTrue(
            (self.run.output_data or {}).get(
                "cancelled_after_side_effect"
            )
        )
        publish_node = (
            self.db.query(ExecutionNodeRun)
            .filter_by(run_id=self.run.id, node_id="publish")
            .one()
        )
        self.assertEqual(publish_node.status, "succeeded")
        mutation = (
            self.db.query(ExecutionFileMutation)
            .filter_by(mutation_id=mutation_id)
            .one()
        )
        self.assertEqual(mutation.status, "published")
        self.assertIsNone(mutation.publish_fence_token)
        self.assertEqual(self.db.query(ExecutionPublishReceipt).count(), 1)
        self.assertTrue(
            (self.publish / "results" / "worker-e2e.xlsx").exists()
        )

    def test_public_mutation_stages_are_audited_and_idempotent(self):
        auth = AuthContext(session=None, user=self.user)
        mutation_id = self.run.input_data["mutation_id"]
        source_before = self.source_workbook.read_bytes()
        self._run_until_node_status(
            self.run.id,
            "copy",
            {"ready", "running"},
        )
        preflight_payload = MutationPreflightRequest(
            mutation_id=mutation_id,
            source={
                "root_id": "special_wool_records",
                "relative_path": "26X0001.xlsx",
            },
        )

        first_plan = preflight_run_mutation(
            self.run.id,
            preflight_payload,
            auth=auth,
            db=self.db,
        )
        repeated_plan = preflight_run_mutation(
            self.run.id,
            preflight_payload,
            auth=auth,
            db=self.db,
        )
        self.assertFalse(first_plan["reused"])
        self.assertTrue(repeated_plan["reused"])
        self.assertEqual(
            self.db.query(ExecutionFileMutation)
            .filter_by(mutation_id=mutation_id)
            .count(),
            1,
        )

        first_copy = copy_run_mutation(
            self.run.id,
            MutationStageRequest(),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        repeated_copy = copy_run_mutation(
            self.run.id,
            MutationStageRequest(),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        self.assertFalse(first_copy["reused"])
        self.assertTrue(repeated_copy["reused"])

        self._run_until_node_status(
            self.run.id,
            "write",
            {"ready", "running"},
        )
        write_payload = MutationWriteRequest(
            writes=[
                {
                    "sheet": "原始记录",
                    "cell": "C3",
                    "value": 77,
                }
            ]
        )
        first_write = write_run_mutation(
            self.run.id,
            write_payload,
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        repeated_write = write_run_mutation(
            self.run.id,
            write_payload,
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        self.assertFalse(first_write["reused"])
        self.assertTrue(repeated_write["reused"])

        self._run_until_node_status(
            self.run.id,
            "verify",
            {"ready", "running"},
        )
        target = self.run.input_data["target"]
        verified = verify_run_mutation(
            self.run.id,
            MutationVerifyRequest(target=target),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        approval_context = verified["verification"]["approval_context"]

        with self.assertRaises(ExecutionApiError) as not_reached:
            publish_run_mutation(
                self.run.id,
                MutationPublishRequest(
                    target=target,
                    approval_context=approval_context,
                ),
                mutation_id=mutation_id,
                auth=auth,
                db=self.db,
            )
        self.assertEqual(
            not_reached.exception.code,
            "mutation_node_not_reached",
        )

        self._run_until_node_status(
            self.run.id,
            "confirm",
            {"waiting_human"},
        )
        task = self.db.query(ExecutionHumanTask).one()
        task = claim_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            actor=self.user,
        )
        self.db.commit()
        submit_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"approved": True},
            actor=self.user,
        )
        self.db.commit()

        with self.assertRaises(ExecutionApiError) as blocked:
            publish_run_mutation(
                self.run.id,
                MutationPublishRequest(
                    target=target,
                    approval_context={
                        **approval_context,
                        "change_plan_checksum": "0" * 64,
                    },
                ),
                mutation_id=mutation_id,
                auth=auth,
                db=self.db,
            )
        self.assertEqual(
            blocked.exception.code,
            "publish_confirmation_required",
        )
        self.assertFalse(
            (self.publish / "results" / "worker-e2e.xlsx").exists()
        )

        published = publish_run_mutation(
            self.run.id,
            MutationPublishRequest(
                target=target,
                approval_context=approval_context,
            ),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        published_again = publish_run_mutation(
            self.run.id,
            MutationPublishRequest(
                target=target,
                approval_context=approval_context,
            ),
            mutation_id=mutation_id,
            auth=auth,
            db=self.db,
        )
        self.assertFalse(published["reused"])
        self.assertTrue(published_again["reused"])
        self.assertEqual(self.db.query(ExecutionPublishReceipt).count(), 1)
        self.assertEqual(self.source_workbook.read_bytes(), source_before)

        detail = run_mutation_detail(
            self.run.id,
            mutation_id,
            auth=auth,
            db=self.db,
        )["mutation"]
        self.assertEqual(detail["status"], "published")
        self.assertEqual(len(detail["publish_receipts"]), 1)
        self.assertEqual(
            detail["publish_receipts"][0]["content_sha256"],
            detail["working_copy"]["content_sha256"],
        )
        actions = {
            item.action
            for item in self.db.query(ExecutionAuditLog)
            .filter(ExecutionAuditLog.resource_id == mutation_id)
            .all()
        }
        self.assertTrue(
            {
                "file_mutation.preflight",
                "file_mutation.copy",
                "file_mutation.write",
                "file_mutation.verify",
                "file_mutation.publish",
                "file_mutation.publish_failed",
            }.issubset(actions)
        )

    def test_public_mutation_is_bound_to_its_run_and_owner(self):
        auth = AuthContext(session=None, user=self.user)
        mutation_id = self.run.input_data["mutation_id"]
        self._run_until_node_status(
            self.run.id,
            "copy",
            {"ready", "running"},
        )
        preflight_run_mutation(
            self.run.id,
            MutationPreflightRequest(
                mutation_id=mutation_id,
                source={
                    "root_id": "special_wool_records",
                    "relative_path": "26X0001.xlsx",
                },
            ),
            auth=auth,
            db=self.db,
        )
        other = ExecutionUser(
            username="other-operator",
            display_name="其他操作员",
            password_hash=hash_password("other-test-password"),
            role="user",
        )
        self.db.add(other)
        self.db.flush()
        bind_user_role(
            self.db,
            other,
            "user",
            created_by_id=self.user.id,
        )
        other_run, _ = create_run(
            self.db,
            workflow=self.workflow,
            actor=other,
            inspection_number="26X0002",
            input_data={
                **self.run.input_data,
                "inspection_number": "26X0002",
            },
            global_data={},
            idempotency_key="other-mutation-run",
        )
        self.db.commit()
        other_auth = AuthContext(session=None, user=other)

        with self.assertRaises(ExecutionApiError) as hidden:
            run_mutation_detail(
                self.run.id,
                mutation_id,
                auth=other_auth,
                db=self.db,
            )
        self.assertEqual(hidden.exception.status_code, 404)

        with self.assertRaises(ExecutionApiError) as wrong_run:
            run_mutation_detail(
                other_run.id,
                mutation_id,
                auth=other_auth,
                db=self.db,
            )
        self.assertEqual(wrong_run.exception.status_code, 404)

        self._run_until_node_status(
            other_run.id,
            "copy",
            {"ready", "running"},
        )
        with self.assertRaises(ExecutionApiError) as conflict:
            preflight_run_mutation(
                other_run.id,
                MutationPreflightRequest(
                    mutation_id=mutation_id,
                    source={
                        "root_id": "special_wool_records",
                        "relative_path": "26X0001.xlsx",
                    },
                ),
                auth=other_auth,
                db=self.db,
            )
        self.assertEqual(conflict.exception.code, "mutation_id_conflict")


if __name__ == "__main__":
    unittest.main()
