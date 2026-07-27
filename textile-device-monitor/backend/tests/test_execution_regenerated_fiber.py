from __future__ import annotations

import tempfile
import unittest
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import xlwt
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.execution import (
    AuthContext,
    router,
    workflow_recommendations,
    workflows,
)
from app.config import settings
from app.database import Base
from app.execution.catalog import (
    _default_definition,
    bind_user_role,
    ensure_default_catalog,
    ensure_default_rbac,
)
from app.execution.engine import (
    claim_human_task,
    claim_next_node,
    create_run,
    execute_claimed_node,
    submit_human_task,
)
from app.execution.errors import ExecutionApiError
from app.execution.models import (
    ExecutionFileIndexEntry,
    ExecutionHumanTask,
    ExecutionIndexJob,
    ExecutionRun,
    ExecutionStorageRoot,
    ExecutionUser,
    ExecutionWorkflow,
    ExecutionWorkflowVersion,
    utcnow,
)
from app.execution.persistence import (
    IndexedFile,
    enqueue_due_index_jobs,
    persist_scan,
    queue_refresh,
    register_persistence_executors,
)
from app.execution.regenerated_fiber import (
    REGENERATED_FIBER_RULES,
    _generic_filename_match,
    _read_profile,
    _safe_candidate_preview,
    catalog_recommendations,
    match_regenerated_fiber_workbooks,
)
from app.execution.security import hash_password
from app.execution.storage import ArtifactRef, FileGateway, StorageRoot
from app.execution.validation import (
    definition_checksum,
    workflow_contract_checksum,
)


COUNT_NODE = "file.regenerated_fiber_count_method"
AREA_NODE = "file.regenerated_fiber_area_method"


class RegeneratedFiberBackendTests(unittest.TestCase):
    def setUp(self):
        register_persistence_executors()
        self.tempdir = tempfile.TemporaryDirectory()
        self.root_path = Path(self.tempdir.name) / "2026-再生纤"
        self.root_path.mkdir(parents=True)
        self.engine = create_engine("sqlite:///:memory:")
        self.Session = sessionmaker(bind=self.engine, autoflush=False)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        ensure_default_rbac(self.db)
        ensure_default_catalog(self.db)
        self.user = ExecutionUser(
            username="regenerated-admin",
            display_name="再生纤测试",
            password_hash=hash_password("regenerated-password"),
            role="admin",
        )
        self.db.add(self.user)
        self.db.flush()
        bind_user_role(
            self.db,
            self.user,
            "admin",
            created_by_id=self.user.id,
        )
        self.root = ExecutionStorageRoot(
            root_id="regenerated_fiber_records",
            name="再生纤原始记录",
            local_path=str(self.root_path),
            access_mode="read",
            category_key="regenerated_fiber",
            is_active=True,
            is_available=True,
            last_scan_finished_at=utcnow(),
        )
        self.db.add(self.root)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        self.tempdir.cleanup()

    def _xlsx(
        self,
        relative_path: str,
        *,
        sheet_name: str,
        values: list[object],
    ) -> Path:
        path = self.root_path / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        for column, value in enumerate(values, start=2):
            sheet.cell(row=14, column=column, value=value)
        workbook.save(path)
        workbook.close()
        return path

    def _xls(
        self,
        relative_path: str,
        *,
        sheet_name: str,
        values: list[object],
    ) -> Path:
        path = self.root_path / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet(sheet_name)
        for column, value in enumerate(values, start=1):
            sheet.write(13, column, value)
        workbook.save(str(path))
        return path

    def _index(
        self,
        path: Path,
        *,
        modified_at: datetime | None = None,
    ) -> ExecutionFileIndexEntry:
        stat = path.stat()
        entry = ExecutionFileIndexEntry(
            storage_root_id=self.root.id,
            relative_path=path.relative_to(self.root_path).as_posix(),
            filename=path.name,
            extension=path.suffix.casefold(),
            file_kind="workbook",
            inspection_number=None,
            size_bytes=stat.st_size,
            modified_at=modified_at
            or datetime.fromtimestamp(stat.st_mtime, timezone.utc),
            fingerprint=f"{stat.st_size}:{stat.st_mtime_ns}",
            metadata_json={
                "metadata_version": 1,
                "parse_status": "deferred",
            },
        )
        self.db.add(entry)
        self.db.flush()
        return entry

    def _execute_one(self):
        node = claim_next_node(
            self.db,
            worker_id="regenerated-worker",
            lease_seconds=30,
        )
        self.assertIsNotNone(node)
        node_id = node.id
        lease_token = node.lease_token
        self.db.commit()
        execute_claimed_node(self.db, node_id, lease_token)
        self.db.commit()
        return node

    def test_saved_value_blank_semantics_for_modern_and_legacy_workbooks(self):
        rule = REGENERATED_FIBER_RULES[COUNT_NODE]
        zero = self._xlsx(
            "260001-zero.xlsx",
            sheet_name=rule.worksheet,
            values=[0],
        )
        self.assertEqual(_read_profile(zero, rule)["status"], "matched")

        false = self._xls(
            "260001-false.xls",
            sheet_name=rule.worksheet,
            values=[False],
        )
        self.assertEqual(_read_profile(false, rule)["status"], "matched")

        whitespace = self._xlsx(
            "260001-space.xlsx",
            sheet_name=rule.worksheet,
            values=["   ", "\t"],
        )
        self.assertEqual(
            _read_profile(whitespace, rule)["status"],
            "content_range_empty",
        )

        # openpyxl does not create a cached result for a newly written formula.
        # data_only=True must therefore regard it as blank rather than execute it.
        formula = self._xlsx(
            "260001-formula.xlsx",
            sheet_name=rule.worksheet,
            values=["=1+1"],
        )
        self.assertEqual(
            _read_profile(formula, rule)["status"],
            "content_range_empty",
        )

    def test_count_method_accepts_numbered_report_sheet_but_not_summary(self):
        rule = REGENERATED_FIBER_RULES[COUNT_NODE]
        numbered = self._xls(
            "260011-numbered.xls",
            sheet_name="根数法报告1",
            values=[1],
        )
        profile = _read_profile(numbered, rule)
        self.assertEqual(profile["status"], "matched")
        self.assertEqual(profile["matched_worksheets"], ["根数法报告1"])

        summary = self._xls(
            "260011-summary.xls",
            sheet_name="根数法报告汇总",
            values=[1],
        )
        self.assertEqual(
            _read_profile(summary, rule)["status"],
            "worksheet_missing",
        )

    def test_same_file_must_satisfy_all_conditions(self):
        empty = self._xlsx(
            "260002-empty.xlsx",
            sheet_name="根数法报告",
            values=[""],
        )
        wrong_sheet = self._xlsx(
            "260002-wrong.xlsx",
            sheet_name="其它报告",
            values=[123],
        )
        self._index(empty)
        self._index(wrong_sheet)
        self.db.commit()

        result = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260002",
        )
        self.assertEqual(result["filename_match_count"], 2)
        self.assertEqual(result["worksheet_match_count"], 1)
        self.assertEqual(result["full_match_count"], 0)
        self.assertEqual(result["candidates"], [])

    def test_no_recent_cutoff_depth_limit_and_final_six_item_limit(self):
        old = datetime.now(timezone.utc) - timedelta(days=400)
        for index in range(8):
            path = self._xlsx(
                f"batch/260003-{index}.xlsx",
                sheet_name="截面统计报告1",
                values=[index],
            )
            self._index(path, modified_at=old + timedelta(minutes=index))
        for index in range(2):
            path = self._xlsx(
                f"260003-invalid-{index}.xlsx",
                sheet_name="其它报告",
                values=[index],
            )
            self._index(path)
        deep = self._xlsx(
            "archive/old/260003-deep.xlsx",
            sheet_name="截面统计报告1",
            values=[1],
        )
        self._index(deep)
        self.db.commit()

        result = match_regenerated_fiber_workbooks(
            self.db,
            node_type=AREA_NODE,
            inspection_number="260003",
            result_limit=100,
        )
        self.assertEqual(result["filename_match_count"], 10)
        self.assertEqual(result["full_match_count"], 8)
        self.assertEqual(len(result["candidates"]), 6)
        self.assertNotIn(
            "archive/old/260003-deep.xlsx",
            {item["relative_path"] for item in result["candidates"]},
        )

    def test_profile_cache_reuses_fingerprint_and_invalidates_on_change(self):
        path = self._xlsx(
            "260004.xlsx",
            sheet_name="根数法报告",
            values=[1],
        )
        entry = self._index(path)
        self.db.commit()

        first = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260004",
        )
        self.assertEqual(first["full_match_count"], 1)
        self.db.commit()
        with patch(
            "app.execution.regenerated_fiber._read_profile",
            side_effect=AssertionError("cache was not reused"),
        ):
            second = match_regenerated_fiber_workbooks(
                self.db,
                node_type=COUNT_NODE,
                inspection_number="260004",
            )
        self.assertEqual(second["full_match_count"], 1)

        self._xlsx(
            "260004.xlsx",
            sheet_name="根数法报告",
            values=[""],
        )
        stat = path.stat()
        entry = self.db.get(ExecutionFileIndexEntry, entry.id)
        entry.size_bytes = stat.st_size
        entry.fingerprint = f"{stat.st_size}:{stat.st_mtime_ns}"
        self.db.commit()
        changed = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260004",
        )
        self.assertEqual(changed["full_match_count"], 0)

    def test_incomplete_or_overwide_query_never_opens_shared_workbooks(self):
        path = self._xlsx(
            "260008.xlsx",
            sheet_name="根数法报告",
            values=[1],
        )
        self._index(path)
        self.db.commit()
        with patch(
            "app.execution.regenerated_fiber._read_profile",
            side_effect=AssertionError("incomplete query opened a workbook"),
        ):
            incomplete = match_regenerated_fiber_workbooks(
                self.db,
                node_type=COUNT_NODE,
                inspection_number="260",
            )
        self.assertEqual(incomplete["query_state"], "incomplete")
        self.assertEqual(incomplete["filename_match_count"], 1)
        self.assertEqual(incomplete["full_match_count"], 0)

        for index in range(51):
            candidate = self.root_path / f"260009-{index}.xlsx"
            candidate.write_bytes(b"indexed only")
            self._index(candidate)
        self.db.commit()
        with patch(
            "app.execution.regenerated_fiber._read_profile",
            side_effect=AssertionError("overwide query opened a workbook"),
        ):
            overwide = match_regenerated_fiber_workbooks(
                self.db,
                node_type=COUNT_NODE,
                inspection_number="260009",
            )
        self.assertEqual(overwide["query_state"], "too_many_matches")
        self.assertEqual(overwide["filename_match_count"], 51)
        self.assertEqual(overwide["full_match_count"], 0)

    def test_transient_profile_failures_are_retried_not_cached(self):
        path = self._xlsx(
            "260010.xlsx",
            sheet_name="根数法报告",
            values=[1],
        )
        entry = self._index(path)
        self.db.commit()
        transient = {
            "rule_version": 1,
            "worksheet": "根数法报告",
            "cell_range": "B14:J14",
            "worksheet_exists": False,
            "content_range_nonempty": False,
            "status": "read_failed",
            "error_type": "PermissionError",
        }
        with patch(
            "app.execution.regenerated_fiber._read_profile",
            return_value=transient,
        ) as reader:
            match_regenerated_fiber_workbooks(
                self.db,
                node_type=COUNT_NODE,
                inspection_number="260010",
            )
            match_regenerated_fiber_workbooks(
                self.db,
                node_type=COUNT_NODE,
                inspection_number="260010",
            )
        self.assertEqual(reader.call_count, 2)
        self.db.refresh(entry)
        self.assertNotIn(
            COUNT_NODE,
            " ".join(
                (entry.metadata_json.get("workbook_profiles") or {}).keys()
            ),
        )

    def test_malformed_and_unsupported_workbooks_are_diagnostic_not_matches(self):
        malformed = self.root_path / "260005-broken.xlsx"
        malformed.write_bytes(b"not a workbook")
        unsupported = self.root_path / "260005-text.csv"
        unsupported.write_text("value", encoding="utf-8")
        self._index(malformed)
        self._index(unsupported)
        self.db.commit()

        result = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260005",
        )
        self.assertEqual(result["filename_match_count"], 2)
        self.assertEqual(result["worksheet_match_count"], 0)
        self.assertEqual(result["full_match_count"], 0)

    def test_specialized_executor_reports_match_counts_when_empty(self):
        path = self._xlsx(
            "260006.xlsx",
            sheet_name="根数法报告",
            values=[""],
        )
        self._index(path)
        self.db.commit()
        workflow = (
            self.db.query(ExecutionWorkflow)
            .filter_by(slug="regenerated-fiber-count-method")
            .one()
        )
        run, _ = create_run(
            self.db,
            workflow=workflow,
            actor=self.user,
            inspection_number="260006",
            input_data={},
            global_data={},
            idempotency_key="regenerated-no-match",
        )
        self.db.commit()
        self._execute_one()
        query_node = self._execute_one()
        self.db.refresh(query_node)
        self.assertEqual(query_node.status, "failed")
        self.assertEqual(query_node.error_code, "matching_workbook_not_found")
        self.assertIn("文件名命中 1", query_node.error_message)
        self.assertIn("工作表命中 1", query_node.error_message)
        self.assertIn("完整命中 0", query_node.error_message)

    def test_area_default_workflow_runs_through_human_selection(self):
        path = self._xlsx(
            "260144785-面积法.xlsx",
            sheet_name="截面统计报告1",
            values=[10],
        )
        entry = self._index(path)
        self.db.commit()
        workflow = (
            self.db.query(ExecutionWorkflow)
            .filter_by(slug="regenerated-fiber-area-method")
            .one()
        )
        run, duplicate = create_run(
            self.db,
            workflow=workflow,
            actor=self.user,
            inspection_number="260144785",
            input_data={},
            global_data={},
            idempotency_key="regenerated-area-run",
        )
        self.db.commit()
        self.assertFalse(duplicate)

        self._execute_one()  # start
        self._execute_one()  # matcher
        self._execute_one()  # create human task
        task = (
            self.db.query(ExecutionHumanTask)
            .filter_by(run_id=run.id)
            .one()
        )
        offered = task.node_run.input_data["candidates"]
        self.assertEqual([item["id"] for item in offered], [entry.id])
        task = claim_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            actor=self.user,
        )
        self.db.commit()
        submit_human_task(
            self.db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"selected_files": [entry.id]},
            actor=self.user,
        )
        self.db.commit()
        self._execute_one()  # aggregate
        self._execute_one()  # end
        self.db.refresh(run)
        self.assertEqual(run.status, "completed")
        self.assertEqual(
            run.output_data["result"]["selection"]["selected_files"][0]["id"],
            entry.id,
        )

    def test_recommendations_rank_full_match_and_apply_category_soft_score(self):
        area = self._xlsx(
            "7月/260144785.xlsx",
            sheet_name="截面统计报告1",
            values=[1],
        )
        self._index(area)
        self.db.commit()
        items, cache_updated = catalog_recommendations(
            self.db,
            inspection_number="260144785",
            preferred_categories=["regenerated_fiber"],
            include_hidden=False,
        )
        self.assertTrue(cache_updated)
        area_workflow = (
            self.db.query(ExecutionWorkflow)
            .filter_by(slug="regenerated-fiber-area-method")
            .one()
        )
        self.assertEqual(items[0]["workflow_id"], area_workflow.id)
        self.assertEqual(items[0]["state"], "full_match")
        self.assertEqual(items[0]["score"], 5)
        self.assertEqual(
            items[0]["matched_conditions"],
            [
                "category",
                "source_root",
                "filename",
                "worksheet",
                "content_range",
            ],
        )
        self.assertEqual(items[0]["candidate_count"], 1)
        self.assertEqual(
            items[0]["candidate_preview"],
            {
                "name": "260144785.xlsx",
                "relative_path": "7月/260144785.xlsx",
                "suffix": ".xlsx",
            },
        )
        self.assertNotIn(
            str(self.root_path),
            items[0]["candidate_preview"]["relative_path"],
        )

        no_number, _ = catalog_recommendations(
            self.db,
            inspection_number="",
            preferred_categories=["regenerated_fiber"],
            include_hidden=False,
        )
        first_workflow = self.db.get(
            ExecutionWorkflow,
            no_number[0]["workflow_id"],
        )
        self.assertEqual(first_workflow.category.key, "regenerated_fiber")
        self.assertIsNone(no_number[0]["candidate_preview"])

    def test_generic_preview_is_stable_and_absolute_paths_are_not_exposed(self):
        older = self._xlsx(
            "6月/260177-old.xlsx",
            sheet_name="其它",
            values=[],
        )
        newer = self._xlsx(
            "7月/260177-new.xlsm",
            sheet_name="其它",
            values=[],
        )
        self._index(
            older,
            modified_at=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        self._index(
            newer,
            modified_at=datetime(2026, 7, 1, tzinfo=timezone.utc),
        )
        self.db.commit()

        count, preview = _generic_filename_match(
            self.db,
            root_ids=[self.root.root_id],
            inspection_number="260177",
        )
        self.assertEqual(count, 2)
        self.assertEqual(
            preview,
            {
                "name": "260177-new.xlsm",
                "relative_path": "7月/260177-new.xlsm",
                "suffix": ".xlsm",
            },
        )
        self.assertEqual(
            _safe_candidate_preview(
                name="260177-new.xlsm",
                relative_path=(
                    r"C:\records\7月\260177-new.xlsm"
                ),
                suffix="xlsm",
            ),
            {
                "name": "260177-new.xlsm",
                "relative_path": "260177-new.xlsm",
                "suffix": ".xlsm",
            },
        )
        self.assertEqual(
            _safe_candidate_preview(
                name="260177-new.xlsm",
                relative_path=(
                    r"\\server\share\260177-new.xlsm"
                ),
                suffix=".xlsm",
            )["relative_path"],
            "260177-new.xlsm",
        )

    def test_recommendation_reports_unavailable_and_api_repeats_categories(self):
        self.root.is_available = False
        self.db.commit()
        auth = AuthContext(session=None, user=self.user)
        response = workflow_recommendations(
            inspection_number="260000",
            preferred_categories=["regenerated_fiber", "regenerated_fiber"],
            auth=auth,
            db=self.db,
        )
        self.assertEqual(
            response["preferred_categories"],
            ["regenerated_fiber"],
        )
        self.assertEqual(response["query_state"], "complete")
        specialized = [
            item
            for item in response["items"]
            if self.db.get(ExecutionWorkflow, item["workflow_id"]).slug
            in {
                "regenerated-fiber-count-method",
                "regenerated-fiber-area-method",
            }
        ]
        self.assertTrue(specialized)
        self.assertTrue(
            all(item["state"] == "index_unavailable" for item in specialized)
        )
        paths = {
            (method, route.path)
            for route in router.routes
            for method in route.methods
        }
        self.assertIn(
            ("GET", "/execution/v1/catalog/recommendations"),
            paths,
        )

    def test_failed_and_partial_index_jobs_are_not_reported_ready(self):
        failed = ExecutionIndexJob(
            storage_root_id=self.root.id,
            status="failed",
            errors=["share unavailable"],
            finished_at=utcnow(),
        )
        self.db.add(failed)
        self.db.commit()
        result = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260011",
        )
        self.assertEqual(result["index_state"], "failed")

        failed.status = "completed_with_errors"
        self.root.last_scan_error = "one directory was unavailable"
        self.db.commit()
        degraded = match_regenerated_fiber_workbooks(
            self.db,
            node_type=COUNT_NODE,
            inspection_number="260011",
        )
        self.assertEqual(degraded["index_state"], "degraded")

    def test_auto_index_whitelist_does_not_disable_manual_refresh(self):
        other_path = Path(self.tempdir.name) / "2026-特种毛"
        other_path.mkdir()
        other = ExecutionStorageRoot(
            root_id="special_wool_records",
            name="特种毛原始记录",
            local_path=str(other_path),
            access_mode="read",
            category_key="special_wool",
            is_active=True,
            is_available=True,
        )
        self.root.last_scan_finished_at = None
        self.db.add(other)
        self.db.commit()
        original = settings.EXECUTION_AUTO_INDEX_ROOT_IDS
        settings.EXECUTION_AUTO_INDEX_ROOT_IDS = "regenerated_fiber_records"
        try:
            created = enqueue_due_index_jobs(
                self.db,
                interval_seconds=300,
            )
            self.assertEqual(created, 1)
            automatic = self.db.query(ExecutionIndexJob).one()
            self.assertEqual(automatic.storage_root_id, self.root.id)

            manual, duplicate = queue_refresh(
                self.db,
                root_id="special_wool_records",
                actor=self.user,
            )
            self.assertFalse(duplicate)
            self.assertEqual(manual.storage_root_id, other.id)
        finally:
            settings.EXECUTION_AUTO_INDEX_ROOT_IDS = original

    def test_system_deprecated_workflows_are_hidden_even_from_admin(self):
        admin_auth = AuthContext(session=None, user=self.user)
        before = workflows(
            categories=[],
            query=None,
            auth=admin_auth,
            db=self.db,
        )
        test_workflow = (
            self.db.query(ExecutionWorkflow)
            .filter_by(slug="system-controlled-xlsx-write-test")
            .one()
        )
        self.assertIn(
            test_workflow.id,
            {item["id"] for item in before["items"]},
        )
        published = test_workflow.versions[-1]
        published.capabilities = {
            **(published.capabilities or {}),
            "system_deprecated": True,
        }
        self.db.commit()

        after = workflows(
            categories=[],
            query=None,
            auth=admin_auth,
            db=self.db,
        )
        recommendations, _ = catalog_recommendations(
            self.db,
            inspection_number="260012",
            preferred_categories=[],
            include_hidden=True,
        )
        self.assertNotIn(
            test_workflow.id,
            {item["id"] for item in after["items"]},
        )
        self.assertNotIn(
            test_workflow.id,
            {item["workflow_id"] for item in recommendations},
        )

    def test_regenerated_background_index_never_opens_workbook(self):
        path = self._xlsx(
            "260007.xlsx",
            sheet_name="根数法报告",
            values=[1],
        )
        stat = path.stat()
        indexed = IndexedFile(
            ref=ArtifactRef(
                "regenerated_fiber_records",
                "260007.xlsx",
            ),
            name=path.name,
            suffix=path.suffix,
            size=stat.st_size,
            modified_ns=stat.st_mtime_ns,
            category="regenerated_fiber",
        )
        gateway = FileGateway(
            [
                StorageRoot(
                    root_id="regenerated_fiber_records",
                    path=self.root_path,
                )
            ]
        )
        with patch(
            "app.execution.persistence.extract_index_metadata",
            side_effect=AssertionError("workbook should not be opened"),
        ):
            added, updated, removed = persist_scan(
                self.db,
                root=self.root,
                files=[indexed],
                scan_complete=True,
                gateway=gateway,
            )
        self.assertEqual((added, updated, removed), (1, 0, 0))
        self.db.flush()
        entry = (
            self.db.query(ExecutionFileIndexEntry)
            .filter_by(filename="260007.xlsx")
            .one()
        )
        self.assertEqual(entry.metadata_json["parse_status"], "deferred")


class RegeneratedFiberCatalogMigrationTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        self.Session = sessionmaker(bind=self.engine, autoflush=False)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        ensure_default_catalog(self.db)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def _insert_legacy(self, *, edited: bool) -> ExecutionWorkflow:
        category = next(
            item
            for item in self.db.query(ExecutionWorkflow).all()
            if item.category.key == "regenerated_fiber"
        ).category
        definition = _default_definition(
            slug="regenerated-fiber-source-selection",
            name="再生纤原始资料发现与选择",
            category_key="regenerated_fiber",
            root_id="regenerated_fiber_records",
        )
        workflow = ExecutionWorkflow(
            slug="regenerated-fiber-source-selection",
            category_id=category.id,
            name="再生纤原始资料发现与选择",
            description="legacy",
            draft_definition=deepcopy(definition),
            draft_revision=2 if edited else 1,
            published_version_number=1,
            capabilities={"read": True, "write": False},
            required_input_count=1,
            is_enabled=True,
        )
        capabilities = {"read": True, "write": False}
        self.db.add(workflow)
        self.db.flush()
        self.db.add(
            ExecutionWorkflowVersion(
                workflow_id=workflow.id,
                version_number=1,
                schema_version="1.0",
                definition=deepcopy(definition),
                checksum=definition_checksum(definition),
                capabilities=capabilities,
                contract_checksum=workflow_contract_checksum(
                    definition,
                    capabilities,
                ),
                release_note="legacy",
            )
        )
        self.db.commit()
        return workflow

    def test_default_methods_seed_idempotently(self):
        ensure_default_catalog(self.db)
        ensure_default_catalog(self.db)
        self.db.commit()
        for slug in (
            "regenerated-fiber-count-method",
            "regenerated-fiber-area-method",
        ):
            workflows = (
                self.db.query(ExecutionWorkflow)
                .filter_by(slug=slug)
                .all()
            )
            self.assertEqual(len(workflows), 1)
            self.assertEqual(len(workflows[0].versions), 1)

    def test_untouched_legacy_default_is_hidden_but_edited_copy_is_preserved(self):
        legacy = self._insert_legacy(edited=False)
        ensure_default_catalog(self.db)
        self.db.commit()
        self.db.refresh(legacy)
        self.assertFalse(legacy.is_enabled)
        self.assertEqual(legacy.availability_code, "system_replaced")
        self.assertEqual(legacy.published_version_number, 2)
        self.assertTrue(legacy.versions[-1].capabilities["hidden"])

        # A separate database verifies the opposite branch without reusing the
        # unique legacy slug.
        other_engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(other_engine)
        OtherSession = sessionmaker(bind=other_engine, autoflush=False)
        other = OtherSession()
        try:
            ensure_default_catalog(other)
            other.commit()
            category = (
                other.query(ExecutionWorkflow)
                .filter_by(slug="regenerated-fiber-count-method")
                .one()
                .category
            )
            definition = _default_definition(
                slug="regenerated-fiber-source-selection",
                name="再生纤原始资料发现与选择",
                category_key="regenerated_fiber",
                root_id="regenerated_fiber_records",
            )
            edited = ExecutionWorkflow(
                slug="regenerated-fiber-source-selection",
                category_id=category.id,
                name="管理员修改过的再生纤流程",
                draft_definition=deepcopy(definition),
                draft_revision=2,
                published_version_number=1,
                capabilities={"read": True},
                required_input_count=1,
                is_enabled=True,
            )
            other.add(edited)
            other.flush()
            capabilities = {"read": True}
            other.add(
                ExecutionWorkflowVersion(
                    workflow_id=edited.id,
                    version_number=1,
                    schema_version="1.0",
                    definition=deepcopy(definition),
                    checksum=definition_checksum(definition),
                    capabilities=capabilities,
                    contract_checksum=workflow_contract_checksum(
                        definition,
                        capabilities,
                    ),
                )
            )
            other.commit()
            ensure_default_catalog(other)
            other.commit()
            other.refresh(edited)
            self.assertTrue(edited.is_enabled)
            self.assertEqual(edited.draft_revision, 2)
            self.assertEqual(edited.published_version_number, 1)
        finally:
            other.close()
            Base.metadata.drop_all(other_engine)
            other_engine.dispose()


if __name__ == "__main__":
    unittest.main()
