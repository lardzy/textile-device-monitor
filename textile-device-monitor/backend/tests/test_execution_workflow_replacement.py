"""HTTP/Worker acceptance for P2 migration and local workflow replacement."""

from copy import deepcopy
from datetime import datetime, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.execution import router as v1_router
from app.api.execution_v2 import router as v2_router
from app.config import settings
from app.database import Base, get_db
from app.execution.catalog import (
    bind_user_role,
    ensure_default_catalog,
    ensure_default_rbac,
)
from app.execution.engine import claim_next_node, execute_claimed_node
from app.execution.models import (
    ExecutionAuditLog,
    ExecutionFileIndexEntry,
    ExecutionHumanTask,
    ExecutionFileMutation,
    ExecutionPublishReceipt,
    ExecutionNodeRun,
    ExecutionRun,
    ExecutionStorageRoot,
    ExecutionUser,
    ExecutionWorkflow,
    ExecutionWorkflowRelease,
    ExecutionWorkflowVersion,
)
from app.execution.project_rules import ensure_default_project_rules
from app.execution.release_v2 import _release_digest
from app.execution.security import SESSION_COOKIE, create_session, hash_password
from app.execution.validation import definition_checksum
from app.execution.worker import ExecutionWorker
from app.execution.worker_state import record_worker_heartbeat


CHOICES = [
    "electron-source-selection",
    "hemp-cotton-source-selection",
    "special-wool-source-selection",
]
CANARY = "system-controlled-xlsx-write-test"


@pytest.fixture
def environment(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXECUTION_CONTRACT_MODE", "enforced")
    monkeypatch.setattr(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_publish")
    monkeypatch.setattr(settings, "EXECUTION_RUNTIME_ROOT", str(tmp_path / "runtime"))
    monkeypatch.setattr(settings, "EXECUTION_PUBLISH_ROOT", str(tmp_path / "publish"))
    engine = create_engine(
        f"sqlite:///{tmp_path / 'replacement_test.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False)
    db = Session()
    ensure_default_rbac(db)
    ensure_default_catalog(db)
    ensure_default_project_rules(db)
    admin = ExecutionUser(
        username="replacement-admin",
        display_name="接替验收",
        role="admin",
        password_hash=hash_password("test-password"),
    )
    operator = ExecutionUser(
        username="replacement-user",
        display_name="操作员",
        role="user",
        password_hash=hash_password("test-password"),
    )
    db.add_all([admin, operator])
    db.flush()
    for user in (admin, operator):
        bind_user_role(db, user, user.role, created_by_id=admin.id)
    roots = {}
    for root_id in (
        "electron_microscopy_records",
        "hemp_cotton_records",
        "special_wool_records",
        "execution_staging",
        "execution_publish",
    ):
        path = tmp_path / root_id
        path.mkdir()
        root = ExecutionStorageRoot(
            root_id=root_id,
            name=root_id,
            local_path=str(path),
            access_mode=(
                "publish"
                if root_id.endswith("publish")
                else "write" if root_id.endswith("staging") else "read"
            ),
            is_active=True,
            is_available=True,
            scan_generation=1,
        )
        db.add(root)
        roots[root_id] = root
    worker = ExecutionWorker(worker_id="replacement-worker")
    record_worker_heartbeat(
        db, worker_id=worker.worker_id, capability_document=worker.capability_document
    )
    _session, token, csrf = create_session(db, admin)
    _user_session, user_token, user_csrf = create_session(db, operator)
    db.commit()
    app = FastAPI()
    app.include_router(v1_router, prefix="/api")
    app.include_router(v2_router, prefix="/api")

    def session_dependency():
        with Session() as request_db:
            yield request_db

    app.dependency_overrides[get_db] = session_dependency
    with TestClient(app) as client, TestClient(app) as user_client:
        client.cookies.set(SESSION_COOKIE, token)
        client.headers["X-CSRF-Token"] = csrf
        user_client.cookies.set(SESSION_COOKIE, user_token)
        user_client.headers["X-CSRF-Token"] = user_csrf
        yield SimpleNamespace(
            db=db,
            Session=Session,
            client=client,
            user_client=user_client,
            admin=admin,
            roots=roots,
            path=tmp_path,
            worker=worker,
        )
    db.close()
    Base.metadata.drop_all(engine)
    engine.dispose()


def request(env, method, path, payload=None, *, status=200):
    response = env.client.request(method, f"/api/execution/{path}", json=payload)
    assert response.status_code == status, response.text
    return response.json()


def source_workflow(env, slug):
    return env.db.query(ExecutionWorkflow).filter_by(slug=slug).one()


def preview(env, slug, *, source="published", profile="native_p2"):
    workflow = source_workflow(env, slug)
    return request(
        env,
        "POST",
        "v2/migrations/v1/preview",
        {
            "workflow_id": workflow.id,
            "source": source,
            "target_profile": profile,
            "target_slug": slug + "-v2",
        },
    )


def stage(env, document):
    content = request(
        env, "POST", "v2/workflow-releases/preflight", {"document": document}
    )
    assert content["content_valid"], content
    release = request(
        env,
        "POST",
        "v2/workflow-releases/apply",
        {"preflight_token": content["preflight_token"]},
        status=201,
    )
    request(
        env,
        "PUT",
        f"v2/workflow-releases/{release['id']}/deployment-binding",
        {
            "environment": "default",
            "expected_revision": 0,
            "root_bindings": [
                {
                    "slot": slot["slot_id"],
                    "storage_root_id": env.roots[slot["slot_id"]].id,
                    "role": slot["access"],
                }
                for slot in document["resources"]["root_slots"]
            ],
        },
    )
    return release


def publish(env, release, source=None):
    report = request(
        env,
        "POST",
        f"v2/workflow-releases/{release['id']}/preflight",
        {"replacement_source": source} if source else {},
    )
    assert report["publish_ready"], report
    return request(
        env,
        "POST",
        f"v2/workflow-releases/{release['id']}/publish",
        {"preflight_token": report["preflight_token"]},
    )


def replacement_payload(state, reason="隔离环境验收"):
    return {
        "expected_source_revision": state["source"]["draft_revision"],
        "expected_target_revision": state["target"]["draft_revision"],
        "expected_source_version": state["source"]["version_number"],
        "expected_target_version": state["target"]["version_number"],
        "reason": reason,
        **(
            {"activation_audit_id": state["last_action"]["id"]}
            if state["last_action"]
            else {}
        ),
    }


def switch(env, workflow_id, action, *, status=200):
    state = request(env, "GET", f"v2/workflows/{workflow_id}/replacement")
    return request(
        env,
        "POST",
        f"v2/workflows/{workflow_id}/replacement/{action}",
        replacement_payload(state),
        status=status,
    )


def run(env, workflow_id, *, inputs=None, key=None):
    return request(
        env,
        "POST",
        "v1/runs",
        {
            "workflow_id": workflow_id,
            "inspection_number": "P2-001",
            "input_data": inputs or {},
            "global_data": {},
            "idempotency_key": key or str(uuid4()),
        },
        status=201,
    )["run"]


def drain(env):
    """Run actual Worker claim/lease/execute until the DAG suspends or ends."""
    for _ in range(30):
        with env.Session() as db:
            claimed = claim_next_node(db, worker_id=env.worker.worker_id)
            if claimed is None:
                return
            node_id, lease_token = claimed.id, claimed.lease_token
            db.commit()
            execute_claimed_node(db, node_id, lease_token)
            db.commit()
    raise AssertionError("DAG did not quiesce")


def index_file(env, root_id, *, suffix=".xlsx"):
    path = env.path / root_id / ("P2-001" + suffix)
    path.write_bytes(b"indexed selection fixture")
    stat = path.stat()
    row = ExecutionFileIndexEntry(
        storage_root_id=env.roots[root_id].id,
        relative_path=path.name,
        filename=path.name,
        extension=suffix,
        file_kind="workbook" if suffix == ".xlsx" else "image",
        inspection_number="P2-001",
        group_key="P2-001",
        size_bytes=stat.st_size,
        modified_at=datetime.now(timezone.utc),
        fingerprint=f"{stat.st_size}:{stat.st_mtime_ns}",
        scan_generation=1,
    )
    env.db.add(row)
    env.db.commit()
    return row.id


def waiting_task(env, run_id):
    env.db.expire_all()
    task = env.db.query(ExecutionHumanTask).filter_by(run_id=run_id).one_or_none()
    assert task is not None, [
        (node.node_id, node.status, node.error_code, node.error_message)
        for node in env.db.query(ExecutionNodeRun).filter_by(run_id=run_id).all()
    ]
    return request(env, "GET", f"v1/human-tasks/{task.id}")["task"]


@pytest.mark.parametrize("profile", ["compat_v1", "native_p2"])
def test_parked_nodes_are_excluded_before_dependency_resolution(environment, profile):
    env = environment
    source = source_workflow(env, CHOICES[1])
    draft = deepcopy(source.draft_definition)
    draft["nodes"].append(
        {
            "id": "parked",
            "type": "not.installed",
            "name": "停放",
            "disabled": True,
            "config": {},
        }
    )
    draft["edges"].append({"id": "parked-edge", "source": "start", "target": "parked"})
    source.draft_definition = draft
    source.draft_revision += 1
    env.db.commit()
    first = preview(env, source.slug, source="draft", profile=profile)
    second = preview(env, source.slug, source="draft", profile=profile)
    assert first == second
    assert first["content_valid"]
    assert first["candidate"]["migration"]["source_digest"] == definition_checksum(
        draft
    )
    assert first["diff"]["excluded_node_ids"] == ["parked"]
    assert first["diff"]["excluded_edge_ids"] == ["parked-edge"]
    assert not any(
        node["id"] == "parked" for node in first["candidate"]["definition"]["nodes"]
    )
    env.db.refresh(source)
    assert source.draft_definition == draft
    assert first["replacement_source"] is None
    compatible = request(
        env,
        "POST",
        "v2/migrations/v1/preview",
        {"workflow_id": source.id, "source": "draft", "target_profile": profile},
    )
    assert compatible["candidate"]["release"]["slug"] == source.slug
    assert (
        compatible["candidate"]["integrity"]["digest"]
        != first["candidate"]["integrity"]["digest"]
    )


def test_parked_required_node_returns_migration_blocker(environment):
    env = environment
    source = source_workflow(env, CHOICES[1])
    draft = deepcopy(source.draft_definition)
    next(node for node in draft["nodes"] if node["id"] == "select")["disabled"] = True
    source.draft_definition = draft
    source.draft_revision += 1
    env.db.commit()
    result = preview(env, source.slug, source="draft")
    assert not result["publish_ready"]
    assert result["diff"]["excluded_node_ids"] == ["select"]
    assert any(
        item["code"] == "migration_required_node_inactive"
        and item["node_id"] == "select"
        for item in result["blockers"]
    )
    env.db.refresh(source)
    assert source.draft_definition == draft


@pytest.mark.parametrize("slug", CHOICES)
def test_selection_candidate_replaces_runs_and_reverts_through_http_and_worker(
    environment, slug, monkeypatch
):
    env = environment
    monkeypatch.setattr(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_human")
    migration = preview(env, slug)
    source = source_workflow(env, slug)
    prior_enabled = source.is_enabled
    original_definition = deepcopy(source.draft_definition)
    original_versions = [
        (v.id, v.checksum, v.contract_checksum) for v in source.versions
    ]
    root_id = migration["candidate"]["resources"]["root_slots"][0]["slot_id"]
    entry_id = index_file(
        env, root_id, suffix=".bmp" if slug == CHOICES[0] else ".xlsx"
    )
    old_run = run(env, source.id) if prior_enabled else None
    if old_run:
        drain(env)
        old_task = waiting_task(env, old_run["id"])
    release = publish(
        env, stage(env, migration["candidate"]), migration["replacement_source"]
    )
    target_id = release["workflow_id"]
    assert target_id != source.id
    assert release["migration_source"] == migration["replacement_source"]
    state = request(env, "GET", f"v2/workflows/{target_id}/replacement")
    assert not state["target"]["is_enabled"] and state["status"] == "pending"
    normal = request(env, "GET", "v1/workflows")["items"]
    assert target_id not in {item["id"] for item in normal}
    assert (
        request(
            env,
            "POST",
            "v1/runs",
            {
                "workflow_id": target_id,
                "inspection_number": "P2-001",
                "idempotency_key": "disabled",
            },
            status=409,
        )["code"]
        == "workflow_disabled"
    )
    activated = switch(env, target_id, "activate")
    assert activated["source"]["archived_at"] and not activated["source"]["is_enabled"]
    assert activated["target"]["is_enabled"]
    assert (
        request(
            env,
            "POST",
            "v1/runs",
            {
                "workflow_id": source.id,
                "inspection_number": "P2-001",
                "idempotency_key": "archived",
            },
            status=409,
        )["code"]
        == "workflow_archived"
    )
    for path, body in (
        (
            f"v1/workflows/{source.id}/test",
            {"inspection_number": "P2-001", "idempotency_key": "archived-test"},
        ),
        (
            f"v1/workflows/{source.id}/publish",
            {"revision": activated["source"]["draft_revision"]},
        ),
        (
            f"v1/workflows/{source.id}/draft",
            {
                "revision": activated["source"]["draft_revision"],
                "definition": original_definition,
            },
        ),
    ):
        assert (
            request(
                env,
                "PUT" if path.endswith("/draft") else "POST",
                path,
                body,
                status=409,
            )["code"]
            == "workflow_archived"
        )
    ordinary = env.user_client.get("/api/execution/v1/workflows").json()["items"]
    assert target_id in {item["id"] for item in ordinary} and source.id not in {
        item["id"] for item in ordinary
    }
    detail = env.user_client.get(f"/api/execution/v1/workflows/{source.id}")
    assert detail.status_code == 200
    assert detail.json()["replacement_workflow"]["id"] == target_id
    admin_items = request(env, "GET", "v1/workflows?include_archived=true")["items"]
    assert {source.id, target_id}.issubset({item["id"] for item in admin_items})
    recommendations = request(
        env, "GET", "v1/catalog/recommendations?inspection_number=P2-001"
    )["items"]
    assert source.id not in {item["workflow_id"] for item in recommendations}
    env.db.expire_all()
    ensure_default_catalog(env.db)
    env.db.commit()
    env.db.refresh(source)
    assert source.archived_at and not source.is_enabled
    assert source.draft_definition == original_definition
    assert [
        (v.id, v.checksum, v.contract_checksum) for v in source.versions
    ] == original_versions
    if old_run:
        request(
            env,
            "POST",
            f"v1/human-tasks/{old_task['id']}/submit",
            {
                "revision": old_task["revision"],
                "data": {"selected_files": [entry_id], "primary_file_id": entry_id},
            },
        )
        drain(env)
        assert request(env, "GET", f"v1/runs/{old_run['id']}")["status"] == "completed"
    current_run = run(env, target_id)
    drain(env)
    task = waiting_task(env, current_run["id"])
    assert task["renderer_contract"]["capability"] == "human.select"
    assert (
        task["renderer_contract"]["contract_digest"]
        == "e355eec60016099fda22268ccd7a02b805ab77682cbee8e6b96c9930f9af9a41"
    )
    assert task["renderer_contract"]["node_contract_digest"] == task["contract_digest"]
    reverted = switch(env, target_id, "revert")
    assert reverted["source"]["is_enabled"] == prior_enabled
    assert (
        reverted["source"]["archived_at"] is None
        and not reverted["target"]["is_enabled"]
    )
    detail = request(env, "GET", f"v1/human-tasks/{task['id']}")
    candidates = [item["id"] for item in detail["node_run"]["input_data"]["items"]]
    request(
        env,
        "POST",
        f"v1/human-tasks/{task['id']}/submit",
        {
            "revision": task["revision"],
            "data": {"selected_ids": [candidates[0]], "primary_id": candidates[0]},
        },
    )
    drain(env)
    assert request(env, "GET", f"v1/runs/{current_run['id']}")["status"] == "completed"
    exported = request(env, "GET", f"v2/workflows/{target_id}/versions/1/export")
    assert exported == migration["candidate"]
    assert source.id not in str(exported) and target_id not in str(exported)
    env.db.expire_all()
    logs = (
        env.db.query(ExecutionAuditLog)
        .filter(ExecutionAuditLog.action.like("workflow_replacement.%"))
        .all()
    )
    assert len(logs) == 2
    assert logs[1].details["activation_audit_id"] == logs[0].id


def test_source_evidence_and_revision_are_frozen_until_publish(environment):
    env = environment
    migration = preview(env, CHOICES[1])
    release = stage(env, migration["candidate"])
    report = request(
        env,
        "POST",
        f"v2/workflow-releases/{release['id']}/preflight",
        {"replacement_source": migration["replacement_source"]},
    )
    source = source_workflow(env, CHOICES[1])
    source.draft_revision += 1
    env.db.commit()
    rejected = request(
        env,
        "POST",
        f"v2/workflow-releases/{release['id']}/publish",
        {"preflight_token": report["preflight_token"]},
        status=409,
    )
    assert rejected["code"] == "replacement_source_changed"
    assert (
        env.db.query(ExecutionWorkflow).filter_by(slug=CHOICES[1] + "-v2").count() == 0
    )


def test_new_release_can_refresh_source_version_and_revert_uses_activation_evidence(
    environment,
):
    env = environment
    first = preview(env, CHOICES[1])
    first_release = publish(
        env, stage(env, first["candidate"]), first["replacement_source"]
    )
    old_report = request(
        env, "POST", f"v2/workflow-releases/{first_release['id']}/preflight", {}
    )
    assert old_report["publish_ready"]
    target_id = first_release["workflow_id"]
    source = source_workflow(env, CHOICES[1])
    draft = deepcopy(source.draft_definition)
    draft["nodes"][0]["name"] = "来源流程更新后的开始"
    source.draft_definition = draft
    source.draft_revision += 1
    env.db.commit()
    request(
        env,
        "POST",
        f"v1/workflows/{source.id}/publish",
        {"revision": source.draft_revision},
    )
    assert (
        request(
            env,
            "POST",
            f"v2/workflow-releases/{first_release['id']}/publish",
            {"preflight_token": old_report["preflight_token"]},
            status=409,
        )["code"]
        == "replacement_source_changed"
    )
    assert (
        switch(env, target_id, "activate", status=409)["code"]
        == "replacement_source_changed"
    )
    current = preview(env, CHOICES[1])
    later = stage(env, current["candidate"])
    assert later["migration_source"] == current["replacement_source"]
    changed_pair = {
        **current["replacement_source"],
        "workflow_id": source_workflow(env, CHOICES[2]).id,
    }
    invalid = request(
        env,
        "POST",
        f"v2/workflow-releases/{later['id']}/preflight",
        {"replacement_source": changed_pair},
    )
    assert not invalid["publish_ready"]
    next_release = publish(env, later, current["replacement_source"])
    assert next_release["workflow_id"] == target_id
    assert (
        request(env, "GET", f"v2/workflow-releases/{first_release['id']}")[
            "migration_source"
        ]
        == first["replacement_source"]
    )
    activated = switch(env, target_id, "activate")
    assert (
        activated["source"]["version_number"]
        == current["replacement_source"]["version_number"]
    )
    # A normal v2 version rollback must not make business revert use an older
    # release's source version instead of the source captured at activation.
    request(
        env,
        "POST",
        f"v2/workflows/{target_id}/rollback",
        {"target_local_version": 1, "reason": "验证版本回滚后的业务回切"},
    )
    reverted = switch(env, target_id, "revert")
    assert reverted["source"]["is_enabled"] and not reverted["target"]["is_enabled"]
    assert (
        reverted["last_action"]["details"]["migration_source"]
        == current["replacement_source"]
    )


def test_historical_renderer_digest_is_projected_from_frozen_run_without_rewriting_task(
    environment,
):
    env = environment
    migration = preview(env, CHOICES[1])
    index_file(env, "hemp_cotton_records")
    published = publish(
        env, stage(env, migration["candidate"]), migration["replacement_source"]
    )
    switch(env, published["workflow_id"], "activate")
    current_run = run(env, published["workflow_id"])
    drain(env)
    task = waiting_task(env, current_run["id"])
    row = env.db.get(ExecutionHumanTask, task["id"])
    old = deepcopy(task["renderer_contract"])
    old["contract_digest"] = old.pop("node_contract_digest")
    row.renderer_contract = old
    env.db.commit()
    projected = waiting_task(env, current_run["id"])
    assert projected["renderer_contract"] == task["renderer_contract"]
    env.db.refresh(row)
    assert row.renderer_contract == old
    for key, unknown in (("contract_digest", "0" * 64), ("protocol", "unknown.v99")):
        damaged = {**old, key: unknown}
        row.renderer_contract = damaged
        env.db.commit()
        assert waiting_task(env, current_run["id"])["renderer_contract"] == damaged


def test_later_publish_and_rollback_preserve_disabled_or_active_state(environment):
    env = environment
    migration = preview(env, CHOICES[1])
    published = publish(
        env, stage(env, migration["candidate"]), migration["replacement_source"]
    )
    target_id = published["workflow_id"]
    for number in (2, 3, 4):
        if number == 3:
            switch(env, target_id, "activate")
        if number == 4:
            switch(env, target_id, "revert")
        document = deepcopy(migration["candidate"])
        document["release"]["release_version"] = number
        document["integrity"]["digest"] = _release_digest(document)
        later = stage(env, document)
        assert later["migration_source"] == migration["replacement_source"]
        publish(env, later)
        state = request(env, "GET", f"v2/workflows/{target_id}/replacement")
        assert state["target"]["is_enabled"] is (number == 3)
    request(
        env,
        "POST",
        f"v2/workflows/{target_id}/rollback",
        {"target_local_version": 1, "reason": "版本回退不会启用入口"},
    )
    assert not request(env, "GET", f"v2/workflows/{target_id}/replacement")["target"][
        "is_enabled"
    ]


def test_switch_permissions_csrf_conflicts_and_readiness(environment, monkeypatch):
    env = environment
    migration = preview(env, CHOICES[1])
    published = publish(
        env, stage(env, migration["candidate"]), migration["replacement_source"]
    )
    target_id = published["workflow_id"]
    state = request(env, "GET", f"v2/workflows/{target_id}/replacement")
    path = f"/api/execution/v2/workflows/{target_id}/replacement/activate"
    payload = replacement_payload(state)
    assert env.user_client.post(path, json=payload).status_code == 403
    assert (
        env.client.post(
            path, json=payload, headers={"X-CSRF-Token": "wrong"}
        ).status_code
        == 403
    )
    assert (
        env.client.post(
            path, json={**payload, "expected_source_revision": 999}
        ).status_code
        == 409
    )
    monkeypatch.setattr(settings, "EXECUTION_CONTRACT_MODE", "legacy")
    assert (
        env.client.post(path, json=payload).json()["code"]
        == "execution_contract_mode_not_enforced"
    )
    monkeypatch.setattr(settings, "EXECUTION_CONTRACT_MODE", "enforced")
    root = env.roots["hemp_cotton_records"]
    root.is_available = False
    env.db.commit()
    assert env.client.post(path, json=payload).status_code == 409
    env.db.refresh(source_workflow(env, CHOICES[1]))
    assert not source_workflow(env, CHOICES[1]).archived_at
    root.is_available = True
    env.db.commit()
    assert env.client.post(path, json=payload).status_code == 200
    assert (
        env.client.post(path, json=payload).json()["code"]
        == "replacement_revision_changed"
    )


def test_controlled_excel_candidate_stays_internal_and_records_real_publish(
    environment,
):
    env = environment
    migration = preview(env, CANARY)
    staged = stage(env, migration["candidate"])
    internal = publish(env, staged, migration["replacement_source"])
    assert (
        switch(env, internal["workflow_id"], "activate", status=409)["code"]
        == "replacement_internal_acceptance_only"
    )
    # Standalone acceptance copy in a disposable environment; production's
    # registered replacement above remains disabled and can never activate.
    document = deepcopy(migration["candidate"])
    document["release"]["slug"] = "internal-xlsx-acceptance"
    document["integrity"]["digest"] = _release_digest(document)
    published = publish(env, stage(env, document))
    source_path = env.path / "special_wool_records" / "controlled.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A1"] = "before"
    workbook.save(source_path)
    started = run(
        env,
        published["workflow_id"],
        inputs={
            "source": {
                "root_id": "special_wool_records",
                "relative_path": source_path.name,
            },
            "mutation_id": "replacement-canary",
            "writes": [{"sheet": "Sheet1", "cell": "B2", "value": "after"}],
            "target": {
                "root_id": "execution_publish",
                "relative_path": "accepted.xlsx",
            },
        },
    )
    drain(env)
    task = waiting_task(env, started["id"])
    assert task["renderer_contract"]["capability"] == "human.approval"
    assert (
        task["renderer_contract"]["contract_digest"]
        == "c65127408d47900c4b04fb4e7da58b629d6507e6e48d2a86f4949e6fc342564b"
    )
    request(
        env,
        "POST",
        f"v1/human-tasks/{task['id']}/submit",
        {
            "revision": task["revision"],
            "data": {"decision": "approved", "reason": "工作簿重读验收"},
        },
    )
    drain(env)
    assert request(env, "GET", f"v1/runs/{started['id']}")["status"] == "completed"
    with_file = load_workbook(
        env.path / "execution_publish" / "accepted.xlsx", data_only=True
    )
    assert with_file["Sheet1"]["B2"].value == "after"
    with_file.close()
    mutation = env.db.query(ExecutionFileMutation).filter_by(run_id=started["id"]).one()
    assert (
        env.db.query(ExecutionPublishReceipt).filter_by(mutation_id=mutation.id).count()
        == 1
    )
