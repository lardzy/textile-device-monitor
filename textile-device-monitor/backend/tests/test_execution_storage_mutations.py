from __future__ import annotations

import errno
import json
import os
import shutil
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook

from app.execution.mutations import (
    CellWrite,
    FileMutationService,
    InvalidMutationIdError,
    MutationConflictError,
    PublishConflictError,
    UnsupportedWorkbookError,
)
from app.execution.storage import ArtifactRef, FileGateway, StorageRoot


def _services(
    tmp_path: Path,
) -> tuple[FileGateway, FileMutationService, Path, Path, Path]:
    source = tmp_path / "source"
    staging = tmp_path / "staging"
    publish = tmp_path / "publish"
    source.mkdir()
    staging.mkdir()
    publish.mkdir()
    gateway = FileGateway(
        [
            StorageRoot("source", source),
            StorageRoot("staging", staging, writable=True),
            StorageRoot("publish", publish, writable=True, publishable=True),
        ]
    )
    service = FileMutationService(gateway, staging_root_id="staging")
    return gateway, service, source, staging, publish


def _create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "原始记录"
    sheet["A1"] = "原始编号"
    sheet["B2"] = 1
    workbook.save(path)
    workbook.close()


def test_copy_write_verify_publish_is_idempotent_and_preserves_source(
    tmp_path: Path,
):
    gateway, service, source, _, publish = _services(tmp_path)
    source_path = source / "template.xlsx"
    _create_workbook(source_path)
    source_before = gateway.fingerprint(ArtifactRef("source", "template.xlsx"))

    prepared = service.prepare_working_copy(
        "mutation-001",
        ArtifactRef("source", "template.xlsx"),
    )
    prepared_again = service.prepare_working_copy(
        "mutation-001",
        ArtifactRef("source", "template.xlsx"),
    )
    assert prepared_again.reused is True
    assert prepared_again.working_ref == prepared.working_ref

    writes = [
        CellWrite("原始记录", "A1", "26X910095-1"),
        CellWrite("原始记录", "$B$2", 42),
    ]
    written = service.write_xlsx_cells(
        "mutation-001",
        prepared.working_ref,
        writes,
    )
    assert written.verification.verified is True
    written_again = service.write_xlsx_cells(
        "mutation-001",
        prepared.working_ref,
        writes,
    )
    assert written_again.reused is True

    published = service.publish(
        "mutation-001",
        prepared.working_ref,
        ArtifactRef("publish", "results/26X910095-1.xlsx"),
    )
    published_again = service.publish(
        "mutation-001",
        prepared.working_ref,
        ArtifactRef("publish", "results/26X910095-1.xlsx"),
    )
    assert published_again.reused is True
    assert published_again.published_fingerprint == published.published_fingerprint

    assert gateway.fingerprint(ArtifactRef("source", "template.xlsx")) == source_before
    assert (publish / "results" / "26X910095-1.xlsx").is_file()
    workbook = load_workbook(publish / "results" / "26X910095-1.xlsx")
    try:
        assert workbook["原始记录"]["A1"].value == "26X910095-1"
        assert workbook["原始记录"]["B2"].value == 42
    finally:
        workbook.close()


def test_verification_closes_stream_before_atomic_replace(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    _, service, source, _, _ = _services(tmp_path)
    source_path = source / "template.xlsx"
    _create_workbook(source_path)

    verification_streams = []
    original_load_workbook = load_workbook

    def tracked_load_workbook(*args, **kwargs):
        filename = kwargs.get("filename", args[0] if args else None)
        if hasattr(filename, "read"):
            verification_streams.append(filename)
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(
        "app.execution.mutations.load_workbook",
        tracked_load_workbook,
    )

    receipt = service.verify_xlsx_cells(
        ArtifactRef("source", source_path.name),
        [CellWrite("原始记录", "A1", "原始编号")],
    )

    assert receipt.verified is True
    assert verification_streams
    assert all(stream.closed for stream in verification_streams)

    replacement_path = source / "replacement.xlsx"
    shutil.copyfile(source_path, replacement_path)
    os.replace(replacement_path, source_path)
    assert source_path.is_file()
    assert not replacement_path.exists()


def test_same_mutation_id_rejects_different_source_or_writes(tmp_path: Path):
    _, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "one.xlsx")
    _create_workbook(source / "two.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-002",
        ArtifactRef("source", "one.xlsx"),
    )

    with pytest.raises(MutationConflictError):
        service.prepare_working_copy(
            "mutation-002",
            ArtifactRef("source", "two.xlsx"),
        )

    service.write_xlsx_cells(
        "mutation-002",
        prepared.working_ref,
        [CellWrite("原始记录", "A1", "first")],
    )
    with pytest.raises(MutationConflictError):
        service.write_xlsx_cells(
            "mutation-002",
            prepared.working_ref,
            [CellWrite("原始记录", "A1", "second")],
        )


def test_source_change_after_prepare_blocks_write_and_publish(tmp_path: Path):
    _, service, source, _, _ = _services(tmp_path)
    source_path = source / "template.xlsx"
    _create_workbook(source_path)
    prepared = service.prepare_working_copy(
        "mutation-003",
        ArtifactRef("source", "template.xlsx"),
    )
    source_path.write_bytes(source_path.read_bytes() + b"changed")

    with pytest.raises(MutationConflictError, match="source_changed_since_prepare"):
        service.write_xlsx_cells(
            "mutation-003",
            prepared.working_ref,
            [CellWrite("原始记录", "A1", "blocked")],
        )
    with pytest.raises(MutationConflictError, match="source_changed_since_prepare"):
        service.publish(
            "mutation-003",
            prepared.working_ref,
            ArtifactRef("publish", "blocked.xlsx"),
        )


def test_publish_refuses_existing_target_without_overwrite(tmp_path: Path):
    _, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-004",
        ArtifactRef("source", "template.xlsx"),
    )
    existing = publish / "existing.xlsx"
    existing.write_bytes(b"keep-me")

    with pytest.raises(PublishConflictError):
        service.publish(
            "mutation-004",
            prepared.working_ref,
            ArtifactRef("publish", "existing.xlsx"),
        )
    assert existing.read_bytes() == b"keep-me"


def test_publish_target_is_bound_to_mutation_id(tmp_path: Path):
    _, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-005",
        ArtifactRef("source", "template.xlsx"),
    )
    service.publish(
        "mutation-005",
        prepared.working_ref,
        ArtifactRef("publish", "first.xlsx"),
    )

    with pytest.raises(MutationConflictError, match="mutation_publish_target_conflict"):
        service.publish(
            "mutation-005",
            prepared.working_ref,
            ArtifactRef("publish", "second.xlsx"),
        )


def test_non_xlsx_workbook_requires_desktop_adapter(tmp_path: Path):
    _, service, source, _, _ = _services(tmp_path)
    (source / "legacy.xls").write_bytes(b"legacy")
    prepared = service.prepare_working_copy(
        "mutation-006",
        ArtifactRef("source", "legacy.xls"),
    )

    with pytest.raises(UnsupportedWorkbookError):
        service.write_xlsx_cells(
            "mutation-006",
            prepared.working_ref,
            [CellWrite("Sheet1", "A1", "value")],
        )


@pytest.mark.parametrize(
    "mutation_id",
    ["", "../escape", "contains/slash", r"contains\\slash", "x" * 101],
)
def test_invalid_mutation_id_is_rejected(tmp_path: Path, mutation_id: str):
    _, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "template.xlsx")

    with pytest.raises(InvalidMutationIdError):
        service.prepare_working_copy(
            mutation_id,
            ArtifactRef("source", "template.xlsx"),
        )


def test_manual_working_copy_change_is_detected(tmp_path: Path):
    gateway, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-007",
        ArtifactRef("source", "template.xlsx"),
    )
    working_path = gateway.resolve(
        prepared.working_ref,
        for_write=True,
        expected_type="file",
    )
    working_path.write_bytes(working_path.read_bytes() + b"manual-change")

    with pytest.raises(MutationConflictError, match="working_copy_changed_before_write"):
        service.write_xlsx_cells(
            "mutation-007",
            prepared.working_ref,
            [CellWrite("原始记录", "A1", "blocked")],
        )


def test_prepare_reconciles_copy_completed_before_manifest(tmp_path: Path):
    gateway, service, source, _, _ = _services(tmp_path)
    source_path = source / "template.xlsx"
    _create_workbook(source_path)
    source_before = gateway.fingerprint(ArtifactRef("source", "template.xlsx"))

    with patch.object(
        service,
        "_write_manifest",
        side_effect=OSError("simulated-crash-before-manifest"),
    ):
        with pytest.raises(OSError, match="simulated-crash"):
            service.prepare_working_copy(
                "mutation-reconcile-copy",
                ArtifactRef("source", "template.xlsx"),
            )

    service = FileMutationService(gateway, staging_root_id="staging")
    receipt = service.prepare_working_copy(
        "mutation-reconcile-copy",
        ArtifactRef("source", "template.xlsx"),
    )
    assert receipt.reused is True
    assert gateway.fingerprint(receipt.working_ref).sha256 == source_before.sha256
    assert gateway.fingerprint(ArtifactRef("source", "template.xlsx")) == source_before


def test_write_reconciles_replacement_completed_before_final_manifest(
    tmp_path: Path,
):
    gateway, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-reconcile-write",
        ArtifactRef("source", "template.xlsx"),
    )
    writes = [CellWrite("原始记录", "A1", "reconciled")]

    original_write_manifest = service._write_manifest
    call_count = 0

    def fail_after_replacement(mutation_id, payload):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            return original_write_manifest(mutation_id, payload)
        raise OSError("simulated-crash-after-replacement")

    with patch.object(service, "_write_manifest", side_effect=fail_after_replacement):
        with pytest.raises(OSError, match="simulated-crash"):
            service.write_xlsx_cells(
                "mutation-reconcile-write",
                prepared.working_ref,
                writes,
            )

    service = FileMutationService(gateway, staging_root_id="staging")
    receipt = service.write_xlsx_cells(
        "mutation-reconcile-write",
        prepared.working_ref,
        writes,
    )
    assert receipt.reused is True
    assert receipt.verification.verified is True
    manifest_path = gateway.resolve(
        ArtifactRef(
            "staging",
            ".execution-mutations/mutation-reconcile-write/manifest.json",
        ),
        expected_type="file",
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["write"]["status"] == "complete"


def test_write_resumes_when_intent_persisted_before_replacement(tmp_path: Path):
    gateway, service, source, _, _ = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-resume-write",
        ArtifactRef("source", "template.xlsx"),
    )
    working_path = gateway.resolve(
        prepared.working_ref,
        for_write=True,
        expected_type="file",
    )
    writes = [CellWrite("原始记录", "A1", "resumed")]
    original_replace = os.replace

    def fail_workbook_replacement(source_path, target_path):
        if Path(target_path) == working_path:
            raise OSError("simulated-crash-before-replacement")
        return original_replace(source_path, target_path)

    with patch(
        "app.execution.mutations.os.replace",
        side_effect=fail_workbook_replacement,
    ):
        with pytest.raises(OSError, match="simulated-crash"):
            service.write_xlsx_cells(
                "mutation-resume-write",
                prepared.working_ref,
                writes,
            )

    assert service.verify_xlsx_cells(prepared.working_ref, writes).verified is False
    service = FileMutationService(gateway, staging_root_id="staging")
    receipt = service.write_xlsx_cells(
        "mutation-resume-write",
        prepared.working_ref,
        writes,
    )
    assert receipt.verification.verified is True


def test_publish_reconciles_copy_completed_before_final_manifest(tmp_path: Path):
    gateway, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-reconcile-publish",
        ArtifactRef("source", "template.xlsx"),
    )
    published_ref = ArtifactRef("publish", "result.xlsx")

    original_write_manifest = service._write_manifest
    call_count = 0

    def fail_after_publish(mutation_id, payload):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            return original_write_manifest(mutation_id, payload)
        raise OSError("simulated-crash-after-publish")

    with patch.object(service, "_write_manifest", side_effect=fail_after_publish):
        with pytest.raises(OSError, match="simulated-crash"):
            service.publish(
                "mutation-reconcile-publish",
                prepared.working_ref,
                published_ref,
            )

    assert (publish / "result.xlsx").is_file()
    service = FileMutationService(gateway, staging_root_id="staging")
    receipt = service.publish(
        "mutation-reconcile-publish",
        prepared.working_ref,
        published_ref,
    )
    assert receipt.reused is True
    assert receipt.published_fingerprint.sha256 == gateway.fingerprint(
        prepared.working_ref
    ).sha256


def test_publish_resumes_when_intent_persisted_before_copy(tmp_path: Path):
    gateway, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-resume-publish",
        ArtifactRef("source", "template.xlsx"),
    )
    published_ref = ArtifactRef("publish", "resumed.xlsx")

    with patch.object(
        service,
        "_copy_exclusive",
        side_effect=OSError("simulated-crash-before-copy"),
    ):
        with pytest.raises(OSError, match="simulated-crash"):
            service.publish(
                "mutation-resume-publish",
                prepared.working_ref,
                published_ref,
            )

    assert not (publish / "resumed.xlsx").exists()
    service = FileMutationService(gateway, staging_root_id="staging")
    receipt = service.publish(
        "mutation-resume-publish",
        prepared.working_ref,
        published_ref,
    )
    assert receipt.published_fingerprint.sha256
    assert (publish / "resumed.xlsx").is_file()


def test_publish_adopts_legacy_orphan_only_when_hash_matches(tmp_path: Path):
    _, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-legacy-publish",
        ArtifactRef("source", "template.xlsx"),
    )
    working_path = service._gateway.resolve(
        prepared.working_ref,
        expected_type="file",
    )
    target_path = publish / "legacy-result.xlsx"
    shutil.copy2(working_path, target_path)

    receipt = service.publish(
        "mutation-legacy-publish",
        prepared.working_ref,
        ArtifactRef("publish", "legacy-result.xlsx"),
    )
    assert receipt.reused is True


def test_smb_hardlink_fallback_publishes_without_overwrite(tmp_path: Path):
    _, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-smb-fallback",
        ArtifactRef("source", "template.xlsx"),
    )

    with patch(
        "app.execution.mutations.os.link",
        side_effect=OSError(errno.EOPNOTSUPP, "hard links unsupported"),
    ):
        receipt = service.publish(
            "mutation-smb-fallback",
            prepared.working_ref,
            ArtifactRef("publish", "smb-result.xlsx"),
        )

    assert receipt.published_fingerprint.sha256
    assert (publish / "smb-result.xlsx").is_file()


def test_smb_fallback_losing_race_never_overwrites_target(tmp_path: Path):
    _, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-smb-race",
        ArtifactRef("source", "template.xlsx"),
    )
    target_path = publish / "race-result.xlsx"

    def create_rival_then_reject_link(*args, **kwargs):
        target_path.write_bytes(b"rival-content")
        raise OSError(errno.EOPNOTSUPP, "hard links unsupported")

    with patch(
        "app.execution.mutations.os.link",
        side_effect=create_rival_then_reject_link,
    ):
        with pytest.raises(PublishConflictError):
            service.publish(
                "mutation-smb-race",
                prepared.working_ref,
                ArtifactRef("publish", "race-result.xlsx"),
            )

    assert target_path.read_bytes() == b"rival-content"


def test_interrupted_smb_partial_never_exposes_final_and_restart_recovers(
    tmp_path: Path,
):
    gateway, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    prepared = service.prepare_working_copy(
        "mutation-smb-interrupted",
        ArtifactRef("source", "template.xlsx"),
    )
    target_path = publish / "interrupted.xlsx"

    def copy_prefix_then_disconnect(source_handle, target_handle, length):
        target_handle.write(source_handle.read(32))
        target_handle.flush()
        assert not target_path.exists()
        raise OSError("simulated-smb-disconnect")

    with patch(
        "app.execution.mutations.shutil.copyfileobj",
        side_effect=copy_prefix_then_disconnect,
    ):
        with pytest.raises(OSError, match="simulated-smb-disconnect"):
            service.publish(
                "mutation-smb-interrupted",
                prepared.working_ref,
                ArtifactRef("publish", "interrupted.xlsx"),
            )

    assert not target_path.exists()
    assert len(list(publish.glob(".execution-*.claim"))) == 1
    partials = list(publish.glob(".execution-*.partial"))
    assert len(partials) == 1
    assert partials[0].stat().st_size == 32

    service = FileMutationService(gateway, staging_root_id="staging")
    with patch(
        "app.execution.mutations.os.link",
        side_effect=OSError(errno.EOPNOTSUPP, "hard links unsupported"),
    ):
        receipt = service.publish(
            "mutation-smb-interrupted",
            prepared.working_ref,
            ArtifactRef("publish", "interrupted.xlsx"),
        )

    assert receipt.published_fingerprint.sha256
    assert target_path.is_file()
    assert not list(publish.glob(".execution-*.claim"))
    assert not list(publish.glob(".execution-*.partial"))


def test_target_claim_prevents_other_mutation_from_taking_partial(
    tmp_path: Path,
):
    gateway, service, source, _, publish = _services(tmp_path)
    _create_workbook(source / "template.xlsx")
    first = service.prepare_working_copy(
        "mutation-claim-owner",
        ArtifactRef("source", "template.xlsx"),
    )

    def stop_mid_copy(source_handle, target_handle, length):
        target_handle.write(source_handle.read(16))
        raise OSError("owner-disconnected")

    with patch(
        "app.execution.mutations.shutil.copyfileobj",
        side_effect=stop_mid_copy,
    ):
        with pytest.raises(OSError, match="owner-disconnected"):
            service.publish(
                "mutation-claim-owner",
                first.working_ref,
                ArtifactRef("publish", "claimed.xlsx"),
            )

    second = service.prepare_working_copy(
        "mutation-claim-rival",
        ArtifactRef("source", "template.xlsx"),
    )
    with pytest.raises(
        MutationConflictError,
        match="target_claimed_by_other_mutation",
    ):
        service.publish(
            "mutation-claim-rival",
            second.working_ref,
            ArtifactRef("publish", "claimed.xlsx"),
        )

    assert not (publish / "claimed.xlsx").exists()
    assert len(list(publish.glob(".execution-*.partial"))) == 1
    service = FileMutationService(gateway, staging_root_id="staging")
    service.publish(
        "mutation-claim-owner",
        first.working_ref,
        ArtifactRef("publish", "claimed.xlsx"),
    )
    assert (publish / "claimed.xlsx").is_file()
