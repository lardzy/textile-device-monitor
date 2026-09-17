"""P3 service/API/native workflow equivalence using the same XLS/XLSX corpus."""

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import pytest
import xlwt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from PIL import Image

from app.config import settings
from app.execution.models import (
    ExecutionArtifact, ExecutionFileIndexEntry, ExecutionHumanTask,
    ExecutionNodeRun, ExecutionRun, ExecutionStorageRoot,
)
from app.execution.regenerated_fiber_results import read_regenerated_fiber_result
from tests.test_execution_workflow_replacement import (
    environment, drain, preview, publish, request, source_workflow, switch, waiting_task,
)

NUMBER = "260144785"


@pytest.fixture
def records_env(environment, monkeypatch):
    env = environment
    monkeypatch.setattr(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_local_write")
    path = env.path / "regenerated_fiber_records"
    path.mkdir()
    root = ExecutionStorageRoot(root_id="regenerated_fiber_records", name="再生纤", local_path=str(path),
        access_mode="read", category_key="regenerated_fiber", is_active=True, is_available=True, scan_generation=1)
    env.db.add(root)
    env.db.commit()
    env.roots[root.root_id] = root
    return env


def workbook_file(env, method, suffix=".xlsx", index=1, *, corrupt=False):
    root = env.roots["regenerated_fiber_records"]
    path = Path(root.local_path) / f"{NUMBER}-{index}{suffix}"
    sheet_name = "截面统计报告1" if method == "area" else "根数法报告1"
    offset = 2 if method == "area" else 0
    if suffix == ".xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet["B14"] = "已保存"
        sheet["I8"] = "张检验员"
        sheet.cell(24 + offset, 2, "正面")
        sheet.cell(25 + offset, 2, "棉")
        sheet.cell(26 + offset, 2, 69.9526).number_format = "0.0"
        sheet.cell(25 + offset, 3, "粘纤")
        sheet.cell(26 + offset, 3, 30.0474).number_format = "0.0"
        sheet.cell(27 + offset, 2, "同批样本")
        image = env.path / f"image-{index}.png"
        Image.new("RGB", (12, 8), color="red").save(image)
        sheet.add_image(WorkbookImage(str(image)), "L24")
        book.save(path)
        book.close()
    else:
        book = xlwt.Workbook()
        sheet = book.add_sheet(sheet_name)
        sheet.write(13, 1, "已保存")
        sheet.write(7, 8, "张检验员")
        sheet.write(23 + offset, 1, "正面")
        sheet.write(24 + offset, 1, "棉")
        sheet.write(25 + offset, 1, 69.9526, xlwt.easyxf(num_format_str="0.0"))
        sheet.write(24 + offset, 2, "粘纤")
        sheet.write(25 + offset, 2, 30.0474, xlwt.easyxf(num_format_str="0.0"))
        sheet.write(26 + offset, 1, "同批样本")
        image = env.path / f"image-{index}.bmp"
        Image.new("RGB", (12, 8), color="red").save(image)
        sheet.insert_bitmap(str(image), 23, 11)
        book.save(str(path))
    if corrupt:
        path.write_bytes(b"broken Excel")
    stat = path.stat()
    entry = ExecutionFileIndexEntry(storage_root_id=root.id, relative_path=path.name, filename=path.name,
        extension=suffix, file_kind="workbook", inspection_number=NUMBER, group_key=NUMBER,
        size_bytes=stat.st_size, modified_at=datetime.fromtimestamp(stat.st_mtime, timezone.utc),
        fingerprint=f"{stat.st_size}:{stat.st_mtime_ns}", scan_generation=1)
    env.db.add(entry)
    env.db.commit()
    return entry


def stage_candidate(env, migration):
    document = migration["candidate"]
    report = request(env, "POST", "v2/workflow-releases/preflight", {"document": document})
    assert report["content_valid"], report
    release = request(env, "POST", "v2/workflow-releases/apply", {"preflight_token": report["preflight_token"]}, status=201)
    request(env, "PUT", f"v2/workflow-releases/{release['id']}/deployment-binding", {
        "environment": "default", "expected_revision": 0,
        "bindings": migration["binding_suggestions"],
    })
    return release


def create(env, workflow_id):
    return request(env, "POST", "v1/runs", {"workflow_id": workflow_id, "inspection_number": NUMBER,
        "input_data": {}, "global_data": {}, "idempotency_key": str(uuid4())}, status=201)["run"]


def node_output(env, run_id, node_id):
    env.db.expire_all()
    node = env.db.query(ExecutionNodeRun).filter_by(run_id=run_id, node_id=node_id).one()
    assert node.status == "succeeded", (node.error_code, node.error_message)
    return deepcopy(node.output_data)


def comparable(env, output):
    value = deepcopy(output)
    value.pop("items", None)
    for file in value["files"]:
        for image in (file.get("result") or {}).get("images") or []:
            if image.get("artifact_id"):
                artifact = env.db.get(ExecutionArtifact, image.pop("artifact_id"))
                image.pop("preview_url", None)
                image["sha256"] = artifact.content_sha256
                image["size_bytes"] = artifact.size_bytes
            image.pop("index", None)
    return value


@pytest.mark.parametrize("method", ["area", "count"])
@pytest.mark.parametrize("suffix", [".xls", ".xlsx"])
@pytest.mark.parametrize("count", [1, 2])
def test_same_corpus_api_v1_v2_and_replacement(records_env, method, suffix, count):
    env = records_env
    entries = [workbook_file(env, method, suffix, i + 1) for i in range(count)]
    slug = f"regenerated-fiber-{method}-method"
    source = source_workflow(env, slug)
    original_definition = deepcopy(source.draft_definition)
    original_version = source.published_version_number
    api_query = request(env, "POST", "v1/regenerated-fiber/query", {"method": method, "inspection_number": NUMBER})
    assert {file['id'] for file in api_query['candidates']} == {entry.id for entry in entries}
    api_read = request(env, "POST", "v1/regenerated-fiber/read-results", {"method": method, "files": api_query['candidates']})
    assert env.db.query(ExecutionRun).count() == 0
    assert env.db.query(ExecutionArtifact).count() == 0
    old_run = create(env, source.id)
    drain(env)
    old_query = node_output(env, old_run['id'], 'query')
    old_read = node_output(env, old_run['id'], 'read-results')
    assert api_query == old_query
    assert comparable(env, api_read) == comparable(env, old_read)

    migration = preview(env, slug, profile="native_p3")
    assert migration == preview(env, slug, profile="native_p3")
    assert migration['content_valid'], migration.get('issues')
    assert migration['migration_status'] == 'p3_complete', migration
    assert not migration['blockers']
    assert not migration['publish_ready']
    env.db.refresh(source)
    assert source.draft_definition == original_definition and source.published_version_number == original_version
    release = publish(env, stage_candidate(env, migration), migration['replacement_source'])
    switch(env, release['workflow_id'], 'activate')
    current = create(env, release['workflow_id'])
    drain(env)
    new_read = node_output(env, current['id'], 'read-results')
    assert comparable(env, new_read) == comparable(env, old_read)
    if count == 1:
        assert env.db.query(ExecutionHumanTask).filter_by(run_id=current['id']).count() == 0
    else:
        task = waiting_task(env, current['id'])
        items = request(env, 'GET', f"v1/human-tasks/{task['id']}")['node_run']['input_data']['items']
        assert items[0]['metadata']['presentation'] == 'result_file'
        assert items[0]['metadata']['result']['parts'][0]['components'][0]['content'] == 70.0
        request(env, 'POST', f"v1/human-tasks/{task['id']}/submit", {'revision': task['revision'],
            'data': {'selected_ids': [item['id'] for item in items], 'primary_id': items[0]['id']}})
        drain(env)
    detail = request(env, 'GET', f"v1/runs/{current['id']}")
    assert detail['status'] == 'completed', detail
    assert len(detail['output_data']['selected_files']) == count
    old_task = waiting_task(env, old_run['id'])
    request(env, 'POST', f"v1/human-tasks/{old_task['id']}/submit", {'revision': old_task['revision'],
        'data': {'selected_files': [api_query['candidates'][0]['id']], 'primary_file_id': api_query['candidates'][0]['id']}})
    drain(env)
    assert request(env, 'GET', f"v1/runs/{old_run['id']}")['status'] == 'completed'
    recommendations = request(env, 'GET', f'v1/catalog/recommendations?inspection_number={NUMBER}')
    assert any(item['workflow_id'] == release['workflow_id'] for item in recommendations['items'])
    switch(env, release['workflow_id'], 'revert')
    assert request(env, 'GET', f"v1/runs/{current['id']}")['status'] == 'completed'


def test_reader_reuses_parse_and_preserves_partial_failures(records_env):
    env = records_env
    workbook_file(env, 'count')
    result = request(env, 'POST', 'v1/regenerated-fiber/query', {'method': 'count', 'inspection_number': NUMBER})
    good = result['candidates'][0]
    stale = {**good, 'fingerprint': '0:0'}
    with patch('app.execution.regenerated_fiber_results.read_regenerated_fiber_result', wraps=read_regenerated_fiber_result) as parse:
        response = request(env, 'POST', 'v1/regenerated-fiber/read-results', {'method': 'count', 'files': [good, good, stale]})
    assert parse.call_count == 1
    assert response['success_count'] == 2 and response['failed_count'] == 1
    assert response['files'][-1]['error']['code'] == 'result_file_stale'
    bad = request(env, 'POST', 'v1/regenerated-fiber/read-results', {'method': 'count', 'files': [stale]}, status=422)
    assert bad['code'] == 'result_workbooks_unreadable'
    request(env, 'POST', 'v1/regenerated-fiber/query', {'method': 'arbitrary', 'inspection_number': NUMBER}, status=422)


def test_migration_reports_parked_reader_and_supports_renamed_nodes(records_env):
    env = records_env
    source = source_workflow(env, 'regenerated-fiber-area-method')
    original = deepcopy(source.draft_definition)
    renamed = deepcopy(original)
    replacements = {node['id']: f"custom-{node['id']}" for node in renamed['nodes']}
    for node in renamed['nodes']:
        node['id'] = replacements[node['id']]
        for key, value in node.get('input_mapping', {}).items():
            if isinstance(value, str):
                for old, new in replacements.items():
                    value = value.replace(f'$.nodes.{old}.', f'$.nodes.{new}.')
                node['input_mapping'][key] = value
    for edge in renamed['edges']:
        edge['source'] = replacements[edge['source']]
        edge['target'] = replacements[edge['target']]
    source.draft_definition = renamed
    env.db.commit()
    migrated = preview(env, source.slug, source='draft', profile='native_p3')
    assert migrated['migration_status'] == 'p3_complete', migrated
    parked = deepcopy(original)
    next(node for node in parked['nodes'] if node['id'] == 'read-results')['disabled'] = True
    source.draft_definition = parked
    env.db.commit()
    blocked = preview(env, source.slug, source='draft', profile='native_p3')
    assert blocked['blockers'][0]['code'] == 'migration_required_node_inactive'
    assert blocked['migration_status'] != 'p3_complete'
