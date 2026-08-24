from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import settings
from app.database import Base
from app.execution.catalog import (
    bind_user_role,
    ensure_default_catalog,
    ensure_default_rbac,
)
from app.execution.engine import (
    claim_next_node,
    complete_node,
    create_run,
    execute_claimed_node,
    submit_human_task,
)
from app.execution.errors import ExecutionApiError
from app.execution.models import (
    ExecutionEdgeRun,
    ExecutionFileMutation,
    ExecutionFileIndexEntry,
    ExecutionHumanApprovalReceipt,
    ExecutionHumanTask,
    ExecutionNodeAttempt,
    ExecutionNodeRun,
    ExecutionRun,
    ExecutionStorageRoot,
    ExecutionUser,
    ExecutionWorkflow,
)
from app.execution.project_rules import ensure_default_project_rules
from app.execution.release_v2 import (
    apply_release,
    preflight_release,
    publish_release,
    put_deployment_binding,
    rollout_profile_blockers,
)
from app.execution.security import hash_password
from app.execution.v2.examples import (
    build_controlled_xlsx_write_canary_release,
    build_native_human_file_selection_smoke_release,
)
from app.execution.v2.native_handlers import (
    NodeExecutionResult,
    _batch_place,
    _data_assign,
    _flow_branch,
    _flow_join,
)
from app.execution.v2.registry import (
    reset_installed_registry_cache,
    worker_capability_document,
)
from app.execution.worker import ExecutionWorker
from app.execution.worker_state import record_worker_heartbeat


@contextmanager
def _environment():
    with TemporaryDirectory(prefix="execution-p2-") as directory:
        root = Path(directory)
        source_path = root / "source"
        staging_path = root / "staging"
        publish_path = root / "publish"
        source_path.mkdir()
        staging_path.mkdir()
        publish_path.mkdir()
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        Session = sessionmaker(bind=engine, autoflush=False)
        db = Session()
        ensure_default_rbac(db)
        ensure_default_catalog(db)
        ensure_default_project_rules(db)
        admin = ExecutionUser(
            username="p2-primitives-admin",
            display_name="P2 primitives admin",
            password_hash=hash_password("test-password"),
            role="admin",
            is_active=True,
        )
        roots = {
            "source": ExecutionStorageRoot(
                root_id="p2_source",
                name="P2 source",
                local_path=str(source_path),
                access_mode="read",
                category_key="other",
                is_active=True,
                is_available=True,
                scan_generation=1,
            ),
            "staging": ExecutionStorageRoot(
                root_id="execution_staging",
                name="P2 staging",
                local_path=str(staging_path),
                access_mode="write",
                category_key="other",
                is_active=True,
                is_available=True,
                scan_generation=1,
            ),
            "publish": ExecutionStorageRoot(
                root_id="execution_publish",
                name="P2 publish",
                local_path=str(publish_path),
                access_mode="publish",
                category_key="other",
                is_active=True,
                is_available=True,
                scan_generation=1,
            ),
        }
        db.add(admin)
        db.add_all(roots.values())
        db.flush()
        bind_user_role(db, admin, "admin", created_by_id=admin.id)
        db.commit()
        try:
            yield db, admin, roots, root
        finally:
            db.close()
            Base.metadata.drop_all(engine)
            engine.dispose()


def _publish_document(db, admin, document, roots, *, profile):
    with (
        patch.object(settings, "EXECUTION_CONTRACT_MODE", "enforced"),
        patch.object(settings, "EXECUTION_V2_ROLLOUT_PROFILE", profile),
    ):
        content = preflight_release(
            db,
            document=document,
            actor=admin,
            scope="content",
        )
        assert content["content_valid"], content["issues"]
        release = apply_release(
            db,
            document=None,
            preflight_token=content["preflight_token"],
            actor=admin,
        )
        bindings = {
            "root_slots": {
                slot_id: {
                    "root_id": root.root_id,
                    "revision": root.binding_revision,
                }
                for slot_id, root in roots.items()
            },
            "credential_slots": {},
            "role_slots": {},
            "rule_slots": {},
        }
        binding = put_deployment_binding(
            db,
            release_id=release.id,
            environment=settings.EXECUTION_ENVIRONMENT_ID,
            expected_revision=0,
            bindings=bindings,
            actor=admin,
        )
        publish_preflight = preflight_release(
            db,
            document=release.portable_document,
            actor=admin,
            release=release,
            scope="publish",
        )
        assert publish_preflight["publish_ready"], publish_preflight["issues"]
        release, version, _receipt = publish_release(
            db,
            release_id=release.id,
            preflight_token=publish_preflight["preflight_token"],
            actor=admin,
            reason="P2 primitive test",
        )
        db.commit()
    return db.get(ExecutionWorkflow, version.workflow_id)


def _start_worker(db, worker_id):
    ExecutionWorker(worker_id=worker_id)
    reset_installed_registry_cache()
    capabilities = worker_capability_document()
    record_worker_heartbeat(
        db,
        worker_id=worker_id,
        capability_document=capabilities,
    )
    db.commit()


def _execute_next(db, worker_id):
    claimed = claim_next_node(db, worker_id=worker_id)
    if claimed is None:
        states = [
            {
                "node_id": row.node_id,
                "status": row.status,
                "error_code": row.error_code,
                "error_message": row.error_message,
                "execution_binding_digest": row.execution_binding_digest,
            }
            for row in db.query(ExecutionNodeRun)
            .order_by(ExecutionNodeRun.created_at, ExecutionNodeRun.node_id)
            .all()
        ]
        raise AssertionError(states)
    node_id = claimed.node_id
    claimed_id = claimed.id
    lease_token = claimed.lease_token
    db.commit()
    execute_claimed_node(db, claimed_id, lease_token)
    db.commit()
    return node_id


def _prepare_canary_for_publish(
    db,
    admin,
    roots,
    root_path,
    *,
    suffix,
):
    workbook_path = root_path / "source" / f"controlled-{suffix}.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A1"] = "before"
    workbook.save(workbook_path)
    workflow = _publish_document(
        db,
        admin,
        build_controlled_xlsx_write_canary_release(),
        {
            "special_wool_records": roots["source"],
            "execution_staging": roots["staging"],
            "execution_publish": roots["publish"],
        },
        profile="p2_publish",
    )
    worker_id = f"p2-write-{suffix}"
    _start_worker(db, worker_id)
    mutation_id = f"p2-xlsx-{suffix}"
    target_relative_path = f"P2-XLSX-SECURITY/{suffix}.xlsx"
    with patch.object(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_publish"):
        run, existed = create_run(
            db,
            workflow=workflow,
            actor=admin,
            inspection_number=f"P2-XLSX-{suffix}",
            input_data={
                "inspection_number": f"P2-XLSX-{suffix}",
                "source": {
                    "root_id": roots["source"].root_id,
                    "relative_path": workbook_path.name,
                },
                "mutation_id": mutation_id,
                "writes": [{"sheet": "Sheet1", "cell": "B2", "value": "after"}],
                "target": {
                    "root_id": roots["publish"].root_id,
                    "relative_path": target_relative_path,
                },
            },
            global_data={},
            idempotency_key=mutation_id,
        )
    assert existed is False
    db.commit()
    assert [_execute_next(db, worker_id) for _ in range(7)] == [
        "start",
        "classify",
        "extract",
        "copy",
        "write",
        "verify",
        "confirm",
    ]
    task = (
        db.query(ExecutionHumanTask)
        .filter(ExecutionHumanTask.run_id == run.id)
        .one()
    )
    submit_human_task(
        db,
        task_id=task.id,
        expected_revision=task.revision,
        data={"decision": "approved", "reason": "security boundary"},
        actor=admin,
    )
    db.commit()
    receipt = (
        db.query(ExecutionHumanApprovalReceipt)
        .filter(ExecutionHumanApprovalReceipt.run_id == run.id)
        .one()
    )
    return (
        run,
        task,
        receipt,
        worker_id,
        root_path / "publish" / target_relative_path,
    )


def test_native_human_file_selection_release_runs_with_stable_ids():
    with _environment() as (db, admin, roots, root_path):
        document = build_native_human_file_selection_smoke_release()
        workflow = _publish_document(
            db,
            admin,
            document,
            {"source": roots["source"]},
            profile="p2_human",
        )
        source_file = root_path / "source" / "P2-001.xlsx"
        source_file.write_bytes(b"indexed fixture")
        stat = source_file.stat()
        entry = ExecutionFileIndexEntry(
            storage_root_id=roots["source"].id,
            relative_path=source_file.name,
            filename=source_file.name,
            extension=".xlsx",
            file_kind="workbook",
            inspection_number="P2-001",
            group_key="P2-001",
            size_bytes=stat.st_size,
            modified_at=datetime.now(timezone.utc),
            fingerprint=f"{stat.st_size}:{stat.st_mtime_ns}",
            scan_generation=roots["source"].scan_generation,
        )
        db.add(entry)
        db.commit()
        worker_id = "p2-human-worker"
        _start_worker(db, worker_id)
        with (
            patch.object(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p1_readonly"),
            pytest.raises(ExecutionApiError) as blocked,
        ):
            create_run(
                db,
                workflow=workflow,
                actor=admin,
                inspection_number="P2-001",
                input_data={"inspection_number": "P2-001"},
                global_data={},
                idempotency_key="p2-native-human-blocked",
            )
        assert blocked.value.code == "rollout_profile_blocked"
        db.rollback()
        with patch.object(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_human"):
            run, existed = create_run(
                db,
                workflow=workflow,
                actor=admin,
                inspection_number="P2-001",
                input_data={"inspection_number": "P2-001"},
                global_data={},
                idempotency_key="p2-native-human-smoke",
            )
        assert existed is False
        db.commit()
        assert [_execute_next(db, worker_id) for _ in range(3)] == [
            "start",
            "query",
            "select",
        ]
        task = (
            db.query(ExecutionHumanTask)
            .filter(ExecutionHumanTask.run_id == run.id)
            .one()
        )
        assert task.status == "open"
        assert task.renderer_contract["capability"] == "human.select"
        submit_human_task(
            db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"selected_ids": [entry.id], "primary_id": entry.id},
            actor=admin,
        )
        db.commit()
        assert [_execute_next(db, worker_id) for _ in range(2)] == [
            "aggregate",
            "end",
        ]
        db.refresh(run)
        assert run.status == "completed"
        assert run.output_data["count"] == 1
        assert run.output_data["selected_items"][0]["id"] == entry.id


def test_controlled_xlsx_canary_records_approval_and_isolated_publish():
    with _environment() as (db, admin, roots, root_path):
        workbook_path = root_path / "source" / "controlled.xlsx"
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.active["A1"] = "before"
        workbook.save(workbook_path)
        document = build_controlled_xlsx_write_canary_release()
        workflow = _publish_document(
            db,
            admin,
            document,
            {
                "special_wool_records": roots["source"],
                "execution_staging": roots["staging"],
                "execution_publish": roots["publish"],
            },
            profile="p2_publish",
        )
        worker_id = "p2-write-worker"
        _start_worker(db, worker_id)
        with (
            patch.object(
                settings,
                "EXECUTION_V2_ROLLOUT_PROFILE",
                "p2_local_write",
            ),
            pytest.raises(ExecutionApiError) as blocked,
        ):
            create_run(
                db,
                workflow=workflow,
                actor=admin,
                inspection_number="P2-XLSX-BLOCKED",
                input_data={
                    "inspection_number": "P2-XLSX-BLOCKED",
                    "source": {
                        "root_id": roots["source"].root_id,
                        "relative_path": workbook_path.name,
                    },
                    "mutation_id": "p2-xlsx-canary-blocked",
                    "writes": [
                        {"sheet": "Sheet1", "cell": "B2", "value": "after"}
                    ],
                    "target": {
                        "root_id": roots["publish"].root_id,
                        "relative_path": "blocked/published.xlsx",
                    },
                },
                global_data={},
                idempotency_key="p2-xlsx-canary-blocked",
            )
        assert blocked.value.code == "rollout_profile_blocked"
        db.rollback()
        with patch.object(settings, "EXECUTION_V2_ROLLOUT_PROFILE", "p2_publish"):
            run, existed = create_run(
                db,
                workflow=workflow,
                actor=admin,
                inspection_number="P2-XLSX-001",
                input_data={
                    "inspection_number": "P2-XLSX-001",
                    "source": {
                        "root_id": roots["source"].root_id,
                        "relative_path": workbook_path.name,
                    },
                    "mutation_id": "p2-xlsx-canary-001",
                    "writes": [
                        {"sheet": "Sheet1", "cell": "B2", "value": "after"}
                    ],
                    "target": {
                        "root_id": roots["publish"].root_id,
                        "relative_path": "P2-XLSX-001/published.xlsx",
                    },
                },
                global_data={},
                idempotency_key="p2-xlsx-canary-001",
            )
        assert existed is False
        db.commit()
        assert [_execute_next(db, worker_id) for _ in range(7)] == [
            "start",
            "classify",
            "extract",
            "copy",
            "write",
            "verify",
            "confirm",
        ]
        task = (
            db.query(ExecutionHumanTask)
            .filter(ExecutionHumanTask.run_id == run.id)
            .one()
        )
        submit_human_task(
            db,
            task_id=task.id,
            expected_revision=task.revision,
            data={"decision": "approved", "reason": "isolated canary"},
            actor=admin,
        )
        db.commit()
        receipt = (
            db.query(ExecutionHumanApprovalReceipt)
            .filter(ExecutionHumanApprovalReceipt.run_id == run.id)
            .one()
        )
        assert receipt.decision == "approved"
        assert receipt.subject_digest == task.result_data[
            "approval_receipt"
        ]["subject_digest"]
        assert [_execute_next(db, worker_id) for _ in range(2)] == [
            "publish",
            "end",
        ]
        db.refresh(run)
        assert run.status == "completed"
        published_path = (
            root_path / "publish" / "P2-XLSX-001" / "published.xlsx"
        )
        assert published_path.is_file()
        published_workbook = __import__("openpyxl").load_workbook(
            published_path,
            data_only=True,
        )
        try:
            assert published_workbook["Sheet1"]["B2"].value == "after"
        finally:
            published_workbook.close()


@pytest.mark.parametrize(
    ("tamper", "expected_code"),
    [
        ("receipt_id", "approval_receipt_invalid"),
        ("subject", "approval_subject_mismatch"),
    ],
)
def test_controlled_publish_rejects_forged_receipt_or_subject_drift(
    tamper,
    expected_code,
):
    with _environment() as (db, admin, roots, root_path):
        run, _task, _receipt, worker_id, target_path = _prepare_canary_for_publish(
            db,
            admin,
            roots,
            root_path,
            suffix=tamper,
        )
        if tamper == "receipt_id":
            approval_node = (
                db.query(ExecutionNodeRun)
                .filter(
                    ExecutionNodeRun.run_id == run.id,
                    ExecutionNodeRun.node_id == "confirm",
                )
                .one()
            )
            approval_node.output_data = {
                **approval_node.output_data,
                "approval_receipt": {
                    **approval_node.output_data["approval_receipt"],
                    "id": "forged-receipt-id",
                },
            }
        else:
            mutation = (
                db.query(ExecutionFileMutation)
                .filter(ExecutionFileMutation.run_id == run.id)
                .one()
            )
            mutation.verification_result = {
                **mutation.verification_result,
                "subject_digest": "0" * 64,
            }
        db.commit()

        assert _execute_next(db, worker_id) == "publish"
        publish_node = (
            db.query(ExecutionNodeRun)
            .filter(
                ExecutionNodeRun.run_id == run.id,
                ExecutionNodeRun.node_id == "publish",
            )
            .one()
        )
        assert publish_node.status == "failed"
        assert publish_node.error_code == expected_code
        assert not target_path.exists()


def test_join_modes_and_atomic_data_assign_conflict():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False)
    db = Session()
    try:
        run = ExecutionRun(
            id="run-join",
            workflow_id="workflow-join",
            created_by_id="user-join",
            idempotency_key="run-join",
            inspection_number="JOIN",
            mode="test",
            status="running",
            definition_snapshot={"schema_version": "1.0", "nodes": [], "edges": []},
            definition_checksum="1" * 64,
            capabilities_snapshot={},
            contract_checksum="2" * 64,
            input_data={},
            global_data={"existing": 1},
            output_data={},
        )
        first = ExecutionNodeRun(
            id="source-a",
            run_id=run.id,
            node_id="a",
            node_type="data.aggregate",
            node_type_version=1,
            node_name="a",
            status="succeeded",
            output_data={"value": "a"},
        )
        second = ExecutionNodeRun(
            id="source-b",
            run_id=run.id,
            node_id="b",
            node_type="data.aggregate",
            node_type_version=1,
            node_name="b",
            status="succeeded",
            output_data={"value": "b"},
        )
        join = ExecutionNodeRun(
            id="join",
            run_id=run.id,
            node_id="join",
            node_type="flow.join",
            node_type_version=1,
            node_name="join",
            status="running",
        )
        branch = ExecutionNodeRun(
            id="branch",
            run_id=run.id,
            node_id="branch",
            node_type="flow.branch",
            node_type_version=1,
            node_name="branch",
            status="running",
        )
        assign = ExecutionNodeRun(
            id="assign",
            run_id=run.id,
            node_id="assign",
            node_type="data.assign",
            node_type_version=1,
            node_name="assign",
            status="running",
            lease_token="assign-lease",
        )
        db.add_all([run, first, second, join, branch, assign])
        db.add_all(
            [
                ExecutionEdgeRun(
                    run_id=run.id,
                    edge_id="edge-b",
                    source_node_id="b",
                    target_node_id="join",
                    status="selected",
                    resolved_at=datetime(2026, 8, 21, 0, 0, 2),
                ),
                ExecutionEdgeRun(
                    run_id=run.id,
                    edge_id="edge-a",
                    source_node_id="a",
                    target_node_id="join",
                    status="selected",
                    resolved_at=datetime(2026, 8, 21, 0, 0, 1),
                ),
                ExecutionEdgeRun(
                    run_id=run.id,
                    edge_id="branch-a",
                    source_node_id="branch",
                    target_node_id="target-a",
                    status="pending",
                    condition={
                        "path": "$.globals.existing",
                        "operator": "eq",
                        "value": 1,
                    },
                ),
                ExecutionEdgeRun(
                    run_id=run.id,
                    edge_id="branch-b",
                    source_node_id="branch",
                    target_node_id="target-b",
                    status="pending",
                    condition={
                        "path": "$.globals.existing",
                        "operator": "gte",
                        "value": 1,
                    },
                ),
                ExecutionEdgeRun(
                    run_id=run.id,
                    edge_id="branch-default",
                    source_node_id="branch",
                    target_node_id="target-default",
                    status="pending",
                    condition="default",
                ),
                ExecutionNodeAttempt(
                    node_run_id=assign.id,
                    attempt_number=1,
                    worker_id="worker",
                    lease_token="assign-lease",
                    status="running",
                ),
            ]
        )
        db.commit()
        context = SimpleNamespace(
            db=db,
            run=run,
            node_run=join,
            node={"config": {"mode": "first_selected"}},
            input_data={},
        )
        assert _flow_join(context) == {"items": [{"value": "a"}]}
        context.node["config"]["mode"] = "collect_selected"
        assert _flow_join(context) == {
            "items": [{"value": "a"}, {"value": "b"}]
        }
        context.node["config"]["mode"] = "all_selected"
        context.input_data = {"items": [{"mapped": "a"}, {"mapped": "b"}]}
        assert _flow_join(context) == {
            "items": [{"mapped": "a"}, {"mapped": "b"}]
        }
        branch_result = _flow_branch(
            SimpleNamespace(
                db=db,
                run=run,
                node_run=branch,
                input_data={"values": {"existing": 1}},
            )
        )
        assert branch_result.selected_edge_ids == ("branch-a", "branch-b")
        for edge in db.query(ExecutionEdgeRun).filter(
            ExecutionEdgeRun.source_node_id == "branch"
        ):
            if edge.edge_id != "branch-default":
                edge.condition = {
                    "path": "$.globals.existing",
                    "operator": "eq",
                    "value": 99,
                }
        db.flush()
        default_result = _flow_branch(
            SimpleNamespace(
                db=db,
                run=run,
                node_run=branch,
                input_data={"values": {"existing": 1}},
            )
        )
        assert default_result.selected_edge_ids == ("branch-default",)
        assigned = _data_assign(
            SimpleNamespace(
                # Model the stale executor view from a concurrent attempt;
                # complete_node must repeat the check after locking the Run.
                run=SimpleNamespace(global_data={}),
                node={
                    "config": {
                        "values": {"existing": 2},
                        "conflict": "error",
                    }
                },
                input_data={},
            )
        )
        try:
            complete_node(
                db,
                node_run_id=assign.id,
                lease_token="assign-lease",
                output_data=assigned.output,
                globals_patch=assigned.globals_patch,
                globals_conflict=assigned.globals_conflict,
            )
        except ExecutionApiError as exc:
            assert exc.code == "data_assign_conflict"
        else:
            raise AssertionError("data.assign conflict must fail under the run lock")
        db.refresh(assign)
        assert assign.status == "running"
        assert run.global_data == {"existing": 1}
    finally:
        db.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_rollout_profiles_progressively_allow_only_native_p2_capabilities():
    instances = [
        {
            "node_id": "human",
            "execution_kind": "human",
            "side_effect_class": "none",
            "source": "resource",
        },
        {
            "node_id": "local-write",
            "execution_kind": "automatic",
            "side_effect_class": "reversible_local_write",
            "source": "resource",
        },
        {
            "node_id": "publish",
            "execution_kind": "automatic",
            "side_effect_class": "durable_write",
            "source": "resource",
        },
        {
            "node_id": "compat-write",
            "execution_kind": "automatic",
            "side_effect_class": "reversible_local_write",
            "source": "v1_registry_adapter",
        },
        {
            "node_id": "external",
            "execution_kind": "external_side_effect",
            "side_effect_class": "external_write",
            "source": "resource",
        },
    ]

    expected = {
        "p1_readonly": {
            "native_human_profile_required",
            "local_write_profile_required",
            "publish_profile_required",
            "compatibility_write_blocked_p2",
            "external_write_blocked_until_p4",
        },
        "p2_human": {
            "local_write_profile_required",
            "publish_profile_required",
            "compatibility_write_blocked_p2",
            "external_write_blocked_until_p4",
        },
        "p2_local_write": {
            "publish_profile_required",
            "compatibility_write_blocked_p2",
            "external_write_blocked_until_p4",
        },
        "p2_publish": {
            "compatibility_write_blocked_p2",
            "external_write_blocked_until_p4",
        },
    }
    for profile, expected_codes in expected.items():
        assert {
            blocker["code"]
            for blocker in rollout_profile_blockers(instances, profile=profile)
        } == expected_codes


def test_batch_place_suspends_on_conflict_and_revalidates_exact_plan():
    with _environment() as (db, _admin, roots, root_path):
        source_path = root_path / "source" / "source.jpg"
        source_path.write_bytes(b"p2-batch-place")
        stat = source_path.stat()
        item = {
            "id": "image-1",
            "root_id": roots["source"].root_id,
            "relative_path": source_path.name,
            "fingerprint": f"{stat.st_size}:{stat.st_mtime_ns}",
            "name": source_path.name,
        }
        target_dir = root_path / "publish" / "images" / "P2-BATCH"
        target_dir.mkdir(parents=True)
        target_path = target_dir / "P2-BATCH-sample.jpg"
        target_path.write_bytes(b"existing")
        base_context = {
            "db": db,
            "run": SimpleNamespace(inspection_number="P2-BATCH"),
            "node": {
                "config": {
                    "target_root_id": roots["publish"].root_id,
                    "target_directory": "images",
                    "naming": "inspection-sample-index",
                }
            },
        }
        suspended = _batch_place(
            SimpleNamespace(
                **base_context,
                input_data={
                    "inspection_number": "P2-BATCH",
                    "sample_identity": "sample",
                    "items": [item],
                },
            )
        )
        assert isinstance(suspended, NodeExecutionResult)
        assert suspended.suspension["resume_protocol"] == "retry_with_decision"
        plan_digest = suspended.suspension["state"]["plan_digest"]

        cancelled = _batch_place(
            SimpleNamespace(
                **base_context,
                input_data={
                    "inspection_number": "P2-BATCH",
                    "sample_identity": "sample",
                    "items": [item],
                    "_native_suspension_request": {
                        "decision": "cancel",
                        "plan_digest": plan_digest,
                    },
                },
            )
        )
        assert cancelled["placement_cancelled"] is True
        assert target_path.read_bytes() == b"existing"

        target_path.unlink()
        with pytest.raises(ExecutionApiError) as changed:
            _batch_place(
                SimpleNamespace(
                    **base_context,
                    input_data={
                        "inspection_number": "P2-BATCH",
                        "sample_identity": "sample",
                        "items": [item],
                        "_native_suspension_request": {
                            "decision": "overwrite",
                            "plan_digest": plan_digest,
                        },
                    },
                )
            )
        assert changed.value.code == "file_batch_place_plan_changed"

        completed = _batch_place(
            SimpleNamespace(
                **base_context,
                input_data={
                    "inspection_number": "P2-BATCH",
                    "sample_identity": "sample",
                    "items": [item],
                },
            )
        )
        assert completed["placement_cancelled"] is False
        assert completed["placed_count"] == 1
        assert target_path.read_bytes() == source_path.read_bytes()
