from __future__ import annotations

import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import xlwt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PillowImage
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.execution.engine import _normalize_human_submission
from app.execution.errors import ExecutionApiError
from app.execution.models import (
    ExecutionArtifact,
    ExecutionFileIndexEntry,
    ExecutionStorageRoot,
)
from app.execution.regenerated_fiber_results import (
    _result_executor,
    read_regenerated_fiber_result,
    summarize_result_inspectors,
)


COUNT_RESULT_NODE = "result.regenerated_fiber_count_method"
AREA_RESULT_NODE = "result.regenerated_fiber_area_method"


class RegeneratedFiberResultReaderTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.base = Path(self.tempdir.name)

    def tearDown(self):
        self.tempdir.cleanup()

    def _png(self, name: str, color: str) -> Path:
        path = self.base / name
        PillowImage.new("RGB", (12, 8), color=color).save(path)
        return path

    def test_count_method_reads_parts_remarks_and_all_embedded_images(self):
        path = self.base / "count.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "根数法报告1"
        sheet["I8"] = "张检验员"
        sheet["B24"] = "正面"
        sheet["G24"] = "反面"
        sheet["B25"] = "棉"
        sheet["C25"] = 0
        sheet["B26"] = 97.6
        sheet["C26"] = 2.4
        sheet["G25"] = "粘纤"
        sheet["G26"] = 0
        sheet["B27"] = "备注一"
        sheet["B28"] = "  "
        sheet["B29"] = "备注二"
        first_image = OpenpyxlImage(str(self._png("first.png", "red")))
        first_image.anchor = "L24"
        sheet.add_image(first_image)
        other = workbook.create_sheet("其它工作表")
        second_image = OpenpyxlImage(str(self._png("second.png", "blue")))
        second_image.anchor = "A1"
        other.add_image(second_image)
        workbook.save(path)
        workbook.close()

        result, images = read_regenerated_fiber_result(
            path,
            node_type=COUNT_RESULT_NODE,
        )

        self.assertTrue(result["has_parts"])
        self.assertEqual(
            result["inspector"],
            {"cell": "I8", "name": "张检验员"},
        )
        self.assertEqual(
            [(part["name"], part["total"]) for part in result["parts"]],
            [("正面", 97.6), ("反面", 0.0)],
        )
        self.assertEqual(
            result["parts"][0]["components"],
            [
                {
                    "name": "棉",
                    "content": 97.6,
                    "display_content": 97.6,
                    "raw_content": 97.6,
                    "name_cell": "B25",
                    "content_cell": "B26",
                }
            ],
        )
        self.assertEqual(
            result["remarks"],
            [
                {"cell": "B27", "text": "备注一"},
                {"cell": "B29", "text": "备注二"},
            ],
        )
        self.assertEqual(result["image_count"], 2)
        self.assertEqual(
            [(item.sheet_name, item.anchor) for item in images],
            [("根数法报告1", "L24"), ("其它工作表", "A1")],
        )
        self.assertTrue(
            any(
                warning["code"] == "content_without_component_name"
                for warning in result["warnings"]
            )
        )

    def test_no_part_names_keep_two_independent_result_groups(self):
        path = self.base / "count-two-results.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "根数法报告1"
        sheet["B25"] = "棉"
        sheet["C25"] = "粘纤"
        sheet["B26"] = 69.9526
        sheet["B26"].number_format = "0.0_ "
        sheet["C26"] = 30
        sheet["G25"] = "棉"
        sheet["H25"] = "粘纤"
        sheet["G26"] = 60
        sheet["H26"] = 40
        workbook.save(path)
        workbook.close()

        result, _images = read_regenerated_fiber_result(
            path,
            node_type=COUNT_RESULT_NODE,
        )

        self.assertFalse(result["has_parts"])
        self.assertEqual(
            [
                (part["key"], part["label"], part["total"])
                for part in result["parts"]
            ],
            [
                ("left", "结果1", 100.0),
                ("right", "结果2", 100.0),
            ],
        )
        self.assertEqual(result["parts"][0]["display_total"], 100.0)
        self.assertEqual(result["parts"][0]["raw_total"], 99.9526)

    def test_canonical_contents_keep_exact_total_when_display_rounding_is_100_1(
        self,
    ):
        path = self.base / "count-rounded-total.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "根数法报告1"
        for column, name, content in (
            ("B", "棉", 33.35),
            ("C", "粘纤", 33.35),
            ("D", "莱赛尔", 33.3),
        ):
            sheet[f"{column}25"] = name
            sheet[f"{column}26"] = content
            sheet[f"{column}26"].number_format = "0.0"
        workbook.save(path)
        workbook.close()

        result, _images = read_regenerated_fiber_result(
            path,
            node_type=COUNT_RESULT_NODE,
        )

        part = result["parts"][0]
        self.assertEqual(part["total"], 100.0)
        self.assertEqual(part["raw_total"], 100.0)
        self.assertEqual(part["display_total"], 100.1)
        self.assertEqual(
            [component["content"] for component in part["components"]],
            [33.3, 33.4, 33.3],
        )
        self.assertEqual(
            [
                component["display_content"]
                for component in part["components"]
            ],
            [33.4, 33.4, 33.3],
        )
        self.assertEqual(
            part["rounding_adjustment"],
            {
                "component_name": "棉",
                "delta": -0.1,
                "excel_display_total": 100.1,
            },
        )

    def test_area_method_keeps_one_unnamed_group_and_numeric_percentages(self):
        path = self.base / "area.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "截面统计报告1"
        sheet["I8"] = "李检验员"
        sheet["B27"] = "棉"
        sheet["C27"] = "粘纤"
        sheet["B28"] = 0.976
        sheet["B28"].number_format = "0.0%"
        sheet["C28"] = "2.4%"
        sheet["B29"] = "面积法备注"
        workbook.save(path)
        workbook.close()

        result, _images = read_regenerated_fiber_result(
            path,
            node_type=AREA_RESULT_NODE,
        )

        self.assertFalse(result["has_parts"])
        self.assertEqual(
            result["inspector"],
            {"cell": "I8", "name": "李检验员"},
        )
        self.assertEqual(len(result["parts"]), 1)
        self.assertEqual(result["parts"][0]["label"], "结果")
        self.assertEqual(result["parts"][0]["total"], 100.0)
        self.assertEqual(
            [
                component["content"]
                for component in result["parts"][0]["components"]
            ],
            [97.6, 2.4],
        )

    def test_legacy_xls_uses_cached_values_and_preserves_zero_content(self):
        path = self.base / "legacy.xls"
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("根数法报告1")
        sheet.write(7, 8, "王检验员")
        sheet.write(23, 1, "部位")
        sheet.write(24, 1, "棉")
        sheet.write(24, 2, 0)
        sheet.write(25, 1, 0)
        sheet.write(26, 1, "旧版备注")
        workbook.save(str(path))

        with patch(
            "app.execution.regenerated_fiber_results."
            "_legacy_images_via_libreoffice",
            return_value=([], []),
        ):
            result, images = read_regenerated_fiber_result(
                path,
                node_type=COUNT_RESULT_NODE,
            )

        self.assertEqual(images, [])
        self.assertEqual(
            result["inspector"],
            {"cell": "I8", "name": "王检验员"},
        )
        self.assertEqual(
            result["parts"][0]["components"],
            [
                {
                    "name": "棉",
                    "content": 0.0,
                    "display_content": 0.0,
                    "raw_content": 0.0,
                    "name_cell": "B25",
                    "content_cell": "B26",
                }
            ],
        )
        self.assertEqual(
            result["remarks"],
            [{"cell": "B27", "text": "旧版备注"}],
        )

    def test_ooxml_workbook_renamed_to_xls_uses_actual_container_format(self):
        original = self.base / "renamed-source.xlsx"
        path = self.base / "renamed.xls"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "根数法报告1"
        sheet["I8"] = "赵检验员"
        sheet["B25"] = "棉"
        sheet["B26"] = 100
        workbook.save(original)
        workbook.close()
        original.replace(path)

        result, images = read_regenerated_fiber_result(
            path,
            node_type=COUNT_RESULT_NODE,
        )

        self.assertEqual(images, [])
        self.assertEqual(
            result["inspector"],
            {"cell": "I8", "name": "赵检验员"},
        )
        self.assertEqual(result["parts"][0]["total"], 100.0)

    def test_inspector_summary_exposes_multi_file_conflict(self):
        summary = summarize_result_inspectors(
            [
                {
                    "id": "first",
                    "read_status": "succeeded",
                    "result": {
                        "inspector": {"cell": "I8", "name": "张三"}
                    },
                },
                {
                    "id": "second",
                    "read_status": "succeeded",
                    "result": {
                        "inspector": {"cell": "I8", "name": " 李四 "}
                    },
                },
                {
                    "id": "third",
                    "read_status": "succeeded",
                    "result": {
                        "inspector": {"cell": "I8", "name": "张三"}
                    },
                },
                {
                    "id": "blank",
                    "read_status": "succeeded",
                    "result": {"inspector": {"cell": "I8", "name": " "}},
                },
                {
                    "id": "failed",
                    "read_status": "failed",
                    "error": {"code": "read_failed"},
                },
            ]
        )

        self.assertTrue(summary["conflict"])
        self.assertIsNone(summary["name"])
        self.assertEqual(summary["names"], ["张三", "李四"])
        self.assertEqual(summary["missing_count"], 1)
        self.assertEqual(
            summary["files"],
            [
                {"file_id": "first", "name": "张三"},
                {"file_id": "second", "name": "李四"},
                {"file_id": "third", "name": "张三"},
                {"file_id": "blank", "name": None},
            ],
        )


class RegeneratedFiberResultExecutorTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        base = Path(self.tempdir.name)
        self.source = base / "source"
        self.staging = base / "staging"
        self.source.mkdir()
        self.staging.mkdir()
        self.engine = create_engine("sqlite:///:memory:")
        self.Session = sessionmaker(bind=self.engine, autoflush=False)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        self.source_root = ExecutionStorageRoot(
            root_id="regenerated_fiber_records",
            name="再生纤",
            local_path=str(self.source),
            access_mode="read",
            category_key="regenerated_fiber",
            is_active=True,
            is_available=True,
        )
        self.staging_root = ExecutionStorageRoot(
            root_id="execution_staging",
            name="暂存",
            local_path=str(self.staging),
            access_mode="write",
            is_active=True,
            is_available=True,
        )
        self.db.add_all([self.source_root, self.staging_root])
        self.db.flush()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        self.tempdir.cleanup()

    def _indexed_workbook(
        self,
        *,
        image_count: int = 1,
    ) -> tuple[Path, ExecutionFileIndexEntry]:
        path = self.source / "262039607.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "截面统计报告1"
        sheet["B27"] = "棉"
        sheet["B28"] = 100
        image_paths: list[Path] = []
        for index in range(image_count):
            png = self.source / f"preview-source-{index}.png"
            PillowImage.new(
                "RGB",
                (10, 6),
                color=("green" if index == 0 else "blue"),
            ).save(png)
            image_paths.append(png)
            embedded_image = OpenpyxlImage(str(png))
            sheet.add_image(embedded_image, f"L{26 + index}")
        workbook.save(path)
        workbook.close()
        for png in image_paths:
            png.unlink()
        stat = path.stat()
        entry = ExecutionFileIndexEntry(
            storage_root_id=self.source_root.id,
            relative_path=path.name,
            filename=path.name,
            extension=".xlsx",
            file_kind="workbook",
            size_bytes=stat.st_size,
            modified_at=datetime.fromtimestamp(stat.st_mtime, timezone.utc),
            fingerprint=f"{stat.st_size}:{stat.st_mtime_ns}",
            metadata_json={},
        )
        self.db.add(entry)
        self.db.flush()
        return path, entry

    def test_executor_registers_preview_artifacts_and_returns_safe_urls(self):
        _path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "name": entry.filename,
            "fingerprint": entry.fingerprint,
        }
        context = SimpleNamespace(
            db=self.db,
            run=SimpleNamespace(id="run-result-reader"),
            node_run=SimpleNamespace(
                id="node-result-reader",
                node_type=AREA_RESULT_NODE,
            ),
            input_data={"files": [candidate]},
        )

        output = _result_executor(context)

        self.assertEqual(output["success_count"], 1)
        image = output["files"][0]["result"]["images"][0]
        self.assertRegex(
            image["preview_url"],
            r"^/api/execution/v1/artifacts/.+/preview$",
        )
        artifact = self.db.get(ExecutionArtifact, image["artifact_id"])
        self.assertIsNotNone(artifact)
        self.assertEqual(artifact.role, "preview")
        self.assertTrue((self.staging / artifact.relative_path).is_file())

    def test_executor_reads_snapshot_and_rejects_source_changed_during_parse(self):
        path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "name": entry.filename,
            "fingerprint": entry.fingerprint,
        }
        context = SimpleNamespace(
            db=self.db,
            run=SimpleNamespace(id="run-snapshot"),
            node_run=SimpleNamespace(
                id="node-snapshot",
                node_type=AREA_RESULT_NODE,
            ),
            input_data={"files": [candidate]},
        )
        seen_paths: list[Path] = []

        def change_source(snapshot: Path, *, node_type: str):
            self.assertEqual(node_type, AREA_RESULT_NODE)
            self.assertNotEqual(snapshot, path)
            self.assertTrue(snapshot.is_file())
            seen_paths.append(snapshot)
            path.write_bytes(path.read_bytes() + b"changed")
            return (
                {
                    "warnings": [],
                    "image_count": 0,
                    "parts": [],
                    "remarks": [],
                },
                [],
            )

        with patch(
            "app.execution.regenerated_fiber_results."
            "read_regenerated_fiber_result",
            side_effect=change_source,
        ):
            with self.assertRaises(ExecutionApiError) as raised:
                _result_executor(context)

        self.assertEqual(raised.exception.code, "result_workbooks_unreadable")
        self.assertEqual(len(seen_paths), 1)
        self.assertFalse(seen_paths[0].exists())
        self.assertEqual(self.db.query(ExecutionArtifact).count(), 0)
        self.assertEqual(
            [item for item in self.staging.rglob("*") if item.is_file()],
            [],
        )

    def test_candidate_image_failure_rolls_back_rows_and_created_files(self):
        _path, entry = self._indexed_workbook(image_count=2)
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "name": entry.filename,
            "fingerprint": entry.fingerprint,
        }
        context = SimpleNamespace(
            db=self.db,
            run=SimpleNamespace(id="run-image-rollback"),
            node_run=SimpleNamespace(
                id="node-image-rollback",
                node_type=AREA_RESULT_NODE,
            ),
            input_data={"files": [candidate]},
        )
        from app.execution import regenerated_fiber_results as result_module

        original_write = result_module._atomic_image_write
        calls = 0

        def fail_second_write(target: Path, data: bytes) -> bool:
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("simulated image write failure")
            return original_write(target, data)

        with patch(
            "app.execution.regenerated_fiber_results._atomic_image_write",
            side_effect=fail_second_write,
        ):
            with self.assertRaises(ExecutionApiError) as raised:
                _result_executor(context)

        self.assertEqual(raised.exception.code, "result_workbooks_unreadable")
        self.assertEqual(self.db.query(ExecutionArtifact).count(), 0)
        self.assertEqual(
            [item for item in self.staging.rglob("*") if item.is_file()],
            [],
        )

    def test_executor_reuses_registered_image_artifact(self):
        _path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "name": entry.filename,
            "fingerprint": entry.fingerprint,
        }
        context = SimpleNamespace(
            db=self.db,
            run=SimpleNamespace(id="run-result-retry"),
            node_run=SimpleNamespace(
                id="node-result-retry",
                node_type=AREA_RESULT_NODE,
            ),
            input_data={"files": [candidate]},
        )

        first = _result_executor(context)
        second = _result_executor(context)

        self.assertEqual(self.db.query(ExecutionArtifact).count(), 1)
        self.assertEqual(
            first["files"][0]["result"]["images"][0]["artifact_id"],
            second["files"][0]["result"]["images"][0]["artifact_id"],
        )

    def test_executor_repairs_missing_file_for_registered_artifact(self):
        _path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "name": entry.filename,
            "fingerprint": entry.fingerprint,
        }
        context = SimpleNamespace(
            db=self.db,
            run=SimpleNamespace(id="run-result-repair"),
            node_run=SimpleNamespace(
                id="node-result-repair",
                node_type=AREA_RESULT_NODE,
            ),
            input_data={"files": [candidate]},
        )
        first = _result_executor(context)
        artifact_id = first["files"][0]["result"]["images"][0]["artifact_id"]
        artifact = self.db.get(ExecutionArtifact, artifact_id)
        target = self.staging / artifact.relative_path
        target.unlink()

        second = _result_executor(context)

        self.assertTrue(target.is_file())
        self.assertEqual(self.db.query(ExecutionArtifact).count(), 1)
        self.assertEqual(
            second["files"][0]["result"]["images"][0]["artifact_id"],
            artifact_id,
        )

    def test_primary_file_is_validated_and_single_selection_is_automatic(self):
        path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "fingerprint": entry.fingerprint,
            "read_status": "succeeded",
        }
        definition = {
            "root_slots": [
                {
                    "root_id": "regenerated_fiber_records",
                    "access": "read",
                }
            ],
            "nodes": [
                {
                    "id": "select",
                    "type": "human.file_selection",
                    "config": {"require_primary": True},
                }
            ],
        }
        run = SimpleNamespace(definition_snapshot=definition)
        node_run = SimpleNamespace(
            node_type="human.file_selection",
            node_id="select",
            input_data={"files": [candidate]},
        )

        normalized = _normalize_human_submission(
            self.db,
            run=run,
            node_run=node_run,
            data={"selected_files": [entry.id]},
        )

        self.assertEqual(normalized["primary_file_id"], entry.id)
        self.assertEqual(normalized["primary_file"]["id"], entry.id)

        with self.assertRaises(ExecutionApiError) as raised:
            _normalize_human_submission(
                self.db,
                run=run,
                node_run=node_run,
                data={
                    "selected_files": [entry.id],
                    "primary_file_id": "not-selected",
                },
            )
        self.assertEqual(raised.exception.code, "primary_file_not_selected")

        path.write_bytes(path.read_bytes() + b"changed-after-read")
        with self.assertRaises(ExecutionApiError) as stale:
            _normalize_human_submission(
                self.db,
                run=run,
                node_run=node_run,
                data={"selected_files": [entry.id]},
            )
        self.assertEqual(stale.exception.code, "file_candidate_stale")

    def test_failed_result_file_cannot_be_selected(self):
        _path, entry = self._indexed_workbook()
        candidate = {
            "id": entry.id,
            "root_id": self.source_root.root_id,
            "relative_path": entry.relative_path,
            "fingerprint": entry.fingerprint,
            "read_status": "failed",
        }
        run = SimpleNamespace(
            definition_snapshot={
                "root_slots": [
                    {"root_id": "regenerated_fiber_records", "access": "read"}
                ],
                "nodes": [
                    {
                        "id": "select",
                        "type": "human.file_selection",
                        "config": {},
                    }
                ],
            }
        )
        node_run = SimpleNamespace(
            node_type="human.file_selection",
            node_id="select",
            input_data={"files": [candidate]},
        )

        with self.assertRaises(ExecutionApiError) as raised:
            _normalize_human_submission(
                self.db,
                run=run,
                node_run=node_run,
                data={"selected_files": [entry.id]},
            )
        self.assertEqual(raised.exception.code, "result_file_not_selectable")


if __name__ == "__main__":
    unittest.main()
