from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.execution.errors import ExecutionApiError
from app.execution.v2.data_handlers import transform, validate, python_compute
from app.execution.v2.designer import starter_document, compile_document
from app.execution.v2.workbook_render import render_file
from tests.test_execution_workflow_replacement import environment, stage, publish, run, drain, request


def context(config, data):
    return SimpleNamespace(node={"config": config}, input_data=data)


def test_transform_projects_filters_maps_defaults_and_deduplicates():
    result = transform(context({"steps": [
        {"op": "select", "path": "#/rows"}, {"op": "filter", "path": "#/accepted", "equals": True},
        {"op": "map", "fields": {"number": {"path": "#/id"}, "unit": {"path": "#/unit", "default": "μm"}}},
        {"op": "unique", "path": "#/number"}, {"op": "sort", "path": "#/number"}, {"op": "limit", "count": 2},
    ]}, {"data": {"rows": [{"id": n, "accepted": n != 2} for n in [3, 1, 1, 2]]}}))
    assert result == {"data": [{"number": 1, "unit": "μm"}, {"number": 3, "unit": "μm"}]}


def test_validation_can_fail_or_route_errors_without_another_approval():
    schema = {"type": "object", "properties": {"count": {"minimum": 1}}, "required": ["count"]}
    assert validate(context({"schema": schema, "fail_on_error": False}, {"data": {"count": 0}}))["valid"] is False
    with pytest.raises(ExecutionApiError):
        validate(context({"schema": schema}, {"data": {"count": 0}}))


def python_config(code="def main(inputs):\n    return {'total': sum(inputs['values'])}"):
    return {"runtime_version": 1, "code": code, "timeout_seconds": 1,
        "input_schema": {"type": "object", "properties": {"values": {"type": "array", "items": {"type": "number"}}}, "required": ["values"], "additionalProperties": False},
        "output_schema": {"type": "object", "properties": {"total": {"type": "number"}}, "required": ["total"], "additionalProperties": False}}


def test_python_is_typed_resource_bounded_and_has_no_application_imports():
    config = python_config("import statistics\ndef main(inputs):\n    return {'total': statistics.mean(inputs['values'])}")
    assert python_compute(context(config, {"values": [1, 3]})) == {"total": 2}
    with pytest.raises(ValueError, match="计算模块"):
        python_compute(context(python_config("import os\ndef main(inputs):\n    return {}"), {"values": []}))
    with pytest.raises(ExecutionApiError, match="超过执行时间"):
        python_compute(context(python_config("def main(inputs):\n    while True: pass"), {"values": []}))
    with pytest.raises(ExecutionApiError):
        python_compute(context(python_config("def main(inputs):\n    return {'total': float('nan')}"), {"values": []}))


def test_new_json_python_graph_publishes_and_runs_with_worker(environment):
    doc = starter_document()
    config = python_config()
    doc["definition"]["input_schema"] = {**config["input_schema"], "properties": {**config["input_schema"]["properties"], "inspection_number": {"type": "string"}}}
    doc["definition"]["output_schema"] = config["output_schema"]
    doc["definition"]["nodes"].insert(1, {"id": "compute", "type": "data.python", "type_version": 1,
        "name": "合计", "config": config, "input_mapping": {"values": "$.inputs.values"}})
    doc["definition"]["nodes"][-1]["input_mapping"] = {"total": "$.nodes.compute.output.total"}
    doc["definition"]["edges"] = [{"id": "a", "source": "start", "target": "compute", "join_policy": "all"},
        {"id": "b", "source": "compute", "target": "end", "join_policy": "all"}]
    compiled = compile_document(doc)
    assert compiled["content_valid"], compiled["issues"]
    published = publish(environment, stage(environment, compiled["document"]))
    task = run(environment, published["workflow_id"], inputs={"values": [1, 2, 3]})
    drain(environment)
    detail = request(environment, "GET", f"v1/runs/{task['id']}")
    assert detail["status"] == "completed", detail
    assert detail["output_data"] == {"total": 6}


def test_workbook_mapping_changes_only_json_and_preserves_template(tmp_path):
    source, target = tmp_path / "template.xlsx", tmp_path / "result.xlsx"
    book = Workbook()
    book.active["A1"] = "原始标题"
    book.active["B1"] = "=1+1"
    book.save(source)
    result = render_file(source, target, values={"title": "新项目", "number": "=not-a-formula"}, fields=[
        {"cell": "A1", "value": {"path": "#/title"}}, {"cell": "A2", "value": {"path": "#/number"}},
    ])
    assert result["cells_verified"] == 2
    book = load_workbook(target)
    assert book.active["A1"].value == "新项目"
    assert book.active["A2"].data_type == "s"
    assert book.active["B1"].value == "=1+1"
    book.close()
    original = load_workbook(source)
    assert original.active["A1"].value == "原始标题"
    original.close()


def test_render_node_publishes_to_bound_directory_with_receipt(environment):
    import hashlib
    from app.execution.models import ExecutionArtifact

    env = environment
    source = env.path / "electron_microscopy_records" / "form.xlsx"
    book = Workbook()
    book.save(source)
    doc = starter_document()
    doc["resources"]["root_slots"] = [
        {"slot_id": "electron_microscopy_records", "name": "模板", "access": "read", "required": True},
        {"slot_id": "execution_staging", "name": "输出", "access": "write", "required": True},
    ]
    doc["definition"]["nodes"].insert(1, {"id": "render", "name": "生成记录", "type": "workbook.render", "type_version": 1,
        "config": {"template": {"root_slot": "electron_microscopy_records", "relative_path": "form.xlsx", "sha256": hashlib.sha256(source.read_bytes()).hexdigest()},
                   "staging_root_slot": "execution_staging", "fields": [{"cell": "B2", "value": {"path": "#/label"}}]},
        "input_mapping": {"values": {"label": "JSON 配置"}}})
    doc["definition"]["edges"] = [{"id": "a", "source": "start", "target": "render", "join_policy": "all"},
        {"id": "b", "source": "render", "target": "end", "join_policy": "all"}]
    compiled = compile_document(doc)
    assert compiled["content_valid"], compiled["issues"]
    published = publish(env, stage(env, compiled["document"]))
    task = run(env, published["workflow_id"])
    drain(env)
    detail = request(env, "GET", f"v1/runs/{task['id']}")
    assert detail["status"] == "completed", detail
    artifact = env.db.query(ExecutionArtifact).filter_by(run_id=task["id"]).one()
    book = load_workbook(env.path / "execution_staging" / artifact.relative_path)
    assert book.active["B2"].value == "JSON 配置"
    book.close()


def test_domain_profile_is_self_contained_and_can_change_family_rules():
    from app.execution.v2.examples import build_domain_records_smoke_release
    from app.execution.microscopy_families import microscopy_family_from_config

    result = compile_document(build_domain_records_smoke_release())
    assert result["content_valid"], result["issues"]
    config = next(n["config"] for n in result["document"]["definition"]["nodes"] if n["type"] == "microscopy.image_candidates")
    config["record_family"] = config["family_profile"]["key"] = "custom_microscopy"
    config["family_profile"]["record_title"] = "新项目原始记录"
    config["family_profile"]["max_selected_images"] = 8
    family = microscopy_family_from_config(config)
    assert family.key == "custom_microscopy"
    assert family.record_title == "新项目原始记录"
    assert family.max_selected_images == 8
