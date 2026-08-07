from __future__ import annotations

import math
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.execution.electron_microscopy import cached_task_snapshot
from app.execution.errors import ExecutionApiError
from app.execution.index_metadata import inspection_numbers_in_text
from app.execution.models import (
    ExecutionFileIndexEntry,
    ExecutionStorageRoot,
    utcnow,
)
from app.execution.registry import node_registry
from app.execution.storage import ArtifactRef, FileGateway, StorageError
from app.execution.workbook_format import (
    SUPPORTED_WORKBOOK_SUFFIXES,
    WorkbookFormat,
    detect_workbook_format,
)


PAPER_FIBER_ROOT_ID = "paper_fiber_records"
PAPER_FIBER_NODE_TYPE = "file.paper_fiber_gbt4688_qualitative"
PAPER_FIBER_WORKFLOW_SLUG = "paper-fiber-gbt4688-2020-qualitative"
PAPER_FIBER_PROJECT_NAME = "纸、纸板和纸浆纤维鉴别分析"
PAPER_FIBER_TEST_METHOD = "GB/T 4688-2020"
PAPER_FIBER_WORKSHEET = "Sheet1"
PAPER_FIBER_RESULT_CELL = "W32"
PAPER_FIBER_PROFILE_VERSION = 1
MAX_WORKBOOK_VALIDATION_MATCHES = 50
_PROFILE_CACHE_KEY = (
    f"{PAPER_FIBER_NODE_TYPE}@{PAPER_FIBER_PROFILE_VERSION}"
)
_DETERMINISTIC_PROFILE_STATUSES = {
    "matched",
    "worksheet_missing",
    "result_empty",
    "unsupported_format",
}
_STANDALONE_100_PATTERN = re.compile(
    r"(?<![\w.])100(?:\.0+)?(?![\w.])"
)


def _normalized_fact(value: object) -> str:
    return " ".join(
        unicodedata.normalize("NFKC", str(value or "")).strip().split()
    )


def contains_standalone_100(value: object) -> bool:
    """Return whether a saved result contains the independent number 100.

    NFKC normalization handles full-width digits. Numeric boundaries prevent
    values such as 1000, 2100 or 100.5 from being treated as a 100 percent
    result, while 100.0 remains equivalent to the integer 100.
    """

    return bool(_STANDALONE_100_PATTERN.search(_normalized_fact(value)))


def _display_value(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, str):
        normalized = value.strip()
        return normalized or None
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return str(int(value))
    normalized = str(value).strip()
    return normalized or None


def _read_result_profile(path: Path) -> dict[str, Any]:
    profile: dict[str, Any] = {
        "profile_version": PAPER_FIBER_PROFILE_VERSION,
        "worksheet": PAPER_FIBER_WORKSHEET,
        "result_cell": PAPER_FIBER_RESULT_CELL,
        "worksheet_exists": False,
        "qualitative_result": None,
        "contains_standalone_100": False,
        "unit": "",
    }
    if path.suffix.casefold() not in SUPPORTED_WORKBOOK_SUFFIXES:
        profile["status"] = "unsupported_format"
        return profile

    try:
        workbook_format = detect_workbook_format(path)
        profile["workbook_format"] = (
            workbook_format.value if workbook_format is not None else None
        )
        if workbook_format is None:
            profile["status"] = "unsupported_format"
            return profile
        if workbook_format is WorkbookFormat.OLE:
            try:
                import xlrd
            except ImportError:
                profile["status"] = "reader_unavailable"
                return profile
            workbook = xlrd.open_workbook(
                str(path),
                on_demand=True,
                formatting_info=False,
            )
            try:
                if PAPER_FIBER_WORKSHEET not in workbook.sheet_names():
                    profile["status"] = "worksheet_missing"
                    return profile
                profile["worksheet_exists"] = True
                sheet = workbook.sheet_by_name(PAPER_FIBER_WORKSHEET)
                raw_value = (
                    sheet.cell_value(31, 22)
                    if sheet.nrows > 31 and sheet.ncols > 22
                    else None
                )
            finally:
                workbook.release_resources()
        else:
            with path.open("rb") as stream:
                workbook = load_workbook(
                    stream,
                    read_only=True,
                    data_only=True,
                    keep_vba=False,
                    keep_links=False,
                )
                try:
                    if PAPER_FIBER_WORKSHEET not in workbook.sheetnames:
                        profile["status"] = "worksheet_missing"
                        return profile
                    profile["worksheet_exists"] = True
                    raw_value = workbook[PAPER_FIBER_WORKSHEET][
                        PAPER_FIBER_RESULT_CELL
                    ].value
                finally:
                    workbook.close()
    except Exception as exc:
        profile["status"] = "read_failed"
        profile["error_type"] = type(exc).__name__
        return profile

    result = _display_value(raw_value)
    if result is None:
        profile["status"] = "result_empty"
        return profile
    has_100 = contains_standalone_100(result)
    profile.update(
        {
            "status": "matched",
            "qualitative_result": result,
            "contains_standalone_100": has_100,
            "unit": "%" if has_100 else "",
        }
    )
    return profile


def _profile_for_entry(
    entry: ExecutionFileIndexEntry,
    *,
    gateway: FileGateway,
) -> tuple[dict[str, Any], bool]:
    metadata = dict(entry.metadata_json or {})
    profiles = dict(metadata.get("workbook_profiles") or {})
    # 鲜活度护栏：后台索引按间隔扫描，共享盘上的最新保存可能尚未反映到
    # entry.fingerprint。候选一旦携带过期指纹，提交时
    # _validate_index_candidate 会以 file_candidate_stale 永久拒绝。
    # 这里先按实时 stat 校正指纹（连带 size/modified_at），让候选始终
    # 与磁盘一致；指纹变化会自动使下方缓存失效并重读结果单元格。
    try:
        live_path = gateway.resolve(
            ArtifactRef(PAPER_FIBER_ROOT_ID, entry.relative_path),
            expected_type="file",
        )
        live_stat = live_path.stat()
    except (ExecutionApiError, StorageError, OSError):
        live_stat = None
    live_fingerprint = (
        f"{live_stat.st_size}:{live_stat.st_mtime_ns}"
        if live_stat is not None
        else None
    )
    if live_fingerprint is not None and live_fingerprint != entry.fingerprint:
        entry.fingerprint = live_fingerprint
        entry.size_bytes = live_stat.st_size
        entry.modified_at = datetime.fromtimestamp(
            live_stat.st_mtime, tz=timezone.utc
        )
    cached = profiles.get(_PROFILE_CACHE_KEY)
    if (
        isinstance(cached, dict)
        and cached.get("profile_version") == PAPER_FIBER_PROFILE_VERSION
        and cached.get("fingerprint") == entry.fingerprint
        and cached.get("status") in _DETERMINISTIC_PROFILE_STATUSES
    ):
        return dict(cached), False

    original_fingerprint = entry.fingerprint
    try:
        path = gateway.resolve(
            ArtifactRef(PAPER_FIBER_ROOT_ID, entry.relative_path),
            expected_type="file",
        )
        profile = _read_result_profile(path)
    except (ExecutionApiError, StorageError, OSError) as exc:
        profile = {
            "profile_version": PAPER_FIBER_PROFILE_VERSION,
            "worksheet": PAPER_FIBER_WORKSHEET,
            "result_cell": PAPER_FIBER_RESULT_CELL,
            "worksheet_exists": False,
            "qualitative_result": None,
            "contains_standalone_100": False,
            "unit": "",
            "status": "file_unavailable",
            "error_type": type(exc).__name__,
        }
    profile.update(
        {
            "profile_version": PAPER_FIBER_PROFILE_VERSION,
            "fingerprint": entry.fingerprint,
            "checked_at": utcnow().isoformat(),
        }
    )
    changed = False
    if (
        entry.fingerprint == original_fingerprint
        and profile.get("status") in _DETERMINISTIC_PROFILE_STATUSES
    ):
        if profiles.get(_PROFILE_CACHE_KEY) != profile:
            profiles[_PROFILE_CACHE_KEY] = profile
            changed = True
    elif _PROFILE_CACHE_KEY in profiles:
        profiles.pop(_PROFILE_CACHE_KEY, None)
        changed = True
    if changed:
        metadata["workbook_profiles"] = profiles
        entry.metadata_json = metadata
    return profile, changed


def _escaped_contains(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("%", "\\%")
        .replace("_", "\\_")
    )
    return f"%{escaped}%"


def _is_complete_inspection_number(value: str) -> bool:
    matches = inspection_numbers_in_text(value)
    return len(matches) == 1 and matches[0].casefold() == value.casefold()


def _first_path_component(db: Session):
    if db.get_bind().dialect.name == "postgresql":
        return func.split_part(ExecutionFileIndexEntry.relative_path, "/", 1)
    first_slash = func.instr(ExecutionFileIndexEntry.relative_path, "/")
    return func.substr(
        ExecutionFileIndexEntry.relative_path,
        1,
        first_slash - 1,
    )


def _indexed_entries(
    db: Session,
    *,
    root: ExecutionStorageRoot,
    inspection_number: str,
) -> list[ExecutionFileIndexEntry]:
    query_text = str(inspection_number or "").strip()
    if not query_text:
        return []
    first_component = _first_path_component(db)
    rows = (
        db.query(ExecutionFileIndexEntry)
        .filter(
            ExecutionFileIndexEntry.storage_root_id == root.id,
            ExecutionFileIndexEntry.missing_since.is_(None),
            ExecutionFileIndexEntry.extension.in_(
                sorted(SUPPORTED_WORKBOOK_SUFFIXES)
            ),
            ExecutionFileIndexEntry.relative_path.like("%/%"),
            first_component.ilike(
                _escaped_contains(query_text),
                escape="\\",
            ),
        )
        .order_by(
            ExecutionFileIndexEntry.modified_at.desc(),
            ExecutionFileIndexEntry.relative_path.asc(),
        )
        .all()
    )
    folded = query_text.casefold()
    return [
        entry
        for entry in rows
        if len(PurePosixPath(entry.relative_path).parts) == 2
        and folded
        in PurePosixPath(entry.relative_path).parts[0].casefold()
    ]


def _index_state(
    db: Session,
    *,
    root: Optional[ExecutionStorageRoot],
    entry_count: int,
) -> str:
    # Reuse the established regenerated-fiber state contract while keeping
    # this module's file matching independent.
    from app.execution.regenerated_fiber import _index_state as state_for_root

    return state_for_root(db, root=root, entry_count=entry_count)


def paper_task_project_match(
    snapshot: Optional[dict[str, Any]],
) -> tuple[list[str], Optional[dict[str, Any]]]:
    best_conditions: list[str] = []
    best_project: Optional[dict[str, Any]] = None
    for project in (snapshot or {}).get("projects") or []:
        if not isinstance(project, dict):
            continue
        conditions: list[str] = []
        if _normalized_fact(project.get("check_item_name")) == _normalized_fact(
            PAPER_FIBER_PROJECT_NAME
        ):
            conditions.append("task_item_name")
        if _normalized_fact(project.get("check_method")) == _normalized_fact(
            PAPER_FIBER_TEST_METHOD
        ):
            conditions.append("test_method")
        if len(conditions) > len(best_conditions):
            best_conditions = conditions
            best_project = dict(project)
        if len(best_conditions) == 2:
            break
    return best_conditions, best_project


def paper_fiber_match(
    db: Session,
    *,
    inspection_number: str,
    gateway: Optional[FileGateway] = None,
    result_limit: int = 6,
) -> dict[str, Any]:
    query_text = str(inspection_number or "").strip()
    root = (
        db.query(ExecutionStorageRoot)
        .filter(ExecutionStorageRoot.root_id == PAPER_FIBER_ROOT_ID)
        .one_or_none()
    )
    root_ready = bool(root and root.is_active and root.is_available)
    base_entry_count = (
        db.query(ExecutionFileIndexEntry.id)
        .filter(
            ExecutionFileIndexEntry.storage_root_id == root.id,
            ExecutionFileIndexEntry.missing_since.is_(None),
        )
        .count()
        if root is not None
        else 0
    )
    index_state = _index_state(
        db,
        root=root,
        entry_count=base_entry_count,
    )
    entries = (
        _indexed_entries(
            db,
            root=root,
            inspection_number=query_text,
        )
        if root_ready and root is not None
        else []
    )
    query_complete = _is_complete_inspection_number(query_text)
    task = cached_task_snapshot(db, inspection_number=query_text)
    task_conditions, matched_project = paper_task_project_match(
        task.get("snapshot")
    )
    matched_conditions: list[str] = []
    if root_ready:
        matched_conditions.append("source_root")
    if entries:
        matched_conditions.append("folder")
    matched_conditions.extend(task_conditions)
    full_match = all(
        item in matched_conditions
        for item in ("source_root", "folder", "task_item_name", "test_method")
    )

    profiles_updated = False
    candidates: list[dict[str, Any]] = []
    query_state = (
        "empty"
        if not query_text
        else "complete"
        if query_complete
        else "incomplete"
    )
    if query_complete and len(entries) > MAX_WORKBOOK_VALIDATION_MATCHES:
        query_state = "too_many_matches"
    elif query_complete and entries:
        if gateway is None:
            from app.execution.persistence import build_file_gateway

            gateway = build_file_gateway(db)
        for entry in entries:
            profile, changed = _profile_for_entry(entry, gateway=gateway)
            profiles_updated = profiles_updated or changed
            if profile.get("status") != "matched":
                continue
            parts = PurePosixPath(entry.relative_path).parts
            result = str(profile["qualitative_result"])
            candidates.append(
                {
                    "id": entry.id,
                    "root_id": PAPER_FIBER_ROOT_ID,
                    "relative_path": entry.relative_path,
                    "name": entry.filename,
                    "suffix": entry.extension,
                    "size": entry.size_bytes,
                    "modified_at": entry.modified_at.isoformat(),
                    "category": root.category_key if root is not None else None,
                    "fingerprint": entry.fingerprint,
                    "folder_name": parts[0],
                    "read_status": "succeeded",
                    "qualitative_result": result,
                    "result": {
                        "worksheet": PAPER_FIBER_WORKSHEET,
                        "cell": PAPER_FIBER_RESULT_CELL,
                        "w32_value": result,
                        "qualitative_result": result,
                        "contains_standalone_100": bool(
                            profile["contains_standalone_100"]
                        ),
                        "unit": str(profile["unit"]),
                    },
                }
            )
        query_state = "validated"

    limit = max(1, min(int(result_limit), 6))
    candidates = candidates[:limit]
    preview_entry = candidates[0] if candidates else (entries[0] if entries else None)
    if isinstance(preview_entry, ExecutionFileIndexEntry):
        candidate_preview = {
            "name": preview_entry.filename,
            "relative_path": preview_entry.relative_path,
            "suffix": preview_entry.extension,
        }
    elif isinstance(preview_entry, dict):
        preview_result = dict(preview_entry.get("result") or {})
        candidate_preview = {
            "name": str(preview_entry["name"]),
            "relative_path": str(preview_entry["relative_path"]),
            "suffix": str(preview_entry["suffix"]),
            "qualitative_result": str(preview_entry["qualitative_result"]),
            "result": {
                "worksheet": str(preview_result["worksheet"]),
                "cell": str(preview_result["cell"]),
                "w32_value": str(preview_result["w32_value"]),
                "unit": str(preview_result["unit"]),
            },
        }
    else:
        candidate_preview = None
    return {
        "root_id": PAPER_FIBER_ROOT_ID,
        "index_state": index_state,
        "query_state": query_state,
        "folder_match_count": len(
            {
                PurePosixPath(entry.relative_path).parts[0]
                for entry in entries
            }
        ),
        "workbook_match_count": len(entries),
        "result_match_count": len(candidates),
        "candidates": candidates,
        "candidate_preview": candidate_preview,
        "matched_conditions": matched_conditions,
        "full_match": full_match,
        "task_cache_state": task["cache_state"],
        "task_snapshot": task.get("snapshot"),
        "matched_task_project": matched_project,
        "cache_updated": bool(task.get("refresh_queued")) or profiles_updated,
    }


def _paper_fiber_executor(context) -> dict[str, Any]:
    config = context.node.get("config") or {}
    inspection_number = str(
        context.input_data.get("inspection_number")
        or context.run.inspection_number
    ).strip()
    if not _is_complete_inspection_number(inspection_number):
        raise ExecutionApiError(
            422,
            "inspection_number_incomplete",
            "请先输入完整检验编号后再读取纸类原始记录",
        )
    match = paper_fiber_match(
        context.db,
        inspection_number=inspection_number,
        result_limit=min(int(config.get("limit", 6)), 6),
    )
    missing = [
        item
        for item in ("source_root", "folder", "task_item_name", "test_method")
        if item not in match["matched_conditions"]
    ]
    if missing and bool(config.get("require_full_task_match", False)):
        raise ExecutionApiError(
            422,
            "paper_fiber_rule_not_matched",
            "当前编号不满足 GB/T 4688-2020 纸类定性分析流程条件",
            details={"missing_conditions": missing},
        )
    if not match["candidates"]:
        raise ExecutionApiError(
            422,
            "paper_fiber_workbook_not_found",
            "匹配目录中没有可读取 Sheet1!W32 结果的原始记录",
            details={
                "folder_match_count": match["folder_match_count"],
                "workbook_match_count": match["workbook_match_count"],
                "result_match_count": match["result_match_count"],
                "index_state": match["index_state"],
            },
        )
    return {
        "candidates": match["candidates"],
        "count": len(match["candidates"]),
        "task": match["task_snapshot"],
        "matched_task_project": match["matched_task_project"],
        "task_validation_state": (
            "matched"
            if "task_item_name" in match["matched_conditions"]
            and "test_method" in match["matched_conditions"]
            else (
                "pending"
                if match["task_cache_state"] in {"pending", "empty"}
                else "warning"
            )
        ),
        "task_cache_state": match["task_cache_state"],
        "missing_conditions": missing,
    }


def register_paper_fiber_executors() -> None:
    node_registry.set_executor(
        PAPER_FIBER_NODE_TYPE,
        1,
        _paper_fiber_executor,
    )
