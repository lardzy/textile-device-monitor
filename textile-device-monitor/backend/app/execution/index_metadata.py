from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


METADATA_VERSION = 1
SUPPORTED_WORKBOOK_SUFFIXES = {
    ".xls",
    ".xlsx",
    ".xlsm",
    ".xlt",
    ".xltx",
    ".xltm",
}
INSPECTION_NUMBER_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])"
    r"((?:\d{2}[A-Za-z]\d{4,}(?:[-_]\d+)?)|(?:\d{6,}(?:[-_]\d+)?))"
    r"(?![A-Za-z0-9])",
    re.IGNORECASE,
)
CATEGORY_KEYWORDS = {
    "special_wool": (
        "特种毛",
        "动物纤维",
        "羊绒",
        "羊毛",
        "兔毛",
        "牦牛",
        "驼绒",
    ),
    "regenerated_fiber": (
        "再生纤维素",
        "粘胶",
        "粘纤",
        "莱赛尔",
        "莫代尔",
        "天丝",
    ),
    "hemp_cotton": (
        "麻棉",
        "苎麻",
        "亚麻",
        "大麻",
        "黄麻",
    ),
}


def inspection_numbers_in_text(value: str) -> list[str]:
    return list(dict.fromkeys(INSPECTION_NUMBER_PATTERN.findall(value or "")))


def _string_values(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    for value in values:
        if isinstance(value, str):
            text = value.strip()
            if text:
                result.append(text[:500])
    return result


def _read_modern_workbook(path: Path) -> tuple[list[str], list[str]]:
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=False,
        keep_links=False,
    )
    try:
        sheet_names = list(workbook.sheetnames)
        values: list[str] = []
        for sheet in workbook.worksheets[:3]:
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or 1, 40),
                max_col=min(sheet.max_column or 1, 20),
                values_only=True,
            ):
                values.extend(_string_values(row))
        return sheet_names, values
    finally:
        workbook.close()


def _read_legacy_workbook(path: Path) -> tuple[list[str], list[str]]:
    import xlrd

    workbook = xlrd.open_workbook(
        str(path),
        on_demand=True,
        formatting_info=False,
    )
    try:
        sheet_names = list(workbook.sheet_names())
        values: list[str] = []
        for sheet_name in sheet_names[:3]:
            sheet = workbook.sheet_by_name(sheet_name)
            for row_index in range(min(sheet.nrows, 40)):
                values.extend(
                    _string_values(
                        sheet.cell_value(row_index, column_index)
                        for column_index in range(min(sheet.ncols, 20))
                    )
                )
            workbook.unload_sheet(sheet_name)
        return sheet_names, values
    finally:
        workbook.release_resources()


def extract_index_metadata(
    path: Path,
    *,
    expected_category: str | None,
) -> dict[str, Any]:
    """Best-effort, bounded metadata extraction for the background indexer."""

    suffix = path.suffix.casefold()
    value: dict[str, Any] = {
        "metadata_version": METADATA_VERSION,
        "expected_category": expected_category,
        "format": suffix,
    }
    if suffix not in SUPPORTED_WORKBOOK_SUFFIXES:
        value["kind"] = "file"
        return value
    value["kind"] = "workbook"
    if path.stat().st_size > 30 * 1024 * 1024:
        value["parse_status"] = "skipped_too_large"
        return value
    try:
        if suffix in {".xls", ".xlt"}:
            sheet_names, cell_text = _read_legacy_workbook(path)
        else:
            sheet_names, cell_text = _read_modern_workbook(path)
    except Exception as exc:
        value["parse_status"] = "failed"
        value["parse_error"] = f"{type(exc).__name__}: {exc}"[:500]
        return value

    combined = "\n".join(cell_text)
    internal_numbers: list[str] = []
    for text in cell_text:
        internal_numbers.extend(inspection_numbers_in_text(text))
    detected_categories = [
        category
        for category, keywords in CATEGORY_KEYWORDS.items()
        if any(keyword in combined for keyword in keywords)
    ]
    if expected_category in detected_categories:
        type_status = "matched"
    elif detected_categories:
        type_status = "mismatch"
    else:
        type_status = "unknown"
    value.update(
        {
            "parse_status": "parsed",
            "sheet_names": sheet_names[:20],
            "internal_inspection_numbers": list(
                dict.fromkeys(internal_numbers)
            )[:20],
            "detected_categories": detected_categories,
            "type_status": type_status,
        }
    )
    return value
