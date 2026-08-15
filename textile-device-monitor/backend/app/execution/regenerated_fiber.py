from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.execution.errors import ExecutionApiError
from app.execution.index_metadata import inspection_numbers_in_text
from app.execution.models import (
    ExecutionFileIndexEntry,
    ExecutionIndexJob,
    ExecutionStorageRoot,
    ExecutionWorkflow,
    ExecutionWorkflowVersion,
    utcnow,
)
from app.execution.registry import node_registry
from app.execution.storage import ArtifactRef, FileGateway, StorageError
from app.execution.workbook_format import (
    SUPPORTED_WORKBOOK_SUFFIXES,
    WorkbookFormat,
    detect_workbook_format,
)


REGENERATED_FIBER_ROOT_ID = "regenerated_fiber_records"
NODE_TYPE_VERSION = 1
WORKBOOK_PROFILE_VERSION = 4
MAX_WORKBOOK_VALIDATION_MATCHES = 50
DETERMINISTIC_PROFILE_STATUSES = {
    "matched",
    "worksheet_missing",
    "content_range_empty",
    "unsupported_format",
}
@dataclass(frozen=True)
class RegeneratedFiberRule:
    node_type: str
    workflow_slug: str
    name: str
    worksheet: str
    cell_range: str = "B14:J14"
    root_id: str = REGENERATED_FIBER_ROOT_ID
    version: int = WORKBOOK_PROFILE_VERSION

    @property
    def cache_key(self) -> str:
        return f"{self.node_type}@{self.version}"


REGENERATED_FIBER_RULES = {
    "file.regenerated_fiber_count_method": RegeneratedFiberRule(
        node_type="file.regenerated_fiber_count_method",
        workflow_slug="regenerated-fiber-count-method",
        name="再生纤-根数法",
        worksheet="根数法报告1",
    ),
    "file.regenerated_fiber_area_method": RegeneratedFiberRule(
        node_type="file.regenerated_fiber_area_method",
        workflow_slug="regenerated-fiber-area-method",
        name="再生纤-面积法",
        worksheet="截面统计报告1",
    ),
}


def _matching_worksheet_names(
    sheet_names: list[str],
    rule: RegeneratedFiberRule,
) -> list[str]:
    return [name for name in sheet_names if name == rule.worksheet]


def _escaped_contains(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("%", "\\%")
        .replace("_", "\\_")
    )
    return f"%{escaped}%"


def _is_nonempty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    # Business values such as 0 and False are deliberately not treated as
    # blank.
    return True


def _read_profile(path: Path, rule: RegeneratedFiberRule) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    value: dict[str, Any] = {
        "rule_version": rule.version,
        "worksheet": rule.worksheet,
        "cell_range": rule.cell_range,
        "worksheet_exists": False,
        "content_range_nonempty": False,
    }
    if suffix not in SUPPORTED_WORKBOOK_SUFFIXES:
        value["status"] = "unsupported_format"
        return value

    try:
        workbook_format = detect_workbook_format(path)
        value["workbook_format"] = (
            workbook_format.value if workbook_format is not None else None
        )
        if workbook_format is None:
            value["status"] = "unsupported_format"
            return value
        if workbook_format is WorkbookFormat.OLE:
            try:
                import xlrd
            except ImportError:
                value["status"] = "reader_unavailable"
                return value
            workbook = xlrd.open_workbook(
                str(path),
                on_demand=True,
                formatting_info=False,
            )
            try:
                worksheet_names = _matching_worksheet_names(
                    workbook.sheet_names(),
                    rule,
                )
                if not worksheet_names:
                    value["status"] = "worksheet_missing"
                    return value
                value["worksheet_exists"] = True
                value["matched_worksheets"] = worksheet_names
                values = []
                for worksheet_name in worksheet_names:
                    sheet = workbook.sheet_by_name(worksheet_name)
                    values.extend(
                        sheet.cell_value(13, column)
                        if sheet.nrows > 13 and sheet.ncols > column
                        else None
                        for column in range(1, 10)
                    )
            finally:
                workbook.release_resources()
        else:
            with path.open("rb") as stream:
                workbook = load_workbook(
                    stream,
                    read_only=True,
                    data_only=True,
                    keep_vba=False,
                    keep_links=False,
                )
                try:
                    worksheet_names = _matching_worksheet_names(
                        list(workbook.sheetnames),
                        rule,
                    )
                    if not worksheet_names:
                        value["status"] = "worksheet_missing"
                        return value
                    value["worksheet_exists"] = True
                    value["matched_worksheets"] = worksheet_names
                    values = [
                        workbook[worksheet_name].cell(
                            row=14,
                            column=column,
                        ).value
                        for worksheet_name in worksheet_names
                        for column in range(2, 11)
                    ]
                finally:
                    workbook.close()
    except Exception as exc:
        # The cache deliberately stores only a bounded technical error. It
        # never records workbook cell content.
        value["status"] = "read_failed"
        value["error_type"] = type(exc).__name__
        return value

    value["content_range_nonempty"] = any(
        _is_nonempty(item) for item in values
    )
    value["status"] = (
        "matched"
        if value["content_range_nonempty"]
        else "content_range_empty"
    )
    return value


def _profile_for_entry(
    entry: ExecutionFileIndexEntry,
    *,
    rule: RegeneratedFiberRule,
    gateway: FileGateway,
) -> tuple[dict[str, Any], bool]:
    metadata = dict(entry.metadata_json or {})
    profiles = dict(metadata.get("workbook_profiles") or {})
    cached = profiles.get(rule.cache_key)
    if (
        isinstance(cached, dict)
        and cached.get("rule_version") == rule.version
        and cached.get("fingerprint") == entry.fingerprint
        and cached.get("status") in DETERMINISTIC_PROFILE_STATUSES
    ):
        return dict(cached), False

    original_fingerprint = entry.fingerprint
    try:
        path = gateway.resolve(
            ArtifactRef(rule.root_id, entry.relative_path),
            expected_type="file",
        )
        profile = _read_profile(path, rule)
    except (ExecutionApiError, StorageError, OSError) as exc:
        profile = {
            "rule_version": rule.version,
            "worksheet": rule.worksheet,
            "cell_range": rule.cell_range,
            "worksheet_exists": False,
            "content_range_nonempty": False,
            "status": "file_unavailable",
            "error_type": type(exc).__name__,
        }
    profile.update(
        {
            "rule_version": rule.version,
            "fingerprint": entry.fingerprint,
            "checked_at": utcnow().isoformat(),
        }
    )
    # No database row lock is held while opening the workbook over SMB.
    return profile, _store_profile_after_read(
        entry,
        rule=rule,
        profile=profile,
        original_fingerprint=original_fingerprint,
    )


def _store_profile_after_read(
    entry: ExecutionFileIndexEntry,
    *,
    rule: RegeneratedFiberRule,
    profile: dict[str, Any],
    original_fingerprint: str,
) -> bool:
    metadata = dict(entry.metadata_json or {})
    profiles = dict(metadata.get("workbook_profiles") or {})
    changed = False
    if (
        entry.fingerprint == original_fingerprint
        and profile.get("status") in DETERMINISTIC_PROFILE_STATUSES
    ):
        if profiles.get(rule.cache_key) != profile:
            profiles[rule.cache_key] = profile
            changed = True
    elif rule.cache_key in profiles:
        # Old versions cached transient read failures forever. Remove such a
        # record so the next complete-number query can retry.
        profiles.pop(rule.cache_key, None)
        changed = True
    if changed:
        metadata["workbook_profiles"] = profiles
        entry.metadata_json = metadata
    return changed


def _is_complete_inspection_number(value: str) -> bool:
    matches = inspection_numbers_in_text(value)
    return (
        len(matches) == 1
        and matches[0].casefold() == value.casefold()
    )


def _index_state(
    db: Session,
    *,
    root: Optional[ExecutionStorageRoot],
    entry_count: int,
) -> str:
    if root is None or not root.is_active or not root.is_available:
        return "unavailable"
    active_job = (
        db.query(ExecutionIndexJob)
        .filter(
            ExecutionIndexJob.storage_root_id == root.id,
            ExecutionIndexJob.status.in_(("queued", "running")),
        )
        .order_by(ExecutionIndexJob.created_at.desc())
        .first()
    )
    if active_job is not None:
        return "indexing"
    latest_job = (
        db.query(ExecutionIndexJob)
        .filter(ExecutionIndexJob.storage_root_id == root.id)
        .order_by(ExecutionIndexJob.created_at.desc())
        .first()
    )
    if latest_job is not None and latest_job.status == "failed":
        return "failed"
    if (
        root.last_scan_error
        or (
            latest_job is not None
            and latest_job.status == "completed_with_errors"
        )
    ):
        return "degraded"
    if root.last_scan_finished_at is None and entry_count == 0:
        return "pending"
    return "ready"


def _eligible_depth(relative_path: str) -> bool:
    # A record may live directly below the root or in exactly one child
    # directory. Deeper archival trees are intentionally ignored.
    return len(Path(relative_path).parts) <= 2


def match_regenerated_fiber_workbooks(
    db: Session,
    *,
    node_type: str,
    inspection_number: str,
    gateway: Optional[FileGateway] = None,
    result_limit: int = 6,
) -> dict[str, Any]:
    rule = REGENERATED_FIBER_RULES.get(node_type)
    if rule is None:
        raise ExecutionApiError(
            422,
            "regenerated_fiber_rule_unknown",
            "未知的再生纤文件识别规则",
        )

    query_text = str(inspection_number or "").strip()
    root = (
        db.query(ExecutionStorageRoot)
        .filter(ExecutionStorageRoot.root_id == rule.root_id)
        .one_or_none()
    )
    if root is None or not root.is_active or not root.is_available:
        return {
            "root_id": rule.root_id,
            "index_state": "unavailable",
            "filename_match_count": 0,
            "worksheet_match_count": 0,
            "full_match_count": 0,
            "candidates": [],
            "best_file_conditions": [],
            "cache_updated": False,
            "query_state": (
                "empty"
                if not query_text
                else (
                    "complete"
                    if _is_complete_inspection_number(query_text)
                    else "incomplete"
                )
            ),
        }

    base_entry_count = (
        db.query(ExecutionFileIndexEntry.id)
        .filter(
            ExecutionFileIndexEntry.storage_root_id == root.id,
            ExecutionFileIndexEntry.missing_since.is_(None),
        )
        .count()
    )
    if not query_text:
        return {
            "root_id": rule.root_id,
            "index_state": _index_state(
                db,
                root=root,
                entry_count=base_entry_count,
            ),
            "filename_match_count": 0,
            "worksheet_match_count": 0,
            "full_match_count": 0,
            "candidates": [],
            "best_file_conditions": [],
            "cache_updated": False,
            "query_state": "empty",
        }

    entries = (
        db.query(ExecutionFileIndexEntry)
        .filter(
            ExecutionFileIndexEntry.storage_root_id == root.id,
            ExecutionFileIndexEntry.missing_since.is_(None),
            ExecutionFileIndexEntry.filename.ilike(
                _escaped_contains(query_text),
                escape="\\",
            ),
        )
        .order_by(
            ExecutionFileIndexEntry.modified_at.desc(),
            ExecutionFileIndexEntry.relative_path.asc(),
        )
        .all()
    )
    entries = [
        entry for entry in entries if _eligible_depth(entry.relative_path)
    ]
    if not entries:
        return {
            "root_id": rule.root_id,
            "index_state": _index_state(
                db,
                root=root,
                entry_count=base_entry_count,
            ),
            "filename_match_count": 0,
            "worksheet_match_count": 0,
            "full_match_count": 0,
            "candidates": [],
            "best_file_conditions": [],
            "cache_updated": False,
            "query_state": (
                "complete"
                if _is_complete_inspection_number(query_text)
                else "incomplete"
            ),
        }

    query_is_complete = _is_complete_inspection_number(query_text)
    if not query_is_complete or len(entries) > MAX_WORKBOOK_VALIDATION_MATCHES:
        return {
            "root_id": rule.root_id,
            "index_state": _index_state(
                db,
                root=root,
                entry_count=base_entry_count,
            ),
            "filename_match_count": len(entries),
            "worksheet_match_count": 0,
            "full_match_count": 0,
            "candidates": [],
            "best_file_conditions": ["filename"],
            "cache_updated": False,
            "query_state": (
                "incomplete"
                if not query_is_complete
                else "too_many_matches"
            ),
        }

    if gateway is None:
        from app.execution.persistence import build_file_gateway

        gateway = build_file_gateway(db)

    worksheet_match_count = 0
    matched: list[dict[str, Any]] = []
    best_file_conditions: list[str] = ["filename"]
    cache_updated = False
    for entry in entries:
        profile, changed = _profile_for_entry(
            entry,
            rule=rule,
            gateway=gateway,
        )
        cache_updated = cache_updated or changed
        conditions = ["filename"]
        if profile.get("worksheet_exists") is True:
            worksheet_match_count += 1
            conditions.append("worksheet")
        if profile.get("content_range_nonempty") is True:
            conditions.append("content_range")
        if len(conditions) > len(best_file_conditions):
            best_file_conditions = conditions
        if profile.get("status") != "matched":
            continue
        matched.append(
            {
                "id": entry.id,
                "root_id": rule.root_id,
                "relative_path": entry.relative_path,
                "name": entry.filename,
                "suffix": entry.extension,
                "size": entry.size_bytes,
                "modified_at": entry.modified_at.isoformat(),
                "category": root.category_key,
                "fingerprint": entry.fingerprint,
            }
        )

    candidates = matched[: max(1, min(int(result_limit), 6))]
    return {
        "root_id": rule.root_id,
        "index_state": _index_state(
            db,
            root=root,
            entry_count=base_entry_count,
        ),
        "filename_match_count": len(entries),
        "worksheet_match_count": worksheet_match_count,
        "full_match_count": len(matched),
        "candidates": candidates,
        "best_file_conditions": best_file_conditions,
        "cache_updated": cache_updated,
        "query_state": "validated",
    }


def _published_version(
    workflow: ExecutionWorkflow,
) -> Optional[ExecutionWorkflowVersion]:
    if workflow.published_version_number is None:
        return None
    return next(
        (
            version
            for version in workflow.versions
            if version.version_number == workflow.published_version_number
        ),
        None,
    )


def _source_root_ids(definition: dict[str, Any]) -> list[str]:
    return [
        str(slot["root_id"])
        for slot in definition.get("root_slots") or []
        if isinstance(slot, dict)
        and slot.get("access", "read") == "read"
        and slot.get("root_id")
    ]


def _specialized_node_type(
    definition: dict[str, Any],
) -> Optional[str]:
    for node in definition.get("nodes") or []:
        if (
            isinstance(node, dict)
            and node.get("disabled") is not True
            and node.get("type")
            in {
                *REGENERATED_FIBER_RULES.keys(),
                "file.electron_microscopy_gbt36422",
                "file.paper_fiber_gbt4688_qualitative",
            }
        ):
            return str(node["type"])
    return None


def _specialized_record_family(
    definition: dict[str, Any],
    node_type: str,
) -> Optional[str]:
    """Return the discover node's pinned ``record_family``, if any."""

    for node in definition.get("nodes") or []:
        if (
            isinstance(node, dict)
            and node.get("disabled") is not True
            and node.get("type") == node_type
        ):
            value = (node.get("config") or {}).get("record_family")
            return str(value).strip() if value else None
    return None


def _safe_candidate_preview(
    *,
    name: object,
    relative_path: object,
    suffix: object,
) -> dict[str, str]:
    """Return catalog-safe file metadata without ever exposing a root path."""

    raw_name = str(name or "").strip()
    safe_name = PurePosixPath(
        PureWindowsPath(raw_name).name
    ).name
    raw_relative_path = str(relative_path or "").strip()
    normalized_relative_path = raw_relative_path.replace("\\", "/")
    path_parts = PurePosixPath(normalized_relative_path).parts
    if (
        not normalized_relative_path
        or normalized_relative_path.startswith("/")
        or PureWindowsPath(raw_relative_path).is_absolute()
        or ".." in path_parts
    ):
        normalized_relative_path = safe_name
    else:
        normalized_relative_path = "/".join(
            part for part in path_parts if part not in {"", "."}
        )
    raw_suffix = str(suffix or "").strip().casefold()
    safe_suffix = (
        raw_suffix
        if raw_suffix.startswith(".")
        else (f".{raw_suffix}" if raw_suffix else "")
    )
    return {
        "name": safe_name,
        "relative_path": normalized_relative_path or safe_name,
        "suffix": safe_suffix,
    }


def _generic_filename_match(
    db: Session,
    *,
    root_ids: list[str],
    inspection_number: str,
) -> tuple[int, Optional[dict[str, str]]]:
    if not root_ids or not inspection_number:
        return 0, None
    query = (
        db.query(ExecutionFileIndexEntry)
        .join(
            ExecutionStorageRoot,
            ExecutionStorageRoot.id
            == ExecutionFileIndexEntry.storage_root_id,
        )
        .filter(
            ExecutionStorageRoot.root_id.in_(root_ids),
            ExecutionStorageRoot.is_active.is_(True),
            ExecutionFileIndexEntry.missing_since.is_(None),
            ExecutionFileIndexEntry.filename.ilike(
                _escaped_contains(inspection_number),
                escape="\\",
            ),
        )
        .order_by(
            ExecutionFileIndexEntry.modified_at.desc(),
            ExecutionFileIndexEntry.relative_path.asc(),
        )
    )
    count = query.order_by(None).count()
    entry = query.first()
    if entry is None:
        return count, None
    return count, _safe_candidate_preview(
        name=entry.filename,
        relative_path=entry.relative_path,
        suffix=entry.extension,
    )


def catalog_recommendations(
    db: Session,
    *,
    inspection_number: str,
    preferred_categories: list[str],
    include_hidden: bool,
) -> tuple[list[dict[str, Any]], bool]:
    inspection_number = str(inspection_number or "").strip()
    preferred = {
        item.strip()
        for item in preferred_categories
        if isinstance(item, str) and item.strip()
    }
    workflows = (
        db.query(ExecutionWorkflow)
        .order_by(
            ExecutionWorkflow.category_id.asc(),
            ExecutionWorkflow.updated_at.desc(),
            ExecutionWorkflow.id.asc(),
        )
        .all()
    )
    roots = {
        root.root_id: root
        for root in db.query(ExecutionStorageRoot).all()
    }
    category_order = {
        workflow.category.key: workflow.category.sort_order
        for workflow in workflows
    }
    original_order = {
        workflow.id: index for index, workflow in enumerate(workflows)
    }
    match_cache: dict[str, dict[str, Any]] = {}
    items: list[dict[str, Any]] = []
    any_cache_updated = False

    for workflow in workflows:
        published = _published_version(workflow)
        published_capabilities = (
            dict(published.capabilities or {}) if published else {}
        )
        if published_capabilities.get("system_deprecated"):
            continue
        if published_capabilities.get("hidden") and not include_hidden:
            continue
        definition = dict(published.definition or {}) if published else {}
        root_ids = _source_root_ids(definition)
        roots_ready = all(
            roots.get(root_id) is not None
            and roots[root_id].is_active
            and roots[root_id].is_available
            for root_id in root_ids
        )
        blocking_availability = workflow.availability_code not in {
            None,
            "root_not_configured",
        }
        runnable = bool(
            workflow.is_enabled
            and published is not None
            and roots_ready
            and not blocking_availability
        )

        conditions: list[str] = []
        score = 0
        if workflow.category.key in preferred:
            conditions.append("category")
            score += 1
        if root_ids and roots_ready:
            conditions.append("source_root")
            score += 1

        node_type = _specialized_node_type(definition)
        candidate_count = 0
        candidate_preview: Optional[dict[str, Any]] = None
        index_state = "ready"
        diagnostics: dict[str, int] = {
            "filename_match_count": 0,
            "worksheet_match_count": 0,
            "full_match_count": 0,
        }
        query_state = (
            "empty"
            if not inspection_number
            else (
                "complete"
                if _is_complete_inspection_number(inspection_number)
                else "incomplete"
            )
        )
        task_cache_state: Optional[str] = None
        if node_type == "file.paper_fiber_gbt4688_qualitative":
            from app.execution.paper_fiber import paper_fiber_match

            if node_type not in match_cache:
                match_cache[node_type] = paper_fiber_match(
                    db,
                    inspection_number=inspection_number,
                )
            match = match_cache[node_type]
            any_cache_updated = any_cache_updated or bool(
                match["cache_updated"]
            )
            index_state = str(match["index_state"])
            query_state = str(match["query_state"])
            for condition in match["matched_conditions"]:
                if condition not in conditions:
                    conditions.append(condition)
                    score += 1
            candidate_count = int(match["result_match_count"])
            candidate_preview = match["candidate_preview"]
            diagnostics = {
                "filename_match_count": int(match["folder_match_count"]),
                "worksheet_match_count": int(match["result_match_count"]),
                "full_match_count": 1 if match["full_match"] else 0,
            }
            full_match = bool(match["full_match"])
            task_cache_state = str(match["task_cache_state"])
        elif node_type == "file.electron_microscopy_gbt36422":
            from app.execution.electron_microscopy import (
                electron_microscopy_match,
            )
            from app.execution.microscopy_families import (
                microscopy_family_for_key,
            )

            family = microscopy_family_for_key(
                _specialized_record_family(definition, node_type)
            )
            cache_key = (node_type, family.key if family else None)
            if cache_key not in match_cache:
                match_cache[cache_key] = electron_microscopy_match(
                    db,
                    inspection_number=inspection_number,
                    family=family,
                )
            match = match_cache[cache_key]
            any_cache_updated = any_cache_updated or bool(
                match["cache_updated"]
            )
            index_state = str(match["index_state"])
            query_state = (
                "complete"
                if _is_complete_inspection_number(inspection_number)
                else ("incomplete" if inspection_number else "empty")
            )
            for condition in match["matched_conditions"]:
                if condition not in conditions:
                    conditions.append(condition)
                    score += 1
            candidate_count = int(match["image_count"])
            if match["folders"]:
                first_folder = match["folders"][0]
                candidate_preview = {
                    "name": str(first_folder["name"]),
                    "relative_path": str(first_folder["relative_path"]),
                    "suffix": "",
                }
            diagnostics = {
                "filename_match_count": int(match["folder_match_count"]),
                "worksheet_match_count": 0,
                "full_match_count": 1 if match["full_match"] else 0,
            }
            full_match = bool(match["full_match"])
            task_cache_state = str(match["task_cache_state"])
        elif node_type is not None:
            if node_type not in match_cache:
                match_cache[node_type] = match_regenerated_fiber_workbooks(
                    db,
                    node_type=node_type,
                    inspection_number=inspection_number,
                )
            match = match_cache[node_type]
            any_cache_updated = (
                any_cache_updated or bool(match["cache_updated"])
            )
            index_state = str(match["index_state"])
            query_state = str(match["query_state"])
            for condition in match["best_file_conditions"]:
                if condition not in conditions:
                    conditions.append(condition)
                    score += 1
            candidate_count = len(match["candidates"])
            if match["candidates"]:
                candidate = match["candidates"][0]
                candidate_preview = _safe_candidate_preview(
                    name=candidate.get("name"),
                    relative_path=candidate.get("relative_path"),
                    suffix=candidate.get("suffix"),
                )
            diagnostics = {
                key: int(match[key])
                for key in diagnostics
            }
            full_match = bool(
                roots_ready
                and inspection_number
                and match["full_match_count"] > 0
            )
        else:
            filename_count, candidate_preview = _generic_filename_match(
                db,
                root_ids=root_ids,
                inspection_number=inspection_number,
            )
            if filename_count:
                conditions.append("filename")
                score += 1
            candidate_count = filename_count
            diagnostics["filename_match_count"] = filename_count
            full_match = False
            root_states = [
                _index_state(
                    db,
                    root=roots.get(root_id),
                    entry_count=(
                        db.query(ExecutionFileIndexEntry.id)
                        .filter(
                            ExecutionFileIndexEntry.storage_root_id
                            == roots[root_id].id,
                            ExecutionFileIndexEntry.missing_since.is_(None),
                        )
                        .count()
                        if roots.get(root_id) is not None
                        else 0
                    ),
                )
                for root_id in root_ids
            ]
            if "unavailable" in root_states:
                index_state = "unavailable"
            elif "indexing" in root_states:
                index_state = "indexing"
            elif "pending" in root_states:
                index_state = "pending"

        if index_state == "unavailable":
            state = "index_unavailable"
        elif index_state == "failed":
            state = "index_failed"
        elif index_state == "degraded" and not full_match:
            state = "index_degraded"
        elif index_state in {"pending", "indexing"} and not full_match:
            state = "index_pending"
        elif full_match:
            state = "full_match"
        elif score:
            state = "partial_match"
        else:
            state = "no_match"

        items.append(
            {
                "workflow_id": workflow.id,
                "state": state,
                "score": score,
                "max_score": 5,
                "full_match": full_match,
                "candidate_count": candidate_count,
                "candidate_preview": candidate_preview,
                "matched_conditions": conditions,
                "index_state": index_state,
                "query_state": query_state,
                "runnable": runnable,
                "diagnostics": diagnostics,
                "task_cache_state": task_cache_state,
                "_category_order": category_order.get(
                    workflow.category.key,
                    0,
                ),
                "_original_order": original_order[workflow.id],
            }
        )

    items.sort(
        key=lambda item: (
            not item["runnable"],
            -item["score"],
            not item["full_match"],
            item["_category_order"],
            item["_original_order"],
            item["workflow_id"],
        )
    )
    for rank, item in enumerate(items, start=1):
        item.pop("_category_order", None)
        item.pop("_original_order", None)
        item["rank"] = rank
    return items, any_cache_updated


def _regenerated_fiber_executor(context) -> dict[str, Any]:
    node_type = context.node_run.node_type
    config = context.node.get("config") or {}
    result = match_regenerated_fiber_workbooks(
        context.db,
        node_type=node_type,
        inspection_number=str(
            context.input_data.get("inspection_number")
            or context.run.inspection_number
        ),
        result_limit=min(int(config.get("limit", 6)), 6),
    )
    if not result["candidates"]:
        details = {
            "filename_match_count": result["filename_match_count"],
            "worksheet_match_count": result["worksheet_match_count"],
            "full_match_count": result["full_match_count"],
            "index_state": result["index_state"],
            "query_state": result["query_state"],
        }
        raise ExecutionApiError(
            422,
            "matching_workbook_not_found",
            (
                "未找到符合规则的再生纤工作簿"
                f"（文件名命中 {details['filename_match_count']}，"
                f"工作表命中 {details['worksheet_match_count']}，"
                f"完整命中 {details['full_match_count']}）"
            ),
            details=details,
        )
    return {
        "candidates": result["candidates"],
        "count": len(result["candidates"]),
        "diagnostics": {
            "filename_match_count": result["filename_match_count"],
            "worksheet_match_count": result["worksheet_match_count"],
            "full_match_count": result["full_match_count"],
        },
        "query_state": result["query_state"],
    }


_EXECUTORS_REGISTERED = False


def register_regenerated_fiber_executors() -> None:
    global _EXECUTORS_REGISTERED
    if _EXECUTORS_REGISTERED:
        return
    for node_type in REGENERATED_FIBER_RULES:
        node_registry.set_executor(
            node_type,
            NODE_TYPE_VERSION,
            _regenerated_fiber_executor,
        )
    from app.execution.regenerated_fiber_results import (
        register_regenerated_fiber_result_executors,
    )

    register_regenerated_fiber_result_executors()
    _EXECUTORS_REGISTERED = True
