from __future__ import annotations

import errno
import hashlib
import json
import math
import os
import re
import shutil
import threading
import uuid
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Iterator, Sequence

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from app.execution.storage import (
    ArtifactFingerprint,
    ArtifactRef,
    FileGateway,
    RootPermissionError,
    fingerprint_file,
)


_MUTATION_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$")
_MANIFEST_SCHEMA_VERSION = 1


class MutationError(RuntimeError):
    pass


class InvalidMutationIdError(MutationError):
    pass


class MutationConflictError(MutationError):
    pass


class UnsupportedWorkbookError(MutationError):
    pass


class WorkbookVerificationError(MutationError):
    pass


class PublishConflictError(MutationConflictError):
    pass


@dataclass(frozen=True, slots=True)
class CellWrite:
    sheet: str
    cell: str
    value: object

    def normalized(self) -> "CellWrite":
        sheet = self.sheet.strip()
        if not sheet:
            raise ValueError("sheet_name_required")
        if len(sheet) > 31:
            raise ValueError("sheet_name_too_long")

        raw_cell = self.cell.strip().replace("$", "").upper()
        try:
            row, column = coordinate_to_tuple(raw_cell)
        except (TypeError, ValueError) as exc:
            raise ValueError("invalid_cell_coordinate") from exc
        canonical_cell = f"{get_column_letter(column)}{row}"
        value = _coerce_cell_value(self.value)
        _encode_cell_value(value)
        return CellWrite(sheet=sheet, cell=canonical_cell, value=value)

    def as_dict(self) -> dict[str, object]:
        value = _coerce_cell_value(self.value)
        return {
            "sheet": self.sheet,
            "cell": self.cell,
            "value": _encode_cell_value(value),
        }


@dataclass(frozen=True, slots=True)
class WorkingCopyReceipt:
    mutation_id: str
    source_ref: ArtifactRef
    source_fingerprint: ArtifactFingerprint
    working_ref: ArtifactRef
    working_fingerprint: ArtifactFingerprint
    created_at: datetime
    reused: bool = False

    def as_dict(self) -> dict[str, object]:
        return {
            "mutation_id": self.mutation_id,
            "source": {
                **self.source_ref.as_dict(),
                "fingerprint": self.source_fingerprint.as_dict(),
            },
            "working_copy": {
                **self.working_ref.as_dict(),
                "fingerprint": self.working_fingerprint.as_dict(),
            },
            "created_at": self.created_at.isoformat(),
            "reused": self.reused,
        }


@dataclass(frozen=True, slots=True)
class WorkbookVerificationReceipt:
    workbook_ref: ArtifactRef
    verified: bool
    checked_cells: int
    mismatches: tuple[dict[str, object], ...]

    def as_dict(self) -> dict[str, object]:
        return {
            "workbook": self.workbook_ref.as_dict(),
            "verified": self.verified,
            "checked_cells": self.checked_cells,
            "mismatches": list(self.mismatches),
        }


@dataclass(frozen=True, slots=True)
class WorkbookWriteReceipt:
    mutation_id: str
    working_ref: ArtifactRef
    before_fingerprint: ArtifactFingerprint
    after_fingerprint: ArtifactFingerprint
    writes: tuple[CellWrite, ...]
    verification: WorkbookVerificationReceipt
    written_at: datetime
    reused: bool = False

    def as_dict(self) -> dict[str, object]:
        return {
            "mutation_id": self.mutation_id,
            "working_copy": {
                **self.working_ref.as_dict(),
                "before_fingerprint": self.before_fingerprint.as_dict(),
                "after_fingerprint": self.after_fingerprint.as_dict(),
            },
            "writes": [item.as_dict() for item in self.writes],
            "verification": self.verification.as_dict(),
            "written_at": self.written_at.isoformat(),
            "reused": self.reused,
        }


@dataclass(frozen=True, slots=True)
class PublishReceipt:
    mutation_id: str
    source_ref: ArtifactRef
    source_fingerprint: ArtifactFingerprint
    published_ref: ArtifactRef
    published_fingerprint: ArtifactFingerprint
    published_at: datetime
    reused: bool = False

    def as_dict(self) -> dict[str, object]:
        return {
            "mutation_id": self.mutation_id,
            "source": {
                **self.source_ref.as_dict(),
                "fingerprint": self.source_fingerprint.as_dict(),
            },
            "published": {
                **self.published_ref.as_dict(),
                "fingerprint": self.published_fingerprint.as_dict(),
            },
            "published_at": self.published_at.isoformat(),
            "reused": self.reused,
        }


def validate_mutation_id(mutation_id: str) -> str:
    if not isinstance(mutation_id, str):
        raise InvalidMutationIdError("mutation_id_must_be_string")
    normalized = mutation_id.strip()
    if not _MUTATION_ID_PATTERN.fullmatch(normalized):
        raise InvalidMutationIdError("invalid_mutation_id")
    return normalized


def _encode_cell_value(value: object) -> dict[str, object]:
    if value is None:
        return {"type": "null", "value": None}
    if isinstance(value, bool):
        return {"type": "bool", "value": value}
    if isinstance(value, int):
        return {"type": "int", "value": value}
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non_finite_cell_value")
        return {"type": "float", "value": value}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise ValueError("unsupported_cell_value")


def _decode_cell_value(payload: dict[str, object]) -> object:
    value_type = str(payload.get("type") or "")
    value = payload.get("value")
    if value_type == "null":
        return None
    if value_type == "bool":
        return bool(value)
    if value_type == "int":
        return int(value)
    if value_type == "float":
        return float(value)
    if value_type == "datetime":
        return datetime.fromisoformat(str(value))
    if value_type == "date":
        return date.fromisoformat(str(value))
    if value_type == "time":
        return time.fromisoformat(str(value))
    if value_type == "string":
        return str(value)
    raise MutationError("invalid_manifest_cell_value")


def _coerce_cell_value(value: object) -> object:
    if isinstance(value, dict) and set(value) == {"type", "value"}:
        return _decode_cell_value(value)
    return value


def _fingerprint_from_dict(payload: dict[str, object]) -> ArtifactFingerprint:
    return ArtifactFingerprint(
        size=int(payload["size"]),
        modified_ns=int(payload["modified_ns"]),
        sha256=str(payload["sha256"]),
    )


def _ref_from_dict(payload: dict[str, object]) -> ArtifactRef:
    return ArtifactRef(
        root_id=str(payload["root_id"]),
        relative_path=str(payload["relative_path"]),
    )


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _parse_utc(value: str) -> datetime:
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _normalize_writes(writes: Sequence[CellWrite]) -> tuple[CellWrite, ...]:
    normalized = tuple(item.normalized() for item in writes)
    if not normalized:
        raise ValueError("at_least_one_cell_write_required")

    seen: set[tuple[str, str]] = set()
    for item in normalized:
        key = (item.sheet.casefold(), item.cell)
        if key in seen:
            raise ValueError("duplicate_cell_write")
        seen.add(key)
    return normalized


def _writes_digest(writes: Sequence[CellWrite]) -> str:
    payload = json.dumps(
        [item.as_dict() for item in writes],
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _values_equal(expected: object, actual: object) -> bool:
    if isinstance(expected, datetime) and isinstance(actual, datetime):
        if expected.tzinfo is None or actual.tzinfo is None:
            return expected.replace(tzinfo=None) == actual.replace(tzinfo=None)
        return expected.astimezone(timezone.utc) == actual.astimezone(timezone.utc)
    return expected == actual


def _fingerprints_match_content(
    left: ArtifactFingerprint,
    right: ArtifactFingerprint,
) -> bool:
    """Compare immutable content while tolerating SMB timestamp rounding."""

    return left.size == right.size and left.sha256 == right.sha256


class FileMutationService:
    """Copy-on-write workbook mutations with a filesystem idempotency ledger."""

    def __init__(self, gateway: FileGateway, *, staging_root_id: str):
        staging_root = gateway.get_root(staging_root_id)
        if not staging_root.writable:
            raise RootPermissionError("staging_root_is_read_only")
        self._gateway = gateway
        self._staging_root_id = staging_root_id
        self._thread_locks: dict[str, threading.RLock] = {}
        self._thread_locks_guard = threading.Lock()

    def prepare_working_copy(
        self,
        mutation_id: str,
        source_ref: ArtifactRef,
        *,
        working_relative_path: str | None = None,
    ) -> WorkingCopyReceipt:
        normalized_id = validate_mutation_id(mutation_id)
        source_path = self._gateway.resolve(
            source_ref,
            expected_type="file",
        )
        source_fingerprint = fingerprint_file(source_path)
        working_ref = ArtifactRef(
            self._staging_root_id,
            working_relative_path
            or f".execution-mutations/{normalized_id}/working/{source_path.name}",
        )

        with self._mutation_lock(normalized_id):
            manifest = self._read_manifest(normalized_id)
            if manifest is not None:
                receipt = self._working_receipt_from_manifest(manifest)
                if receipt.source_ref != source_ref:
                    raise MutationConflictError("mutation_source_conflict")
                if receipt.working_ref != working_ref:
                    raise MutationConflictError("mutation_working_path_conflict")
                if receipt.source_fingerprint != source_fingerprint:
                    raise MutationConflictError("source_changed_since_prepare")

                working_path = self._gateway.resolve(
                    receipt.working_ref,
                    expected_type="file",
                )
                current_fingerprint = fingerprint_file(working_path)
                expected_fingerprint = receipt.working_fingerprint
                write_state = manifest.get("write")
                if write_state:
                    write_payload = dict(write_state)
                    pending_before = _fingerprint_from_dict(
                        dict(write_payload["before_fingerprint"])
                    )
                    pending_after = _fingerprint_from_dict(
                        dict(write_payload["after_fingerprint"])
                    )
                    if str(write_payload.get("status") or "complete") == "pending":
                        if _fingerprints_match_content(
                            current_fingerprint,
                            pending_before,
                        ):
                            expected_fingerprint = pending_before
                        else:
                            expected_fingerprint = pending_after
                    else:
                        expected_fingerprint = pending_after
                if not _fingerprints_match_content(
                    current_fingerprint,
                    expected_fingerprint,
                ):
                    raise MutationConflictError("working_copy_changed")
                return replace(receipt, reused=True)

            working_path = self._gateway.resolve(
                working_ref,
                must_exist=False,
                for_write=True,
            )
            if working_path.exists():
                # Crash recovery for the copy -> manifest window. Adopting an
                # existing file is safe only when its complete content equals
                # the still-unchanged source requested by this mutation.
                copied_fingerprint = fingerprint_file(working_path)
                current_source_fingerprint = fingerprint_file(source_path)
                if (
                    current_source_fingerprint != source_fingerprint
                    or not _fingerprints_match_content(
                        copied_fingerprint,
                        source_fingerprint,
                    )
                ):
                    raise MutationConflictError("working_copy_target_exists")
                created_at = _utc_now()
                manifest = self._new_manifest(
                    normalized_id,
                    source_ref,
                    source_fingerprint,
                    working_ref,
                    copied_fingerprint,
                    created_at,
                )
                self._write_manifest(normalized_id, manifest)
                return WorkingCopyReceipt(
                    mutation_id=normalized_id,
                    source_ref=source_ref,
                    source_fingerprint=source_fingerprint,
                    working_ref=working_ref,
                    working_fingerprint=copied_fingerprint,
                    created_at=created_at,
                    reused=True,
                )
            self._gateway.ensure_parent(working_ref)
            self._copy_exclusive(
                source_path,
                working_path,
                mutation_id=normalized_id,
                expected_source=source_fingerprint,
            )

            copied_fingerprint = fingerprint_file(working_path)
            current_source_fingerprint = fingerprint_file(source_path)
            if (
                current_source_fingerprint != source_fingerprint
                or copied_fingerprint.sha256 != source_fingerprint.sha256
                or copied_fingerprint.size != source_fingerprint.size
            ):
                working_path.unlink(missing_ok=True)
                raise MutationConflictError("source_changed_during_copy")

            created_at = _utc_now()
            manifest = self._new_manifest(
                normalized_id,
                source_ref,
                source_fingerprint,
                working_ref,
                copied_fingerprint,
                created_at,
            )
            self._write_manifest(normalized_id, manifest)
            return WorkingCopyReceipt(
                mutation_id=normalized_id,
                source_ref=source_ref,
                source_fingerprint=source_fingerprint,
                working_ref=working_ref,
                working_fingerprint=copied_fingerprint,
                created_at=created_at,
            )

    def write_xlsx_cells(
        self,
        mutation_id: str,
        working_ref: ArtifactRef,
        writes: Sequence[CellWrite],
    ) -> WorkbookWriteReceipt:
        normalized_id = validate_mutation_id(mutation_id)
        normalized_writes = _normalize_writes(writes)
        operation_digest = _writes_digest(normalized_writes)

        if working_ref.root_id != self._staging_root_id:
            raise RootPermissionError("workbook_must_be_in_staging_root")
        if Path(working_ref.relative_path).suffix.casefold() != ".xlsx":
            raise UnsupportedWorkbookError("xlsx_desktop_adapter_required")

        with self._mutation_lock(normalized_id):
            manifest = self._require_manifest(normalized_id)
            prepared = self._working_receipt_from_manifest(manifest)
            if prepared.working_ref != working_ref:
                raise MutationConflictError("mutation_working_path_conflict")
            self._assert_source_unchanged(prepared)

            existing_write = manifest.get("write")
            working_path = self._gateway.resolve(
                working_ref,
                for_write=True,
                expected_type="file",
            )
            current = fingerprint_file(working_path)
            if existing_write:
                write_payload = dict(existing_write)
                if str(write_payload["operation_digest"]) != operation_digest:
                    raise MutationConflictError("mutation_write_conflict")
                before = _fingerprint_from_dict(
                    dict(write_payload["before_fingerprint"])
                )
                expected_after = _fingerprint_from_dict(
                    dict(write_payload["after_fingerprint"])
                )
                write_status = str(write_payload.get("status") or "complete")

                if _fingerprints_match_content(current, expected_after):
                    verification = self.verify_xlsx_cells(
                        working_ref,
                        normalized_writes,
                    )
                    if not verification.verified:
                        raise WorkbookVerificationError(
                            "workbook_verification_failed"
                        )
                    if write_status == "pending":
                        write_payload["status"] = "complete"
                        write_payload["after_fingerprint"] = current.as_dict()
                        write_payload["reconciled_at"] = _utc_now().isoformat()
                        manifest["write"] = write_payload
                        self._write_manifest(normalized_id, manifest)
                    receipt = self._write_receipt_from_manifest(manifest)
                    return replace(
                        receipt,
                        verification=verification,
                        reused=True,
                    )

                if write_status != "pending":
                    raise MutationConflictError("working_copy_changed_after_write")
                if not _fingerprints_match_content(current, before):
                    raise MutationConflictError("working_copy_changed_during_write")
                # The intent was durable but replacement did not happen. Rebuild
                # the candidate and safely resume from the verified before state.
                before = current
            else:
                before = current
                if before != prepared.working_fingerprint:
                    raise MutationConflictError("working_copy_changed_before_write")

            workbook = load_workbook(
                filename=working_path,
                read_only=False,
                data_only=False,
                keep_links=True,
            )
            temp_path = working_path.with_name(
                f".{working_path.stem}.{uuid.uuid4().hex}.xlsx"
            )
            try:
                for item in normalized_writes:
                    if item.sheet not in workbook.sheetnames:
                        raise MutationError(f"worksheet_not_found:{item.sheet}")
                    workbook[item.sheet][item.cell] = item.value
                workbook.save(temp_path)
            finally:
                workbook.close()

            try:
                temp_verification = self._verify_workbook_path(
                    temp_path,
                    working_ref,
                    normalized_writes,
                )
                if not temp_verification.verified:
                    raise WorkbookVerificationError("workbook_verification_failed")
                expected_after = fingerprint_file(temp_path)
                written_at = _utc_now()
                manifest["write"] = {
                    "status": "pending",
                    "operation_digest": operation_digest,
                    "before_fingerprint": before.as_dict(),
                    "after_fingerprint": expected_after.as_dict(),
                    "writes": [item.as_dict() for item in normalized_writes],
                    "written_at": written_at.isoformat(),
                }
                # Persist the complete intent and expected output hash before the
                # atomic replacement, so a crash can be reconciled on retry.
                self._write_manifest(normalized_id, manifest)
                os.replace(temp_path, working_path)
            finally:
                temp_path.unlink(missing_ok=True)

            verification = self.verify_xlsx_cells(working_ref, normalized_writes)
            if not verification.verified:
                raise WorkbookVerificationError("workbook_verification_failed")
            after = fingerprint_file(working_path)
            if not _fingerprints_match_content(after, expected_after):
                raise MutationConflictError("working_copy_hash_mismatch_after_write")
            write_payload = dict(manifest["write"])
            write_payload["status"] = "complete"
            write_payload["after_fingerprint"] = after.as_dict()
            write_payload["completed_at"] = _utc_now().isoformat()
            manifest["write"] = write_payload
            self._write_manifest(normalized_id, manifest)
            return WorkbookWriteReceipt(
                mutation_id=normalized_id,
                working_ref=working_ref,
                before_fingerprint=before,
                after_fingerprint=after,
                writes=normalized_writes,
                verification=verification,
                written_at=written_at,
            )

    def verify_xlsx_cells(
        self,
        workbook_ref: ArtifactRef,
        writes: Sequence[CellWrite],
    ) -> WorkbookVerificationReceipt:
        normalized_writes = _normalize_writes(writes)
        workbook_path = self._gateway.resolve(
            workbook_ref,
            expected_type="file",
        )
        if workbook_path.suffix.casefold() != ".xlsx":
            raise UnsupportedWorkbookError("xlsx_desktop_adapter_required")
        return self._verify_workbook_path(
            workbook_path,
            workbook_ref,
            normalized_writes,
        )

    def publish(
        self,
        mutation_id: str,
        working_ref: ArtifactRef,
        published_ref: ArtifactRef,
    ) -> PublishReceipt:
        normalized_id = validate_mutation_id(mutation_id)
        if working_ref.root_id != self._staging_root_id:
            raise RootPermissionError("publish_source_must_be_in_staging_root")

        with self._mutation_lock(normalized_id):
            manifest = self._require_manifest(normalized_id)
            prepared = self._working_receipt_from_manifest(manifest)
            if prepared.working_ref != working_ref:
                raise MutationConflictError("mutation_working_path_conflict")
            self._assert_source_unchanged(prepared)

            source_path = self._gateway.resolve(
                working_ref,
                expected_type="file",
            )
            source_fingerprint = fingerprint_file(source_path)
            expected_source_fingerprint = prepared.working_fingerprint
            if manifest.get("write"):
                write_payload = dict(manifest["write"])
                if str(write_payload.get("status") or "complete") != "complete":
                    raise MutationConflictError("workbook_write_not_reconciled")
                expected_source_fingerprint = _fingerprint_from_dict(
                    dict(write_payload["after_fingerprint"])
                )
            if not _fingerprints_match_content(
                source_fingerprint,
                expected_source_fingerprint,
            ):
                raise MutationConflictError("working_copy_changed_before_publish")

            existing_publish = manifest.get("publish")
            target_path = self._gateway.resolve(
                published_ref,
                must_exist=False,
                for_publish=True,
            )
            if existing_publish:
                publish_payload = dict(existing_publish)
                existing_ref = _ref_from_dict(publish_payload)
                if existing_ref != published_ref:
                    raise MutationConflictError("mutation_publish_target_conflict")
                recorded_source = _fingerprint_from_dict(
                    dict(publish_payload["source_fingerprint"])
                )
                if not _fingerprints_match_content(
                    recorded_source,
                    source_fingerprint,
                ):
                    raise MutationConflictError("publish_source_changed")
                publish_status = str(publish_payload.get("status") or "complete")
                if target_path.exists():
                    published_fingerprint = fingerprint_file(target_path)
                    if not _fingerprints_match_content(
                        published_fingerprint,
                        recorded_source,
                    ):
                        error_code = (
                            "published_artifact_changed"
                            if publish_status == "complete"
                            else "publish_target_conflict"
                        )
                        raise PublishConflictError(error_code)
                    if publish_status == "pending":
                        publish_payload["status"] = "complete"
                        publish_payload["published_fingerprint"] = (
                            published_fingerprint.as_dict()
                        )
                        publish_payload["reconciled_at"] = _utc_now().isoformat()
                        manifest["publish"] = publish_payload
                        self._write_manifest(normalized_id, manifest)
                    return replace(
                        self._publish_receipt_from_manifest(manifest),
                        reused=True,
                    )
                if publish_status == "complete":
                    raise MutationConflictError("published_artifact_missing")
                published_at = _parse_utc(str(publish_payload["published_at"]))
            else:
                published_at = _utc_now()
                if target_path.exists():
                    # Crash recovery for the copy -> manifest window used by
                    # older manifests. Identical content is observationally the
                    # completed requested publish and can be safely adopted.
                    published_fingerprint = fingerprint_file(target_path)
                    if not _fingerprints_match_content(
                        published_fingerprint,
                        source_fingerprint,
                    ):
                        raise PublishConflictError("publish_target_exists")
                    manifest["publish"] = {
                        "status": "complete",
                        **published_ref.as_dict(),
                        "source_fingerprint": source_fingerprint.as_dict(),
                        "published_fingerprint": published_fingerprint.as_dict(),
                        "published_at": published_at.isoformat(),
                        "reconciled_at": published_at.isoformat(),
                    }
                    self._write_manifest(normalized_id, manifest)
                    return PublishReceipt(
                        mutation_id=normalized_id,
                        source_ref=working_ref,
                        source_fingerprint=source_fingerprint,
                        published_ref=published_ref,
                        published_fingerprint=published_fingerprint,
                        published_at=published_at,
                        reused=True,
                    )

                manifest["publish"] = {
                    "status": "pending",
                    **published_ref.as_dict(),
                    "source_fingerprint": source_fingerprint.as_dict(),
                    # The target filesystem may round mtime; reconciliation uses
                    # size + SHA-256 while preserving the expected full record.
                    "published_fingerprint": source_fingerprint.as_dict(),
                    "published_at": published_at.isoformat(),
                }
                self._write_manifest(normalized_id, manifest)

            self._gateway.ensure_parent(published_ref)
            try:
                self._copy_exclusive(
                    source_path,
                    target_path,
                    mutation_id=normalized_id,
                    expected_source=source_fingerprint,
                )
            except MutationConflictError as exc:
                if not target_path.exists():
                    raise
                raced_fingerprint = fingerprint_file(target_path)
                if not _fingerprints_match_content(
                    raced_fingerprint,
                    source_fingerprint,
                ):
                    raise PublishConflictError("publish_target_exists") from exc

            published_fingerprint = fingerprint_file(target_path)
            if not _fingerprints_match_content(
                published_fingerprint,
                source_fingerprint,
            ):
                # Never remove or replace a conflicting shared target: another
                # actor may have won the no-overwrite race.
                raise MutationError("published_artifact_verification_failed")

            publish_payload = dict(manifest["publish"])
            publish_payload["status"] = "complete"
            publish_payload["published_fingerprint"] = (
                published_fingerprint.as_dict()
            )
            publish_payload["completed_at"] = _utc_now().isoformat()
            manifest["publish"] = publish_payload
            self._write_manifest(normalized_id, manifest)
            return PublishReceipt(
                mutation_id=normalized_id,
                source_ref=working_ref,
                source_fingerprint=source_fingerprint,
                published_ref=published_ref,
                published_fingerprint=published_fingerprint,
                published_at=published_at,
            )

    def _verify_workbook_path(
        self,
        workbook_path: Path,
        workbook_ref: ArtifactRef,
        writes: Sequence[CellWrite],
    ) -> WorkbookVerificationReceipt:
        mismatches: list[dict[str, object]] = []
        workbook = load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=False,
            keep_links=True,
        )
        try:
            for item in writes:
                if item.sheet not in workbook.sheetnames:
                    actual: object = None
                    reason = "worksheet_not_found"
                else:
                    actual = workbook[item.sheet][item.cell].value
                    reason = "value_mismatch"
                if item.sheet not in workbook.sheetnames or not _values_equal(item.value, actual):
                    mismatches.append(
                        {
                            "sheet": item.sheet,
                            "cell": item.cell,
                            "expected": _encode_cell_value(item.value),
                            "actual": _encode_cell_value(actual),
                            "reason": reason,
                        }
                    )
        finally:
            workbook.close()

        return WorkbookVerificationReceipt(
            workbook_ref=workbook_ref,
            verified=not mismatches,
            checked_cells=len(writes),
            mismatches=tuple(mismatches),
        )

    def _assert_source_unchanged(self, prepared: WorkingCopyReceipt) -> None:
        current = self._gateway.fingerprint(prepared.source_ref)
        if current != prepared.source_fingerprint:
            raise MutationConflictError("source_changed_since_prepare")

    @staticmethod
    def _new_manifest(
        mutation_id: str,
        source_ref: ArtifactRef,
        source_fingerprint: ArtifactFingerprint,
        working_ref: ArtifactRef,
        working_fingerprint: ArtifactFingerprint,
        created_at: datetime,
    ) -> dict[str, object]:
        return {
            "schema_version": _MANIFEST_SCHEMA_VERSION,
            "mutation_id": mutation_id,
            "created_at": created_at.isoformat(),
            "source": {
                **source_ref.as_dict(),
                "fingerprint": source_fingerprint.as_dict(),
            },
            "working_copy": {
                **working_ref.as_dict(),
                "fingerprint": working_fingerprint.as_dict(),
            },
            "write": None,
            "publish": None,
        }

    def _manifest_ref(self, mutation_id: str) -> ArtifactRef:
        return ArtifactRef(
            self._staging_root_id,
            f".execution-mutations/{mutation_id}/manifest.json",
        )

    def _read_manifest(self, mutation_id: str) -> dict[str, object] | None:
        ref = self._manifest_ref(mutation_id)
        path = self._gateway.resolve(
            ref,
            must_exist=False,
            for_write=True,
        )
        if not path.exists():
            return None
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise MutationError("mutation_manifest_invalid") from exc
        if int(payload.get("schema_version") or 0) != _MANIFEST_SCHEMA_VERSION:
            raise MutationError("mutation_manifest_version_unsupported")
        if str(payload.get("mutation_id") or "") != mutation_id:
            raise MutationError("mutation_manifest_id_mismatch")
        return payload

    def _require_manifest(self, mutation_id: str) -> dict[str, object]:
        manifest = self._read_manifest(mutation_id)
        if manifest is None:
            raise MutationError("working_copy_not_prepared")
        return manifest

    def _write_manifest(self, mutation_id: str, payload: dict[str, object]) -> None:
        ref = self._manifest_ref(mutation_id)
        path = self._gateway.resolve(
            ref,
            must_exist=False,
            for_write=True,
        )
        self._gateway.ensure_parent(ref)
        temp_path = path.with_name(f".manifest.{uuid.uuid4().hex}.tmp")
        try:
            with temp_path.open("x", encoding="utf-8") as handle:
                json.dump(
                    payload,
                    handle,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temp_path, path)
        finally:
            temp_path.unlink(missing_ok=True)

    @contextmanager
    def _mutation_lock(self, mutation_id: str) -> Iterator[None]:
        with self._thread_locks_guard:
            thread_lock = self._thread_locks.setdefault(
                mutation_id,
                threading.RLock(),
            )

        with thread_lock:
            lock_directory_ref = ArtifactRef(
                self._staging_root_id,
                ".execution-mutations/.locks",
            )
            lock_directory = self._gateway.ensure_directory(lock_directory_ref)
            lock_path = lock_directory / f"{mutation_id}.lock"
            with lock_path.open("a+b") as handle:
                try:
                    import fcntl
                except ImportError:  # pragma: no cover - Windows fallback
                    fcntl = None
                if fcntl is not None:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
                try:
                    yield
                finally:
                    if fcntl is not None:
                        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)

    @staticmethod
    def _copy_exclusive(
        source: Path,
        target: Path,
        *,
        mutation_id: str,
        expected_source: ArtifactFingerprint | None = None,
    ) -> None:
        """Publish a complete file without ever streaming into the final name.

        A target-scoped claim serializes cooperative publishers. The copy is
        written to a mutation-specific partial in the same directory, fsynced
        and hash-verified, then made visible atomically. A crashed retry may
        only reuse or remove the partial named by its own verified claim.
        """

        normalized_id = validate_mutation_id(mutation_id)
        expected = expected_source or fingerprint_file(source)
        current_source = fingerprint_file(source)
        if not _fingerprints_match_content(current_source, expected):
            raise MutationConflictError("source_changed_before_copy")

        claim_path, partial_path, claim_payload = (
            FileMutationService._copy_protocol_paths(
                target,
                normalized_id,
                expected,
            )
        )
        FileMutationService._acquire_copy_claim(claim_path, claim_payload)

        if target.is_symlink():
            FileMutationService._remove_owned_partial(
                partial_path,
                claim_path,
                claim_payload,
            )
            FileMutationService._release_copy_claim(
                claim_path,
                claim_payload,
            )
            raise MutationConflictError("target_exists")
        if target.exists():
            target_fingerprint = fingerprint_file(target)
            if _fingerprints_match_content(target_fingerprint, expected):
                FileMutationService._remove_owned_partial(
                    partial_path,
                    claim_path,
                    claim_payload,
                )
                FileMutationService._release_copy_claim(
                    claim_path,
                    claim_payload,
                )
                return
            FileMutationService._remove_owned_partial(
                partial_path,
                claim_path,
                claim_payload,
            )
            FileMutationService._release_copy_claim(claim_path, claim_payload)
            raise MutationConflictError("target_exists")

        if partial_path.exists():
            if partial_path.is_symlink() or not partial_path.is_file():
                raise MutationConflictError("copy_partial_unsafe")
            partial_fingerprint = fingerprint_file(partial_path)
            if not _fingerprints_match_content(partial_fingerprint, expected):
                FileMutationService._remove_owned_partial(
                    partial_path,
                    claim_path,
                    claim_payload,
                )
            else:
                with partial_path.open("rb") as partial_handle:
                    FileMutationService._fsync_best_effort(
                        partial_handle.fileno()
                    )

        if not partial_path.exists():
            # Deliberately retain claim + partial after an interrupted copy.
            # A restarted invocation can prove ownership, reject symlinks,
            # remove only this mutation's incomplete partial, and resume.
            with source.open("rb") as source_handle, partial_path.open(
                "xb"
            ) as partial_handle:
                shutil.copyfileobj(
                    source_handle,
                    partial_handle,
                    length=1024 * 1024,
                )
                partial_handle.flush()
                FileMutationService._fsync_best_effort(partial_handle.fileno())
            try:
                shutil.copystat(source, partial_path, follow_symlinks=False)
            except OSError:
                # SMB metadata support varies; SHA-256 is authoritative.
                pass

        partial_fingerprint = fingerprint_file(partial_path)
        source_after_copy = fingerprint_file(source)
        if (
            not _fingerprints_match_content(partial_fingerprint, expected)
            or not _fingerprints_match_content(source_after_copy, expected)
        ):
            FileMutationService._remove_owned_partial(
                partial_path,
                claim_path,
                claim_payload,
            )
            FileMutationService._release_copy_claim(claim_path, claim_payload)
            raise MutationConflictError("source_changed_during_copy")

        if target.is_symlink():
            FileMutationService._remove_owned_partial(
                partial_path,
                claim_path,
                claim_payload,
            )
            FileMutationService._release_copy_claim(
                claim_path,
                claim_payload,
            )
            raise MutationConflictError("target_exists")
        if target.exists():
            # A non-cooperative actor raced the claimed publication. Never use
            # replace in that case and never remove the actor's target.
            target_fingerprint = fingerprint_file(target)
            if not _fingerprints_match_content(target_fingerprint, expected):
                FileMutationService._remove_owned_partial(
                    partial_path,
                    claim_path,
                    claim_payload,
                )
                FileMutationService._release_copy_claim(
                    claim_path,
                    claim_payload,
                )
                raise MutationConflictError("target_exists")
            FileMutationService._remove_owned_partial(
                partial_path,
                claim_path,
                claim_payload,
            )
        else:
            try:
                # Preferred atomic no-overwrite commit.
                os.link(partial_path, target, follow_symlinks=False)
                partial_path.unlink()
            except OSError as link_error:
                if link_error.errno == errno.EEXIST:
                    if not target.exists() or target.is_symlink():
                        FileMutationService._remove_owned_partial(
                            partial_path,
                            claim_path,
                            claim_payload,
                        )
                        FileMutationService._release_copy_claim(
                            claim_path,
                            claim_payload,
                        )
                        raise MutationConflictError("target_exists") from link_error
                    target_fingerprint = fingerprint_file(target)
                    if not _fingerprints_match_content(
                        target_fingerprint,
                        expected,
                    ):
                        FileMutationService._remove_owned_partial(
                            partial_path,
                            claim_path,
                            claim_payload,
                        )
                        FileMutationService._release_copy_claim(
                            claim_path,
                            claim_payload,
                        )
                        raise MutationConflictError("target_exists") from link_error
                    FileMutationService._remove_owned_partial(
                        partial_path,
                        claim_path,
                        claim_payload,
                    )
                    # The identical target is the requested observable result.
                    # Continue to final hash verification and claim release.
                else:
                    unsupported_link = {
                        errno.EACCES,
                        errno.EINVAL,
                        errno.EPERM,
                        errno.EXDEV,
                        getattr(errno, "ENOSYS", errno.EPERM),
                        getattr(errno, "ENOTSUP", errno.EPERM),
                        getattr(errno, "EOPNOTSUPP", errno.EPERM),
                    }
                    if link_error.errno not in unsupported_link:
                        raise
                    if target.is_symlink():
                        FileMutationService._remove_owned_partial(
                            partial_path,
                            claim_path,
                            claim_payload,
                        )
                        FileMutationService._release_copy_claim(
                            claim_path,
                            claim_payload,
                        )
                        raise MutationConflictError("target_exists") from link_error
                    if target.exists():
                        target_fingerprint = fingerprint_file(target)
                        if not _fingerprints_match_content(
                            target_fingerprint,
                            expected,
                        ):
                            FileMutationService._remove_owned_partial(
                                partial_path,
                                claim_path,
                                claim_payload,
                            )
                            FileMutationService._release_copy_claim(
                                claim_path,
                                claim_payload,
                            )
                            raise MutationConflictError(
                                "target_exists"
                            ) from link_error
                        FileMutationService._remove_owned_partial(
                            partial_path,
                            claim_path,
                            claim_payload,
                        )
                    else:
                        # Under the verified target claim, rename/replace only
                        # makes the fully written same-directory partial visible.
                        # Other execution publishers cannot hold this claim.
                        os.replace(partial_path, target)

        final_fingerprint = fingerprint_file(target)
        if not _fingerprints_match_content(final_fingerprint, expected):
            FileMutationService._release_copy_claim(claim_path, claim_payload)
            raise MutationConflictError("target_hash_mismatch_after_commit")
        FileMutationService._fsync_directory_best_effort(target.parent)
        FileMutationService._release_copy_claim(claim_path, claim_payload)

    @staticmethod
    def _copy_protocol_paths(
        target: Path,
        mutation_id: str,
        expected: ArtifactFingerprint,
    ) -> tuple[Path, Path, dict[str, object]]:
        target_key = hashlib.sha256(
            target.name.casefold().encode("utf-8")
        ).hexdigest()[:20]
        operation_key = hashlib.sha256(
            (
                f"{mutation_id}\0{target.name.casefold()}\0"
                f"{expected.size}\0{expected.sha256}"
            ).encode("utf-8")
        ).hexdigest()[:32]
        claim_path = target.parent / f".execution-{target_key}.claim"
        partial_name = f".execution-{operation_key}.partial"
        partial_path = target.parent / partial_name
        payload: dict[str, object] = {
            "schema_version": 1,
            "mutation_id": mutation_id,
            "target_name": target.name,
            "partial_name": partial_name,
            "source_size": expected.size,
            "source_sha256": expected.sha256,
        }
        return claim_path, partial_path, payload

    @staticmethod
    def _acquire_copy_claim(
        claim_path: Path,
        expected_payload: dict[str, object],
    ) -> None:
        encoded = json.dumps(
            expected_payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
        flags |= getattr(os, "O_BINARY", 0)
        try:
            descriptor = os.open(claim_path, flags, 0o600)
        except FileExistsError:
            existing = FileMutationService._read_copy_claim(claim_path)
            if existing != expected_payload:
                raise MutationConflictError("target_claimed_by_other_mutation")
            return
        except OSError as exc:
            if exc.errno != errno.EEXIST:
                raise
            existing = FileMutationService._read_copy_claim(claim_path)
            if existing != expected_payload:
                raise MutationConflictError("target_claimed_by_other_mutation")
            return

        with os.fdopen(descriptor, "wb") as handle:
            handle.write(encoded)
            handle.flush()
            FileMutationService._fsync_best_effort(handle.fileno())

    @staticmethod
    def _read_copy_claim(claim_path: Path) -> dict[str, object]:
        if claim_path.is_symlink() or not claim_path.is_file():
            raise MutationConflictError("target_claim_invalid")
        try:
            payload = json.loads(claim_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise MutationConflictError("target_claim_invalid") from exc
        if not isinstance(payload, dict):
            raise MutationConflictError("target_claim_invalid")
        return payload

    @staticmethod
    def _remove_owned_partial(
        partial_path: Path,
        claim_path: Path,
        expected_claim: dict[str, object],
    ) -> None:
        if FileMutationService._read_copy_claim(claim_path) != expected_claim:
            raise MutationConflictError("target_claim_ownership_lost")
        if partial_path.exists():
            if partial_path.is_symlink() or not partial_path.is_file():
                raise MutationConflictError("copy_partial_unsafe")
            partial_path.unlink()

    @staticmethod
    def _release_copy_claim(
        claim_path: Path,
        expected_claim: dict[str, object],
    ) -> None:
        if not claim_path.exists():
            return
        if FileMutationService._read_copy_claim(claim_path) != expected_claim:
            raise MutationConflictError("target_claim_ownership_lost")
        claim_path.unlink()

    @staticmethod
    def _fsync_best_effort(descriptor: int) -> None:
        try:
            os.fsync(descriptor)
        except OSError as exc:
            unsupported = {
                errno.EINVAL,
                getattr(errno, "ENOTSUP", errno.EINVAL),
                getattr(errno, "EOPNOTSUPP", errno.EINVAL),
            }
            if exc.errno not in unsupported:
                raise

    @staticmethod
    def _fsync_directory_best_effort(directory: Path) -> None:
        flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0)
        try:
            descriptor = os.open(directory, flags)
        except OSError:
            return
        try:
            FileMutationService._fsync_best_effort(descriptor)
        finally:
            os.close(descriptor)

    @staticmethod
    def _working_receipt_from_manifest(
        manifest: dict[str, object],
    ) -> WorkingCopyReceipt:
        source = dict(manifest["source"])
        working = dict(manifest["working_copy"])
        return WorkingCopyReceipt(
            mutation_id=str(manifest["mutation_id"]),
            source_ref=_ref_from_dict(source),
            source_fingerprint=_fingerprint_from_dict(dict(source["fingerprint"])),
            working_ref=_ref_from_dict(working),
            working_fingerprint=_fingerprint_from_dict(dict(working["fingerprint"])),
            created_at=_parse_utc(str(manifest["created_at"])),
        )

    @staticmethod
    def _write_receipt_from_manifest(
        manifest: dict[str, object],
    ) -> WorkbookWriteReceipt:
        write = dict(manifest["write"])
        writes = tuple(
            CellWrite(
                sheet=str(item["sheet"]),
                cell=str(item["cell"]),
                value=_decode_cell_value(dict(item["value"])),
            )
            for item in write["writes"]
        )
        working = _ref_from_dict(dict(manifest["working_copy"]))
        verification = WorkbookVerificationReceipt(
            workbook_ref=working,
            verified=True,
            checked_cells=len(writes),
            mismatches=(),
        )
        return WorkbookWriteReceipt(
            mutation_id=str(manifest["mutation_id"]),
            working_ref=working,
            before_fingerprint=_fingerprint_from_dict(
                dict(write["before_fingerprint"])
            ),
            after_fingerprint=_fingerprint_from_dict(
                dict(write["after_fingerprint"])
            ),
            writes=writes,
            verification=verification,
            written_at=_parse_utc(str(write["written_at"])),
        )

    @staticmethod
    def _publish_receipt_from_manifest(
        manifest: dict[str, object],
    ) -> PublishReceipt:
        publish = dict(manifest["publish"])
        return PublishReceipt(
            mutation_id=str(manifest["mutation_id"]),
            source_ref=_ref_from_dict(dict(manifest["working_copy"])),
            source_fingerprint=_fingerprint_from_dict(
                dict(publish["source_fingerprint"])
            ),
            published_ref=_ref_from_dict(publish),
            published_fingerprint=_fingerprint_from_dict(
                dict(publish["published_fingerprint"])
            ),
            published_at=_parse_utc(str(publish["published_at"])),
        )
