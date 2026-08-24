from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.execution.errors import ExecutionApiError
from app.execution.persistence import build_file_gateway
from app.execution.registry import node_registry
from app.execution.storage import ArtifactRef, StorageError


def _ref(value: Any) -> ArtifactRef:
    if isinstance(value, dict) and "source" in value:
        value = value["source"]
    if not isinstance(value, dict):
        raise ExecutionApiError(422, "artifact_ref_required", "缺少工作簿文件引用")
    try:
        return ArtifactRef(
            str(value["root_id"]),
            str(value["relative_path"]),
        )
    except (KeyError, ValueError, StorageError) as exc:
        raise ExecutionApiError(422, "artifact_ref_invalid", "工作簿文件引用不合法") from exc


class _WorkbookReader:
    def __init__(self, path: Path, *, data_only: bool = True) -> None:
        self.path = path
        self.data_only = data_only
        self.kind = "openpyxl"
        self.book = None

    def __enter__(self):
        suffix = self.path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            self.book = load_workbook(
                self.path,
                read_only=True,
                data_only=self.data_only,
                keep_vba=suffix in {".xlsm", ".xltm"},
                keep_links=True,
            )
        elif suffix in {".xls", ".xlt"}:
            try:
                import xlrd
            except ImportError as exc:
                raise ExecutionApiError(
                    503,
                    "xls_reader_unavailable",
                    "旧版 Excel 读取组件不可用",
                ) from exc
            self.kind = "xlrd"
            self.book = xlrd.open_workbook(
                str(self.path),
                on_demand=True,
                formatting_info=False,
            )
        else:
            raise ExecutionApiError(
                422,
                "workbook_format_unsupported",
                f"不支持读取 {suffix or '无扩展名'} 工作簿",
            )
        return self

    def __exit__(self, *_args):
        if self.book is not None:
            self.book.close()

    @property
    def sheet_names(self) -> list[str]:
        if self.kind == "openpyxl":
            return list(self.book.sheetnames)
        return list(self.book.sheet_names())

    def value(self, sheet: str, cell: str):
        if self.kind == "openpyxl":
            if sheet not in self.book.sheetnames:
                raise KeyError(sheet)
            return self.book[sheet][cell].value
        if sheet not in self.book.sheet_names():
            raise KeyError(sheet)
        from openpyxl.utils.cell import coordinate_to_tuple

        row, column = coordinate_to_tuple(cell.replace("$", "").upper())
        return self.book.sheet_by_name(sheet).cell_value(row - 1, column - 1)


def _path(context, ref: ArtifactRef) -> Path:
    try:
        return build_file_gateway(context.db).resolve(
            ref,
            expected_type="file",
        )
    except StorageError as exc:
        raise ExecutionApiError(
            422,
            "workbook_unavailable",
            "工作簿不存在或不在允许的根目录内",
        ) from exc


def _feature_matches(reader: _WorkbookReader, feature: dict[str, Any]) -> bool:
    try:
        value = reader.value(str(feature["sheet"]), str(feature["cell"]))
    except (KeyError, IndexError):
        return False
    if feature.get("nonempty"):
        return value not in (None, "")
    if "equals" in feature:
        return value == feature["equals"]
    if "contains" in feature:
        return str(feature["contains"]) in str(value or "")
    return False


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _classify_executor(
    context,
    *,
    detached_io: bool = False,
) -> dict[str, Any]:
    ref = _ref(
        context.input_data.get("source")
        or context.input_data.get("file")
        or context.input_data
    )
    path = _path(context, ref)
    config = context.node.get("config") or {}
    if detached_io:
        context.db.flush()
        context.db.commit()
    with _WorkbookReader(path, data_only=bool(config.get("data_only", True))) as reader:
        matches = []
        for candidate in config.get("types") or []:
            features = candidate.get("features") or []
            if features and all(_feature_matches(reader, feature) for feature in features):
                matches.append(str(candidate.get("name") or "未命名类型"))
        return {
            "source": ref.as_dict(),
            "format": path.suffix.casefold(),
            "sheet_names": reader.sheet_names,
            "matched_types": matches,
            "detected_type": matches[0] if len(matches) == 1 else None,
            "ambiguous": len(matches) > 1,
        }


def _extract_executor(
    context,
    *,
    detached_io: bool = False,
) -> dict[str, Any]:
    ref = _ref(
        context.input_data.get("source")
        or context.input_data.get("file")
        or context.input_data
    )
    path = _path(context, ref)
    config = context.node.get("config") or {}
    if detached_io:
        context.db.flush()
        context.db.commit()
    values: dict[str, Any] = {}
    errors: list[dict[str, str]] = []
    with _WorkbookReader(path, data_only=bool(config.get("data_only", True))) as reader:
        for field in config.get("fields") or []:
            name = str(field.get("name") or "")
            if not name:
                continue
            try:
                value = reader.value(str(field["sheet"]), str(field["cell"]))
                values[name] = _json_value(value)
                if field.get("required") and value in (None, ""):
                    errors.append({"field": name, "code": "required_value_missing"})
            except (KeyError, IndexError):
                errors.append({"field": name, "code": "cell_unavailable"})
    if errors and bool(config.get("fail_on_error", True)):
        raise ExecutionApiError(
            422,
            "workbook_summary_incomplete",
            "工作簿关键字段读取失败",
            details={"errors": errors},
        )
    return {"source": ref.as_dict(), "values": values, "errors": errors}


_REGISTERED = False


def register_excel_executors() -> None:
    global _REGISTERED
    if _REGISTERED:
        return
    node_registry.set_executor("excel.classify", 1, _classify_executor)
    node_registry.set_executor("excel.extract_summary", 1, _extract_executor)
    _REGISTERED = True
