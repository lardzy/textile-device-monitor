from __future__ import annotations

import hashlib
import io
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterator, Optional
from uuid import uuid4

from openpyxl import load_workbook
from PIL import Image as PillowImage
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.execution.errors import ExecutionApiError
from app.execution.models import (
    ExecutionArtifact,
    ExecutionFileIndexEntry,
    ExecutionStorageRoot,
)
from app.execution.persistence import build_file_gateway, storage_root_by_key
from app.execution.registry import node_registry
from app.execution.storage import ArtifactRef, FileGateway, StorageError
from app.execution.workbook_format import (
    SUPPORTED_WORKBOOK_SUFFIXES,
    WorkbookFormat,
    detect_workbook_format,
)


RESULT_NODE_TYPE_VERSION = 1
IMAGE_ARTIFACT_ROOT_ID = "execution_staging"
_PERCENT_NUMBER_FORMAT = re.compile(r"(?<!\\)%")


@dataclass(frozen=True)
class ResultBlock:
    key: str
    part_cell: str
    name_row: int
    content_row: int
    start_column: int
    end_column: int


@dataclass(frozen=True)
class RegeneratedFiberResultRule:
    node_type: str
    method: str
    worksheet: str
    inspector_cell: str
    blocks: tuple[ResultBlock, ...]
    remark_rows: tuple[int, ...]


@dataclass(frozen=True)
class CellData:
    value: Any
    number_format: str = ""


@dataclass(frozen=True)
class EmbeddedImage:
    data: bytes
    extension: str
    media_type: str
    sheet_name: str
    index: int
    width: Optional[int]
    height: Optional[int]
    anchor: Optional[str]


RESULT_RULES = {
    "result.regenerated_fiber_count_method": RegeneratedFiberResultRule(
        node_type="result.regenerated_fiber_count_method",
        method="count",
        worksheet="根数法报告1",
        inspector_cell="I8",
        blocks=(
            ResultBlock("left", "B24", 25, 26, 2, 6),
            ResultBlock("right", "G24", 25, 26, 7, 11),
        ),
        remark_rows=(27, 28, 29),
    ),
    "result.regenerated_fiber_area_method": RegeneratedFiberResultRule(
        node_type="result.regenerated_fiber_area_method",
        method="area",
        worksheet="截面统计报告1",
        inspector_cell="I8",
        blocks=(
            ResultBlock("left", "B26", 27, 28, 2, 6),
            ResultBlock("right", "G26", 27, 28, 7, 11),
        ),
        remark_rows=(29, 30, 31),
    ),
}


def _column_name(column: int) -> str:
    value = ""
    current = column
    while current:
        current, remainder = divmod(current - 1, 26)
        value = chr(ord("A") + remainder) + value
    return value


def _cell_ref(row: int, column: int) -> str:
    return f"{_column_name(column)}{row}"


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _component_name(value: Any) -> Optional[str]:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None if value is False else str(value)
    if isinstance(value, (int, float, Decimal)) and value == 0:
        return None
    text = str(value).strip()
    return text or None


def _text_value(value: Any) -> Optional[str]:
    if _is_blank(value):
        return None
    return str(value).strip()


def _number_format_is_percent(number_format: str) -> bool:
    if not isinstance(number_format, str):
        return False
    # Ignore quoted text so a literal "%" in a label does not change the
    # numeric value. Escaped percent signs are ignored by the regex.
    without_quoted_text = re.sub(r'"[^"]*"', "", number_format)
    return _PERCENT_NUMBER_FORMAT.search(without_quoted_text) is not None


def _display_decimal_places(number_format: str) -> Optional[int]:
    if not isinstance(number_format, str):
        return None
    primary = number_format.split(";", 1)[0]
    primary = re.sub(r'"[^"]*"', "", primary)
    primary = re.sub(r"\\.", "", primary)
    match = re.search(r"[0#?]+(?:\.([0#?]+))?", primary)
    if match is None:
        return None
    decimals = match.group(1)
    return len(decimals) if decimals is not None else 0


def _numeric_content(cell: CellData) -> tuple[Optional[float], Optional[float]]:
    value = cell.value
    if value is None or isinstance(value, bool):
        return None, None
    percent_text = False
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None, None
        if text.endswith(("%", "％")):
            percent_text = True
            text = text[:-1].strip()
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None, None
    elif isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        return None, None
    if not number.is_finite():
        return None, None
    if _number_format_is_percent(cell.number_format) and not percent_text:
        number *= Decimal("100")
    raw = number
    decimal_places = _display_decimal_places(cell.number_format)
    if decimal_places is not None:
        quantum = Decimal(1).scaleb(-decimal_places)
        number = number.quantize(quantum, rounding=ROUND_HALF_UP)
    if number == 0:
        number = abs(number)
    if raw == 0:
        raw = abs(raw)
    return float(raw), float(number)


def _balance_component_contents(
    components: list[dict[str, Any]],
) -> tuple[Optional[dict[str, Any]], Optional[dict[str, float]]]:
    """Build report-ready percentages while retaining Excel/raw values.

    Excel rounds every cell independently, so a mathematically valid group can
    visibly add up to 99.9 or 100.1. Conversely, cached legacy formulas can be
    a few hundredths away from 100 while their displayed result is exactly
    100. Use the displayed values as the report result, then apply at most one
    small tail-difference correction to the component closest to its
    normalized raw share. This keeps the business-facing values at exactly
    100 without losing either source representation.
    """
    if not components or any(
        component.get("raw_content") is None
        or component.get("display_content") is None
        for component in components
    ):
        return None, None

    raw_values = [
        Decimal(str(component["raw_content"]))
        for component in components
    ]
    display_values = [
        Decimal(str(component["display_content"]))
        for component in components
    ]
    raw_total = sum(raw_values, Decimal("0"))
    display_total = sum(display_values, Decimal("0"))
    target = Decimal("100")

    for component, display in zip(components, display_values):
        component["content"] = float(display)

    residual = target - display_total
    if residual == 0:
        return None, None
    # A larger difference signals incomplete/corrupt input rather than normal
    # one-decimal rounding and must not be silently normalized.
    if (
        raw_total <= 0
        or abs(raw_total - target) > Decimal("1")
        or abs(residual) > Decimal("1")
    ):
        return None, {
            "raw_total": float(raw_total),
            "display_total": float(display_total),
        }

    normalized_raw = [
        value * target / raw_total
        for value in raw_values
    ]
    adjusted_index = min(
        range(len(components)),
        key=lambda index: (
            abs(
                display_values[index]
                + residual
                - normalized_raw[index]
            ),
            index,
        ),
    )
    adjusted = display_values[adjusted_index] + residual
    if adjusted < 0:
        return None, {
            "raw_total": float(raw_total),
            "display_total": float(display_total),
        }
    components[adjusted_index]["content"] = float(adjusted)
    return {
        "component_name": components[adjusted_index]["name"],
        "delta": float(residual),
        "excel_display_total": float(display_total),
    }, None


def _anchor_label(image: Any) -> Optional[str]:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    try:
        return _cell_ref(int(marker.row) + 1, int(marker.col) + 1)
    except (AttributeError, TypeError, ValueError):
        return None


def _verified_image(
    raw: bytes,
    *,
    sheet_name: str,
    index: int,
    anchor: Optional[str],
) -> EmbeddedImage:
    try:
        with PillowImage.open(io.BytesIO(raw)) as image:
            image.verify()
        with PillowImage.open(io.BytesIO(raw)) as image:
            image_format = str(image.format or "").upper()
            width, height = image.size
    except Exception as exc:
        raise ExecutionApiError(
            422,
            "embedded_image_invalid",
            "工作簿中存在无法读取的插图",
            details={"sheet": sheet_name, "image_index": index},
        ) from exc

    extensions = {
        "BMP": ".bmp",
        "GIF": ".gif",
        "JPEG": ".jpg",
        "PNG": ".png",
        "TIFF": ".tiff",
        "WEBP": ".webp",
    }
    extension = extensions.get(image_format)
    if extension is None:
        raise ExecutionApiError(
            415,
            "embedded_image_format_unsupported",
            "工作簿插图格式暂不支持在线查看",
            details={
                "sheet": sheet_name,
                "image_index": index,
                "format": image_format,
            },
        )
    return EmbeddedImage(
        data=raw,
        extension=extension,
        media_type=mimetypes.guess_type(f"image{extension}")[0]
        or "application/octet-stream",
        sheet_name=sheet_name,
        index=index,
        width=width,
        height=height,
        anchor=anchor,
    )


def _extract_openpyxl_images(workbook: Any) -> tuple[list[EmbeddedImage], list[dict[str, Any]]]:
    images: list[EmbeddedImage] = []
    warnings: list[dict[str, Any]] = []
    image_index = 0
    for worksheet in workbook.worksheets:
        for image in list(getattr(worksheet, "_images", ())):
            image_index += 1
            try:
                raw = image._data()
                images.append(
                    _verified_image(
                        raw,
                        sheet_name=worksheet.title,
                        index=image_index,
                        anchor=_anchor_label(image),
                    )
                )
            except ExecutionApiError as exc:
                warnings.append(
                    {
                        "code": exc.code,
                        "message": exc.message,
                        "details": exc.details,
                    }
                )
            except Exception as exc:
                warnings.append(
                    {
                        "code": "embedded_image_read_failed",
                        "message": "工作簿中的一张插图读取失败",
                        "details": {
                            "sheet": worksheet.title,
                            "image_index": image_index,
                            "error_type": type(exc).__name__,
                        },
                    }
                )
    return images, warnings


def _build_result(
    *,
    rule: RegeneratedFiberResultRule,
    read_cell,
) -> dict[str, Any]:
    part_names = {
        block.key: _text_value(read_cell(block.part_cell).value)
        for block in rule.blocks
    }
    has_parts = any(name is not None for name in part_names.values())
    warnings: list[dict[str, Any]] = []

    def components_for(block: ResultBlock) -> list[dict[str, Any]]:
        components: list[dict[str, Any]] = []
        for column in range(block.start_column, block.end_column + 1):
            name_cell = _cell_ref(block.name_row, column)
            content_cell = _cell_ref(block.content_row, column)
            name = _component_name(read_cell(name_cell).value)
            content_data = read_cell(content_cell)
            if name is None:
                if not _is_blank(content_data.value):
                    warnings.append(
                        {
                            "code": "content_without_component_name",
                            "message": "检测到没有对应纤维名称的含量",
                            "cell": content_cell,
                        }
                    )
                continue
            raw_content, display_content = _numeric_content(content_data)
            if raw_content is None:
                warnings.append(
                    {
                        "code": "component_content_not_numeric",
                        "message": f"纤维“{name}”的含量不是有效数字",
                        "cell": content_cell,
                    }
                )
            components.append(
                {
                    "name": name,
                    "content": display_content,
                    "display_content": display_content,
                    "raw_content": raw_content,
                    "name_cell": name_cell,
                    "content_cell": content_cell,
                }
            )
        return components

    parts: list[dict[str, Any]] = []

    def total_for(
        components: list[dict[str, Any]],
        key: str,
    ) -> float:
        return float(
            sum(
                (
                    Decimal(str(component[key]))
                    for component in components
                    if component[key] is not None
                ),
                Decimal("0"),
            )
        )

    if has_parts:
        for block in rule.blocks:
            part_name = part_names[block.key]
            if part_name is None:
                continue
            components = components_for(block)
            adjustment, invalid_total = _balance_component_contents(components)
            if invalid_total is not None:
                warnings.append(
                    {
                        "code": "component_total_not_100",
                        "message": f"部位“{part_name}”的成分含量合计异常",
                        "details": invalid_total,
                    }
                )
            part = {
                "key": block.key,
                "name": part_name,
                "part_cell": block.part_cell,
                "components": components,
                "total": total_for(components, "content"),
                "display_total": total_for(
                    components,
                    "display_content",
                ),
                "raw_total": total_for(components, "raw_content"),
            }
            if adjustment is not None:
                part["rounding_adjustment"] = adjustment
            parts.append(part)
    else:
        # Real legacy templates can contain two independent 100% result groups
        # even when both part-name cells are blank. Combining the left and
        # right blocks would manufacture a 200% result, while dropping either
        # side would destroy business data. Preserve every populated block as
        # an anonymous result group and let the UI label them 结果1/结果2.
        anonymous = [
            (block, components_for(block))
            for block in rule.blocks
        ]
        anonymous = [
            (block, components)
            for block, components in anonymous
            if components
        ]
        for index, (block, components) in enumerate(anonymous, start=1):
            adjustment, invalid_total = _balance_component_contents(components)
            label = "结果" if len(anonymous) == 1 else f"结果{index}"
            if invalid_total is not None:
                warnings.append(
                    {
                        "code": "component_total_not_100",
                        "message": f"{label}的成分含量合计异常",
                        "details": invalid_total,
                    }
                )
            part = {
                "key": block.key,
                "name": None,
                "label": label,
                "part_cell": block.part_cell,
                "components": components,
                "total": total_for(components, "content"),
                "display_total": total_for(
                    components,
                    "display_content",
                ),
                "raw_total": total_for(components, "raw_content"),
            }
            if adjustment is not None:
                part["rounding_adjustment"] = adjustment
            parts.append(part)

    remarks = [
        {"cell": f"B{row}", "text": value}
        for row in rule.remark_rows
        if (value := _text_value(read_cell(f"B{row}").value)) is not None
    ]
    inspector_name = _component_name(read_cell(rule.inspector_cell).value)
    return {
        "method": rule.method,
        "worksheet": rule.worksheet,
        "inspector": {
            "cell": rule.inspector_cell,
            "name": inspector_name,
        },
        "has_parts": has_parts,
        "parts": parts,
        "remarks": remarks,
        "warnings": warnings,
    }


def _read_modern_workbook(
    path: Path,
    rule: RegeneratedFiberResultRule,
) -> tuple[dict[str, Any], list[EmbeddedImage], list[dict[str, Any]]]:
    with path.open("rb") as stream:
        workbook = load_workbook(
            stream,
            read_only=False,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        try:
            if rule.worksheet not in workbook.sheetnames:
                raise ExecutionApiError(
                    422,
                    "result_worksheet_missing",
                    f"工作簿缺少工作表“{rule.worksheet}”",
                )
            worksheet = workbook[rule.worksheet]

            def read_cell(reference: str) -> CellData:
                cell = worksheet[reference]
                return CellData(cell.value, str(cell.number_format or ""))

            result = _build_result(rule=rule, read_cell=read_cell)
            images, image_warnings = _extract_openpyxl_images(workbook)
            return result, images, image_warnings
        finally:
            workbook.close()


def _xlrd_number_format(workbook: Any, cell: Any) -> str:
    try:
        xf = workbook.xf_list[cell.xf_index]
        number_format = workbook.format_map.get(xf.format_key)
        return str(number_format.format_str if number_format else "")
    except (AttributeError, IndexError, KeyError, TypeError):
        return ""


def _legacy_images_via_libreoffice(
    path: Path,
) -> tuple[list[EmbeddedImage], list[dict[str, Any]]]:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        return [], [
            {
                "code": "legacy_image_reader_unavailable",
                "message": "当前运行环境未安装 LibreOffice，无法读取旧版工作簿插图",
            }
        ]

    with tempfile.TemporaryDirectory(prefix="execution-xls-images-") as directory:
        workdir = Path(directory)
        profile = workdir / "profile"
        source = workdir / "source"
        output = workdir / "output"
        profile.mkdir()
        source.mkdir()
        output.mkdir()
        local_copy = source / path.name
        try:
            shutil.copy2(path, local_copy)
        except OSError as exc:
            return [], [
                {
                    "code": "legacy_image_copy_failed",
                    "message": "无法为旧版工作簿创建只读转换副本",
                    "details": {"error_type": type(exc).__name__},
                }
            ]
        env = os.environ.copy()
        env["HOME"] = str(workdir)
        command = [
            executable,
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            'xlsx:Calc MS Excel 2007 XML',
            "--outdir",
            str(output),
            str(local_copy),
        ]
        try:
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                timeout=90,
                env=env,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            return [], [
                {
                    "code": "legacy_image_conversion_failed",
                    "message": "旧版工作簿插图转换失败",
                    "details": {"error_type": type(exc).__name__},
                }
            ]
        converted = sorted(output.glob("*.xlsx"))
        if completed.returncode != 0 or not converted:
            return [], [
                {
                    "code": "legacy_image_conversion_failed",
                    "message": "旧版工作簿插图转换失败",
                    "details": {"return_code": completed.returncode},
                }
            ]
        with converted[0].open("rb") as stream:
            workbook = load_workbook(
                stream,
                read_only=False,
                data_only=True,
                keep_vba=False,
                keep_links=False,
            )
            try:
                return _extract_openpyxl_images(workbook)
            finally:
                workbook.close()


def _read_legacy_workbook(
    path: Path,
    rule: RegeneratedFiberResultRule,
) -> tuple[dict[str, Any], list[EmbeddedImage], list[dict[str, Any]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ExecutionApiError(
            503,
            "legacy_workbook_reader_unavailable",
            "当前运行环境无法读取旧版 Excel 工作簿",
        ) from exc

    workbook = xlrd.open_workbook(
        str(path),
        on_demand=True,
        formatting_info=True,
    )
    try:
        if rule.worksheet not in workbook.sheet_names():
            raise ExecutionApiError(
                422,
                "result_worksheet_missing",
                f"工作簿缺少工作表“{rule.worksheet}”",
            )
        worksheet = workbook.sheet_by_name(rule.worksheet)

        def read_cell(reference: str) -> CellData:
            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", reference)
            if match is None:
                raise ValueError("invalid_cell_reference")
            column_name, row_text = match.groups()
            column = 0
            for character in column_name:
                column = column * 26 + ord(character) - ord("A") + 1
            row_index = int(row_text) - 1
            column_index = column - 1
            if row_index >= worksheet.nrows or column_index >= worksheet.ncols:
                return CellData(None)
            cell = worksheet.cell(row_index, column_index)
            return CellData(
                cell.value,
                _xlrd_number_format(workbook, cell),
            )

        result = _build_result(rule=rule, read_cell=read_cell)
    finally:
        workbook.release_resources()
    images, image_warnings = _legacy_images_via_libreoffice(path)
    return result, images, image_warnings


def read_regenerated_fiber_result(
    path: Path,
    *,
    node_type: str,
) -> tuple[dict[str, Any], list[EmbeddedImage]]:
    rule = RESULT_RULES.get(node_type)
    if rule is None:
        raise ExecutionApiError(
            422,
            "regenerated_fiber_result_rule_unknown",
            "未知的再生纤结果读取规则",
        )
    suffix = path.suffix.casefold()
    if suffix not in SUPPORTED_WORKBOOK_SUFFIXES:
        raise ExecutionApiError(
            415,
            "result_workbook_format_unsupported",
            "该文件格式不支持结果读取",
            details={"extension": suffix},
        )
    try:
        workbook_format = detect_workbook_format(path)
        if workbook_format is None:
            raise ExecutionApiError(
                415,
                "result_workbook_format_unsupported",
                "无法识别工作簿的实际文件格式",
                details={"extension": suffix},
            )
        if workbook_format is WorkbookFormat.OLE:
            result, images, image_warnings = _read_legacy_workbook(path, rule)
        else:
            result, images, image_warnings = _read_modern_workbook(path, rule)
    except ExecutionApiError:
        raise
    except Exception as exc:
        raise ExecutionApiError(
            422,
            "result_workbook_read_failed",
            "工作簿结果读取失败",
            details={"error_type": type(exc).__name__},
        ) from exc
    result["warnings"].extend(image_warnings)
    result["image_count"] = len(images)
    return result, images


def summarize_result_inspectors(
    files: list[dict[str, Any]],
) -> dict[str, Any]:
    """Summarize inspector names without silently resolving a conflict."""

    file_inspectors: list[dict[str, Optional[str]]] = []
    names: list[str] = []
    normalized_names: set[str] = set()
    missing_count = 0
    for item in files:
        if item.get("read_status") == "failed":
            continue
        result = item.get("result")
        inspector = (
            result.get("inspector")
            if isinstance(result, dict)
            else None
        )
        raw_name = (
            inspector.get("name")
            if isinstance(inspector, dict)
            else None
        )
        name = _text_value(raw_name)
        file_id = str(item.get("id") or "")
        file_inspectors.append({"file_id": file_id, "name": name})
        if name is None:
            missing_count += 1
            continue
        normalized = name.casefold()
        if normalized not in normalized_names:
            normalized_names.add(normalized)
            names.append(name)
    conflict = len(names) > 1
    return {
        "name": names[0] if len(names) == 1 else None,
        "names": names,
        "conflict": conflict,
        "missing_count": missing_count,
        "files": file_inspectors,
    }


def _atomic_image_write(path: Path, data: bytes) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() == data:
            return False
        raise ExecutionApiError(
            409,
            "artifact_file_conflict",
            "插图制品路径已存在但内容不一致",
        )
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        with temporary.open("xb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        return True
    finally:
        temporary.unlink(missing_ok=True)


def _artifact_record(
    db: Session,
    *,
    root: ExecutionStorageRoot,
    run_id: str,
    node_run_id: str,
    source_file_id: str,
    source_fingerprint: str,
    image: EmbeddedImage,
    digest: str,
    relative_path: str,
) -> tuple[ExecutionArtifact, bool]:
    def existing_record() -> Optional[ExecutionArtifact]:
        return (
            db.query(ExecutionArtifact)
            .filter(
                ExecutionArtifact.run_id == run_id,
                ExecutionArtifact.storage_root_id == root.id,
                ExecutionArtifact.relative_path == relative_path,
                ExecutionArtifact.content_sha256 == digest,
            )
            .one_or_none()
        )

    existing = existing_record()
    if existing is not None:
        return existing, False
    artifact = ExecutionArtifact(
        run_id=run_id,
        node_run_id=node_run_id,
        storage_root_id=root.id,
        relative_path=relative_path,
        filename=f"插图-{image.index}{image.extension}",
        role="preview",
        media_type=image.media_type,
        size_bytes=len(image.data),
        content_sha256=digest,
        immutable=True,
        metadata_json={
            "source_file_id": source_file_id,
            "source_fingerprint": source_fingerprint,
            "sheet": image.sheet_name,
            "image_index": image.index,
            "width": image.width,
            "height": image.height,
            "anchor": image.anchor,
        },
    )
    try:
        # The savepoint is deliberately narrower than the candidate savepoint:
        # a concurrent insert may violate the identity constraint, but must not
        # poison the worker's outer Session.
        with db.begin_nested():
            db.add(artifact)
            db.flush()
        return artifact, True
    except IntegrityError as exc:
        existing = existing_record()
        if existing is None:
            raise ExecutionApiError(
                409,
                "artifact_registration_conflict",
                "插图制品并发登记冲突，请重试该节点",
            ) from exc
        return existing, False


def _image_artifact(
    db: Session,
    *,
    gateway: FileGateway,
    run_id: str,
    node_run_id: str,
    source_file_id: str,
    source_fingerprint: str,
    image: EmbeddedImage,
) -> tuple[ExecutionArtifact, Optional[Path]]:
    digest = hashlib.sha256(image.data).hexdigest()
    relative_path = (
        f"embedded-images/{run_id}/{node_run_id}/{source_file_id}/"
        f"{image.index:03d}-{digest[:16]}{image.extension}"
    )
    ref = ArtifactRef(IMAGE_ARTIFACT_ROOT_ID, relative_path)
    target = gateway.resolve(ref, must_exist=False, for_write=True)
    root = storage_root_by_key(db, IMAGE_ARTIFACT_ROOT_ID)
    artifact, created_record = _artifact_record(
        db,
        root=root,
        run_id=run_id,
        node_run_id=node_run_id,
        source_file_id=source_file_id,
        source_fingerprint=source_fingerprint,
        image=image,
        digest=digest,
        relative_path=relative_path,
    )
    created_file = _atomic_image_write(target, image.data)
    # Rebuilding a missing file for an existing artifact is a repair, not a
    # candidate-owned write. It remains in place if a later image fails.
    return artifact, target if created_record and created_file else None


def _source_fingerprint(path: Path) -> str:
    stat = path.stat()
    return f"{stat.st_size}:{stat.st_mtime_ns}"


def _assert_source_fingerprint(
    path: Path,
    *,
    expected: str,
    candidate_id: str,
) -> None:
    try:
        current = _source_fingerprint(path)
    except OSError as exc:
        raise ExecutionApiError(
            409,
            "result_file_stale",
            "待读取文件在读取期间变得不可用，请重新查询",
            details={"candidate_id": candidate_id},
        ) from exc
    if current != expected:
        raise ExecutionApiError(
            409,
            "result_file_stale",
            "待读取文件已变化，请重新查询",
            details={"candidate_id": candidate_id},
        )


@contextmanager
def _stable_workbook_snapshot(
    path: Path,
    *,
    expected_fingerprint: str,
    candidate_id: str,
) -> Iterator[Path]:
    _assert_source_fingerprint(
        path,
        expected=expected_fingerprint,
        candidate_id=candidate_id,
    )
    with tempfile.TemporaryDirectory(
        prefix="execution-result-snapshot-"
    ) as directory:
        snapshot = Path(directory) / f"source{path.suffix.casefold()}"
        try:
            shutil.copyfile(path, snapshot)
        except OSError as exc:
            raise ExecutionApiError(
                409,
                "result_file_stale",
                "无法创建待读取文件的稳定副本，请重新查询",
                details={"candidate_id": candidate_id},
            ) from exc
        _assert_source_fingerprint(
            path,
            expected=expected_fingerprint,
            candidate_id=candidate_id,
        )
        expected_size = int(expected_fingerprint.split(":", 1)[0])
        if snapshot.stat().st_size != expected_size:
            raise ExecutionApiError(
                409,
                "result_file_stale",
                "待读取文件在复制期间发生变化，请重新查询",
                details={"candidate_id": candidate_id},
            )
        try:
            yield snapshot
        finally:
            # Re-check after both cell and embedded-image extraction. The
            # returned result therefore always describes one local snapshot.
            _assert_source_fingerprint(
                path,
                expected=expected_fingerprint,
                candidate_id=candidate_id,
            )


def _candidate_entry(
    db: Session,
    candidate: dict[str, Any],
) -> tuple[ExecutionFileIndexEntry, ExecutionStorageRoot]:
    candidate_id = str(candidate.get("id") or "")
    row = (
        db.query(ExecutionFileIndexEntry, ExecutionStorageRoot)
        .join(
            ExecutionStorageRoot,
            ExecutionStorageRoot.id
            == ExecutionFileIndexEntry.storage_root_id,
        )
        .filter(ExecutionFileIndexEntry.id == candidate_id)
        .one_or_none()
    )
    if row is None:
        raise ExecutionApiError(
            409,
            "result_file_not_indexed",
            "待读取文件已不在索引中",
            details={"candidate_id": candidate_id},
        )
    entry, root = row
    if (
        entry.missing_since is not None
        or root.root_id != "regenerated_fiber_records"
        or candidate.get("root_id") != root.root_id
        or candidate.get("relative_path") != entry.relative_path
        or candidate.get("fingerprint") != entry.fingerprint
    ):
        raise ExecutionApiError(
            409,
            "result_file_stale",
            "待读取文件已变化，请重新查询",
            details={"candidate_id": candidate_id},
        )
    return entry, root


def _result_executor(context) -> dict[str, Any]:
    node_type = context.node_run.node_type
    if node_type not in RESULT_RULES:
        raise ExecutionApiError(
            422,
            "regenerated_fiber_result_rule_unknown",
            "未知的再生纤结果读取规则",
        )
    files = context.input_data.get("files")
    if not isinstance(files, list) or not files:
        raise ExecutionApiError(
            422,
            "result_files_required",
            "结果读取节点至少需要一个文件",
        )
    gateway = build_file_gateway(context.db)
    output_files: list[dict[str, Any]] = []
    success_count = 0
    for raw_candidate in files:
        candidate = dict(raw_candidate) if isinstance(raw_candidate, dict) else {}
        try:
            entry, root = _candidate_entry(context.db, candidate)
            path = gateway.resolve(
                ArtifactRef(root.root_id, entry.relative_path),
                expected_type="file",
            )
            with _stable_workbook_snapshot(
                path,
                expected_fingerprint=entry.fingerprint,
                candidate_id=entry.id,
            ) as snapshot:
                result, images = read_regenerated_fiber_result(
                    snapshot,
                    node_type=node_type,
                )
            image_items: list[dict[str, Any]] = []
            created_files: list[Path] = []
            candidate_transaction = context.db.begin_nested()
            try:
                # One failed image must roll back every artifact row created
                # for this candidate. Keep its savepoint open until filesystem
                # compensation is complete so another worker cannot reuse a
                # path while this worker is removing it.
                for image in images:
                    artifact, created_file = _image_artifact(
                        context.db,
                        gateway=gateway,
                        run_id=context.run.id,
                        node_run_id=context.node_run.id,
                        source_file_id=entry.id,
                        source_fingerprint=entry.fingerprint,
                        image=image,
                    )
                    if created_file is not None:
                        created_files.append(created_file)
                    image_items.append(
                        {
                            "artifact_id": artifact.id,
                            "name": artifact.filename,
                            "media_type": artifact.media_type,
                            "sheet": image.sheet_name,
                            "anchor": image.anchor,
                            "width": image.width,
                            "height": image.height,
                            "preview_url": (
                                "/api/execution/v1/artifacts/"
                                f"{artifact.id}/preview"
                            ),
                        }
                    )
                candidate_transaction.commit()
            except Exception as candidate_error:
                cleanup_error: Optional[OSError] = None
                for created_file in reversed(created_files):
                    try:
                        created_file.unlink(missing_ok=True)
                    except OSError as exc:
                        cleanup_error = cleanup_error or exc
                if candidate_transaction.is_active:
                    candidate_transaction.rollback()
                if cleanup_error is not None:
                    raise ExecutionApiError(
                        500,
                        "artifact_cleanup_failed",
                        "插图制品写入失败，且未能完整清理临时文件",
                    ) from cleanup_error
                raise candidate_error
            result["images"] = image_items
            result["image_count"] = len(image_items)
            output_files.append(
                {
                    **candidate,
                    "read_status": "succeeded",
                    "result": result,
                }
            )
            success_count += 1
        except (ExecutionApiError, StorageError, OSError) as exc:
            code = (
                exc.code
                if isinstance(exc, ExecutionApiError)
                else "result_file_unavailable"
            )
            message = (
                exc.message
                if isinstance(exc, ExecutionApiError)
                else "文件当前不可读取"
            )
            output_files.append(
                {
                    **candidate,
                    "read_status": "failed",
                    "error": {
                        "code": code,
                        "message": message,
                    },
                }
            )
    if success_count == 0:
        error_codes = [
            item["error"]["code"]
            for item in output_files
            if isinstance(item.get("error"), dict)
        ]
        raise ExecutionApiError(
            422,
            "result_workbooks_unreadable",
            "候选工作簿均未能读取结果",
            details={
                "file_count": len(output_files),
                "error_codes": error_codes,
            },
        )
    return {
        "files": output_files,
        "count": len(output_files),
        "success_count": success_count,
        "failed_count": len(output_files) - success_count,
        "inspector_summary": summarize_result_inspectors(output_files),
    }


_EXECUTORS_REGISTERED = False


def register_regenerated_fiber_result_executors() -> None:
    global _EXECUTORS_REGISTERED
    if _EXECUTORS_REGISTERED:
        return
    for node_type in RESULT_RULES:
        node_registry.set_executor(
            node_type,
            RESULT_NODE_TYPE_VERSION,
            _result_executor,
        )
    _EXECUTORS_REGISTERED = True
