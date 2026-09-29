"""Template rendering with JSON field mappings and durable artifact receipts."""

import os
import re
import shutil
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from app.execution.errors import ExecutionApiError
from app.execution.models import ExecutionArtifact
from app.execution.persistence import build_file_gateway, storage_root_by_key
from app.execution.storage import ArtifactRef, fingerprint_file, fsync_file
from app.execution.v2.canonical import canonical_sha256
from app.execution.v2.data_handlers import select_value
from app.execution.v2.templates import resolve_template


def render_file(template, target, *, values, fields, images=(), gateway=None, layout=None, number_formats=None):
    """Keep BIFF caches/formulas intact; use openpyxl for OOXML templates."""
    writes = [{**field, "value": select_value(values, field["value"])} for field in fields]
    if template.suffix.lower() == ".xls":
        import xlrd
        from app.execution.biff_patch import CellEdit, patch_workbook_file

        if images:
            if not layout:
                raise ExecutionApiError(422, "xls_layout_required", "XLS 图片渲染需要配置图片区域和打印范围")
            if any(f.get("sheet", layout["sheet"]) != layout["sheet"] for f in writes):
                raise ExecutionApiError(422, "xls_layout_unsupported", "XLS 图片与字段必须位于配置的工作表")
            from app.execution.workbook_images import render_image_workbook

            with TemporaryDirectory(prefix="render-inputs-", dir=target.parent) as temporary:
                selected = []
                for index, item in enumerate(images):
                    reference = item["artifact"]
                    source = resolve_template_reference(gateway, reference)
                    copied = Path(temporary) / f"{index}{source.suffix}"
                    before = fingerprint_file(source)
                    shutil.copyfile(source, copied)
                    resolve_template_reference(gateway, reference)
                    if fingerprint_file(copied).sha256 != before.sha256:
                        raise ExecutionApiError(409, "render_image_changed", "图片复制期间内容发生变化")
                    selected.append((str(reference.get("id") or index), copied))
                verification = render_image_workbook(template, target, cells={f["cell"]: f["value"] for f in writes},
                    selected=selected, layout=layout, number_formats=number_formats)
            check = xlrd.open_workbook(target)
            try:
                sheet = check.sheet_by_name(layout["sheet"])
                for field in writes:
                    row, column = coordinate_to_tuple(field["cell"])
                    expected = "" if field["value"] is None else str(field["value"])
                    if str(sheet.cell_value(row - 1, column - 1)) != expected:
                        raise ExecutionApiError(422, "render_verification_failed", f"保存后重读不一致：{field['cell']}")
            finally:
                check.release_resources()
            return verification
        book = xlrd.open_workbook(template)
        # The existing BIFF patcher addresses one sheet. Never silently patch
        # matching coordinates in multiple sheets.
        try:
            if book.nsheets != 1 or any(f.get("sheet", book.sheet_names()[0]) != book.sheet_names()[0] for f in writes):
                raise ExecutionApiError(422, "xls_layout_unsupported", "XLS 字段渲染支持单工作表")
        finally:
            book.release_resources()
        edits = []
        for field in writes:
            row, column = coordinate_to_tuple(field["cell"])
            value = field["value"]
            kind = field.get("kind") or ("blank" if value is None else "number" if isinstance(value, (int, float)) and not isinstance(value, bool) else "text")
            edits.append(CellEdit(row - 1, column - 1, kind, value))
        patch_workbook_file(template, target, edits)
        check = xlrd.open_workbook(target)
        try:
            for field in writes:
                row, column = coordinate_to_tuple(field["cell"])
                actual = check.sheet_by_index(0).cell_value(row - 1, column - 1)
                if actual != ("" if field["value"] is None else field["value"]):
                    raise ExecutionApiError(422, "render_verification_failed", f"保存后重读不一致：{field['cell']}")
        finally:
            check.release_resources()
    else:
        from openpyxl.drawing.image import Image

        book = load_workbook(template)
        try:
            for field in writes:
                cell = book[field.get("sheet") or book.sheetnames[0]][field["cell"]]
                cell.value = field["value"]
                if isinstance(field["value"], str):
                    cell.data_type = "s"  # Literal data, including strings beginning with '='.
            for item in images:
                reference = item["artifact"]
                path = resolve_template_reference(gateway, reference)
                image = Image(path)
                image.width, image.height = item.get("width", image.width), item.get("height", image.height)
                book[item.get("sheet") or book.sheetnames[0]].add_image(image, item["cell"])
            book.save(target)
        finally:
            book.close()
        check = load_workbook(target, data_only=False)
        try:
            for field in writes:
                actual = check[field.get("sheet") or check.sheetnames[0]][field["cell"]].value
                expected = field["value"] if field["value"] != "" else None
                if actual != expected:
                    raise ExecutionApiError(422, "render_verification_failed", f"保存后重读不一致：{field['cell']}")
        finally:
            check.close()
    return {"cells_verified": len(writes), "images_written": len(images)}


def resolve_template_reference(gateway, reference):
    path = gateway.resolve(ArtifactRef(reference["root_id"], reference["relative_path"]))
    actual = fingerprint_file(path)
    if ((reference.get("sha256") and actual.sha256 != reference["sha256"])
            or (reference.get("fingerprint") and f"{actual.size}:{actual.modified_ns}" != reference["fingerprint"])):
        raise ExecutionApiError(409, "render_image_changed", "输入图片内容已改变")
    return path


def render(context):
    config, data = context.node["config"], context.input_data
    reference = config.get("template")
    if "templates" in config:
        reference = config["templates"].get(data.get("template_key"))
        if reference is None:
            raise ExecutionApiError(422, "render_template_unknown", "上游选择的模板不在节点模板集合中")
    template = resolve_template(context.db, reference)
    filename = data.get("filename") or config.get("filename") or f"record{template.suffix.lower()}"
    if (len(filename) > 200 or re.search(r'[<>:"/\\|?*\x00-\x1f]', filename)
            or Path(filename).suffix.lower() != template.suffix.lower()):
        raise ExecutionApiError(422, "render_filename_invalid", "工作簿文件名必须是有效名称，并与模板的 XLS/XLSX 格式一致")
    gateway = build_file_gateway(context.db)
    digest = canonical_sha256({"config": config, "input": data})
    relative = f"rendered/{context.run.id}/{context.node_run.id}/{digest}{template.suffix.lower()}"
    root = storage_root_by_key(context.db, config["staging_root_id"])
    existing = context.db.query(ExecutionArtifact).filter_by(node_run_id=context.node_run.id, relative_path=relative).one_or_none()
    target_ref = ArtifactRef(root.root_id, relative)
    target = gateway.resolve(target_ref, must_exist=False, for_write=True)
    if existing is not None:
        if not target.exists() or fingerprint_file(target).sha256 != existing.content_sha256:
            raise ExecutionApiError(409, "render_artifact_changed", "已生成文件缺失或被修改")
        artifact = existing
    else:
        gateway.ensure_parent(target_ref)
        with TemporaryDirectory(prefix="render-", dir=target.parent) as temporary:
            source = Path(temporary) / template.name
            # Render the bytes whose digest was selected, even if the shared
            # template is edited concurrently.
            source.write_bytes(template.read_bytes())
            if fingerprint_file(source).sha256 != reference["sha256"]:
                raise ExecutionApiError(409, "template_changed", "模板读取期间发生变化")
            rendered = Path(temporary) / f"output{template.suffix.lower()}"
            verification = render_file(source, rendered, values=data["values"], fields=config["fields"],
                                       images=data.get("images", []), gateway=gateway, layout=config.get("image_layout"), number_formats=config.get("number_formats"))
            os.replace(rendered, target)
            fsync_file(target)
        fingerprint = fingerprint_file(target)
        artifact = ExecutionArtifact(run_id=context.run.id, node_run_id=context.node_run.id, storage_root_id=root.id,
            relative_path=relative, filename=filename, role="working", immutable=True,
            media_type="application/vnd.ms-excel" if target.suffix == ".xls" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=fingerprint.size, content_sha256=fingerprint.sha256,
            metadata_json={"request_digest": digest, "verification": verification, "template": reference})
        context.db.add(artifact)
        context.db.flush()
    return {"artifact_id": artifact.id, "artifact": {"root_id": root.root_id, "relative_path": relative,
            "sha256": artifact.content_sha256}, "verification": artifact.metadata_json["verification"]}
