from __future__ import annotations

import asyncio
import hmac
import json
import logging
import mimetypes
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

from fastapi import (
    APIRouter,
    Body,
    Depends,
    Header,
    Path as ApiPath,
    Query,
    Request,
    Response,
)
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.routing import APIRoute
from fastapi.exceptions import RequestValidationError
from sqlalchemy import and_, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.config import settings
from app.database import SessionLocal, get_db
from app.execution.catalog import (
    bind_user_role,
    bootstrap_execution_system,
    create_workflow,
    export_workflow,
    get_workflow,
    publish_workflow,
    update_workflow_draft,
)
from app.execution.engine import (
    assert_idempotent_run_matches,
    can_user_handle_human_task,
    claim_human_task,
    claim_publish_node_for_api,
    complete_node,
    create_run,
    fail_node,
    reject_human_task,
    resolve_node_input,
    retry_failed_node,
    save_human_task_draft,
    set_run_control_status,
    submit_human_task,
)
from app.execution.errors import ExecutionApiError, conflict, not_found
from app.execution.events import append_audit_log, append_run_event
from app.execution.external_operations import (
    approve_prepared_external_operation,
    bridge_external_operation,
    claim_approved_external_operation,
    complete_external_attempt,
    ensure_bridge_credential_account,
    fail_external_attempt,
    heartbeat_external_attempt,
    lock_legacy_remote_business_scope,
    public_external_attempt,
    public_external_operation,
    public_external_reconciliation_context,
    reconcile_external_operation,
    record_external_attempt_stage,
    resolve_legacy_target_sample_number,
)
from app.execution.models import (
    ExecutionArtifact,
    ExecutionArtifactRelation,
    ExecutionAuditLog,
    ExecutionCategory,
    ExecutionCredential,
    ExecutionEvent,
    ExecutionExternalOperation,
    ExecutionFileIndexEntry,
    ExecutionFileMutation,
    ExecutionHumanTask,
    ExecutionIndexJob,
    ExecutionNodeRun,
    ExecutionPermission,
    ExecutionPublishReceipt,
    ExecutionRole,
    ExecutionRun,
    ExecutionSession,
    ExecutionStorageRoot,
    ExecutionUser,
    ExecutionUserRole,
    ExecutionWorkflow,
    ExecutionWorkflowVersion,
    utcnow,
)
from app.execution.persistence import (
    build_file_gateway,
    electron_groups_from_index,
    ensure_storage_roots,
    queue_refresh,
    search_index,
)
from app.execution.mutation_runtime import (
    existing_publish_result,
    plan_file_mutation,
    prepare_file_mutation,
    publish_file_mutation,
    require_publish_confirmation,
    verify_file_mutation,
    write_file_mutation,
)
from app.execution.registry import node_registry
from app.execution.regenerated_fiber import catalog_recommendations
from app.execution.electron_microscopy import (
    ELECTRON_IMAGE_SUFFIXES,
    ELECTRON_ROOT_ID,
    claim_task_snapshot_refresh,
    complete_task_snapshot_refresh,
    fail_task_snapshot_refresh,
    request_task_snapshot_refresh,
)
from app.execution.schemas import (
    CredentialUpsert,
    ExternalBridgeClaimRequest,
    ExternalBridgeCompleteRequest,
    ExternalBridgeFailRequest,
    ExternalBridgeHeartbeatRequest,
    ExternalBridgeStageRequest,
    ExternalOperationApprovalRequest,
    ExternalOperationReconciliationRequest,
    FileRefreshRequest,
    HumanTaskClaimRequest,
    HumanTaskDraftRequest,
    HumanTaskRejectRequest,
    HumanTaskSubmitRequest,
    LoginRequest,
    MutationPreflightRequest,
    MutationPublishRequest,
    MutationStageRequest,
    MutationVerifyRequest,
    MutationWriteRequest,
    NodeRetryRequest,
    RunCreate,
    TaskSnapshotBridgeClaimRequest,
    TaskSnapshotBridgeCompleteRequest,
    TaskSnapshotBridgeFailRequest,
    UserCreate,
    UserUpdate,
    WorkflowCreate,
    WorkflowDraftUpdate,
    WorkflowImportRequest,
    WorkflowPublishRequest,
    WorkflowTestRequest,
)
from app.execution.security import (
    DUMMY_PASSWORD_HASH,
    SESSION_COOKIE,
    create_session,
    clear_login_failures,
    csrf_token_for_session,
    encrypt_credential,
    has_permission,
    hash_password,
    login_throttle_key,
    record_login_failure,
    require_login_not_throttled,
    require_csrf,
    require_permission,
    resolve_session,
    verify_password,
)
from app.execution.validation import (
    definition_checksum,
    validate_definition,
    workflow_contract_checksum,
)
from app.execution.storage import ArtifactRef, StorageError


logger = logging.getLogger(__name__)


class ExecutionRoute(APIRoute):
    def get_route_handler(self) -> Callable:
        original = super().get_route_handler()

        async def handler(request: Request):
            try:
                return await original(request)
            except RequestValidationError as exc:
                error = ExecutionApiError(
                    422,
                    "request_validation_failed",
                    "请求参数校验失败",
                    details={
                        "issues": [
                            {
                                "path": ".".join(
                                    str(item)
                                    for item in issue.get("loc", ())
                                ),
                                "message": issue.get("msg", ""),
                                "type": issue.get("type", ""),
                            }
                            for issue in exc.errors()
                        ]
                    },
                )
                return JSONResponse(
                    status_code=error.status_code,
                    content=error.as_dict(),
                )
            except ExecutionApiError as exc:
                return JSONResponse(
                    status_code=exc.status_code,
                    content=exc.as_dict(),
                    headers=exc.headers,
                )
            except Exception:
                logger.exception(
                    "Unhandled execution API error",
                    extra={"path": request.url.path},
                )
                error = ExecutionApiError(
                    500,
                    "internal_error",
                    "执行系统暂时无法处理该请求，请稍后重试",
                )
                return JSONResponse(
                    status_code=error.status_code,
                    content=error.as_dict(),
                )

        return handler


router = APIRouter(
    prefix="/execution/v1",
    tags=["execution"],
    route_class=ExecutionRoute,
)


@dataclass
class AuthContext:
    session: ExecutionSession
    user: ExecutionUser


def _may_access_all_runs(
    db: Session,
    auth: AuthContext,
) -> bool:
    return has_permission(db, auth.user, "audit.read")


def _run_for_auth(
    db: Session,
    *,
    run_id: str,
    auth: AuthContext,
) -> ExecutionRun:
    run = db.get(ExecutionRun, run_id)
    if run is None or (
        run.created_by_id != auth.user.id
        and not _may_access_all_runs(db, auth)
    ):
        # Do not reveal another user's run IDs.
        raise not_found("流程运行", run_id)
    return run


def _external_operation_for_auth(
    db: Session,
    *,
    operation_id: str,
    auth: AuthContext,
) -> tuple[ExecutionExternalOperation, ExecutionRun]:
    statement = db.query(ExecutionExternalOperation).filter(
        ExecutionExternalOperation.id == operation_id
    )
    operation = statement.one_or_none()
    if operation is None:
        raise not_found("外部操作预检单", operation_id)
    run = _run_for_auth(db, run_id=operation.run_id, auth=auth)
    return operation, run


def _external_operation_for_approval(
    db: Session,
    *,
    operation_id: str,
    auth: AuthContext,
) -> tuple[
    ExecutionExternalOperation,
    ExecutionRun,
    ExecutionNodeRun,
]:
    """Lock an approval target in the engine-wide run -> node -> op order."""

    locator = (
        db.query(
            ExecutionExternalOperation.run_id,
            ExecutionExternalOperation.node_run_id,
            ExecutionExternalOperation.request_summary,
        )
        .filter(ExecutionExternalOperation.id == operation_id)
        .one_or_none()
    )
    if locator is None:
        raise not_found("外部操作预检单", operation_id)

    run = (
        db.query(ExecutionRun)
        .filter(ExecutionRun.id == locator.run_id)
        .populate_existing()
        .with_for_update()
        .one_or_none()
    )
    if run is None or (
        run.created_by_id != auth.user.id and auth.user.role != "admin"
    ):
        # audit.read may grant read access to another user's run, but it must
        # never grant authority to approve that user's external side effect.
        raise not_found("外部操作预检单", operation_id)

    node_run = (
        db.query(ExecutionNodeRun)
        .filter(
            ExecutionNodeRun.id == locator.node_run_id,
            ExecutionNodeRun.run_id == run.id,
        )
        .populate_existing()
        .with_for_update()
        .one_or_none()
    )
    if node_run is None:
        raise not_found("外部操作预检单", operation_id)

    target_sample_number = str(
        (locator.request_summary or {}).get("target_sample_number") or ""
    ).strip() or resolve_legacy_target_sample_number(run)
    remote_business_key = lock_legacy_remote_business_scope(
        db,
        sample_number=target_sample_number,
    )
    operation = (
        db.query(ExecutionExternalOperation)
        .filter(
            ExecutionExternalOperation.id == operation_id,
            ExecutionExternalOperation.run_id == run.id,
            ExecutionExternalOperation.node_run_id == node_run.id,
        )
        .populate_existing()
        .with_for_update()
        .one_or_none()
    )
    if operation is None:
        raise not_found("外部操作预检单", operation_id)
    if operation.remote_business_key != remote_business_key:
        raise conflict(
            "external_operation_run_changed",
            "预检单与当前流程运行不一致，请刷新后重试",
            operation_id=operation.id,
        )
    return operation, run, node_run


def _ensure_workflow_visible(
    db: Session,
    *,
    workflow: ExecutionWorkflow,
    auth: AuthContext,
) -> None:
    if (
        _published_workflow_capabilities(workflow).get("hidden")
        and not has_permission(db, auth.user, "workflow.design")
    ):
        raise not_found("流程", workflow.id)


def _human_task_for_auth(
    db: Session,
    *,
    task_id: str,
    auth: AuthContext,
) -> ExecutionHumanTask:
    task = db.get(ExecutionHumanTask, task_id)
    if task is None or (
        not _may_access_all_runs(db, auth)
        and not can_user_handle_human_task(
            db,
            task=task,
            user=auth.user,
        )
    ):
        raise not_found("人工任务", task_id)
    return task


def _auth_context(
    request: Request,
    db: Session = Depends(get_db),
) -> AuthContext:
    session, user = resolve_session(db, request.cookies.get(SESSION_COOKIE))
    return AuthContext(session=session, user=user)


def permission(
    permission_key: str,
    *,
    csrf: bool = False,
) -> Callable:
    def dependency(
        request: Request,
        auth: AuthContext = Depends(_auth_context),
        db: Session = Depends(get_db),
        x_csrf_token: Optional[str] = Header(default=None, alias="X-CSRF-Token"),
    ) -> AuthContext:
        require_permission(db, auth.user, permission_key)
        if csrf:
            require_csrf(auth.session, x_csrf_token)
        return auth

    dependency.execution_permission_key = permission_key
    dependency.execution_csrf_required = csrf
    return dependency


def _user_dict(
    user: ExecutionUser,
    *,
    db: Optional[Session] = None,
) -> dict[str, Any]:
    value = {
        "id": user.id,
        "username": user.username,
        "display_name": user.display_name,
        "role": user.role,
        "is_active": user.is_active,
        "created_at": user.created_at.isoformat(),
        "updated_at": user.updated_at.isoformat(),
    }
    if db is not None:
        value["permissions"] = sorted(
            {
                binding.permission.key
                for role_binding in user.role_bindings
                for binding in role_binding.role.permission_bindings
            }
        )
    return value


def _category_dict(category: ExecutionCategory) -> dict[str, Any]:
    return {
        "id": category.id,
        "key": category.key,
        "name": category.name,
        "description": category.description,
        "sort_order": category.sort_order,
        "is_active": category.is_active,
    }


def _published_workflow_capabilities(
    workflow: ExecutionWorkflow,
) -> dict[str, Any]:
    if workflow.published_version_number is None:
        return {}
    version = next(
        (
            item
            for item in workflow.versions
            if item.version_number == workflow.published_version_number
        ),
        None,
    )
    return version.capabilities or {} if version is not None else {}


def _workflow_dict(
    workflow: ExecutionWorkflow,
    *,
    include_definition: bool = False,
    include_published_definition: bool = False,
) -> dict[str, Any]:
    published_definition: dict[str, Any] = {}
    published_capabilities = _published_workflow_capabilities(workflow)
    if workflow.published_version_number is not None:
        published = next(
            (
                version
                for version in workflow.versions
                if version.version_number
                == workflow.published_version_number
            ),
            None,
        )
        if published is not None:
            published_definition = published.definition or {}
    visible_capabilities = (
        workflow.capabilities or {}
        if include_definition
        else published_capabilities
    )
    value = {
        "id": workflow.id,
        "slug": workflow.slug,
        "name": workflow.name,
        "description": workflow.description,
        "category": _category_dict(workflow.category),
        "draft_revision": workflow.draft_revision,
        "published_version": workflow.published_version_number,
        "capabilities": visible_capabilities,
        "published_capabilities": published_capabilities,
        "required_input_count": workflow.required_input_count,
        "is_enabled": workflow.is_enabled,
        "availability": {
            "available": workflow.is_enabled
            and workflow.availability_code is None,
            "code": workflow.availability_code,
            "message": workflow.availability_message,
        },
        "created_at": workflow.created_at.isoformat(),
        "updated_at": workflow.updated_at.isoformat(),
        # Runtime form contracts come only from the immutable published
        # version. Draft nodes/configuration remain designer-only.
        "input_schema": published_definition.get("input_schema") or {},
        "global_schema": published_definition.get("global_schema") or {},
    }
    if include_definition:
        value["draft_definition"] = workflow.draft_definition
        value["draft_capabilities"] = workflow.capabilities or {}
    if include_published_definition:
        # 运行准备页只读取不可变的已发布版本，绝不把管理员草稿暴露给普通用户。
        value["published_definition"] = published_definition
    return value


def _version_dict(version: ExecutionWorkflowVersion) -> dict[str, Any]:
    return {
        "id": version.id,
        "workflow_id": version.workflow_id,
        "version": version.version_number,
        "schema_version": version.schema_version,
        "checksum": version.checksum,
        "capabilities": version.capabilities or {},
        "contract_checksum": version.contract_checksum,
        "release_note": version.release_note,
        "published_by_id": version.published_by_id,
        "published_at": version.published_at.isoformat(),
    }


def _human_task_dict(task: ExecutionHumanTask) -> dict[str, Any]:
    value = {
        "id": task.id,
        "run_id": task.run_id,
        "node_id": task.node_run.node_id,
        "title": task.title,
        "description": task.description,
        "form_schema": task.form_schema or {},
        "draft_data": task.draft_data or {},
        "result_data": task.result_data or {},
        "status": task.status,
        "revision": task.revision,
        "assigned_user_id": task.assigned_user_id,
        "candidate_role": task.candidate_role_key,
        "claimed_by_id": task.claimed_by_id,
        "claimed_at": task.claimed_at.isoformat() if task.claimed_at else None,
        "completed_by_id": task.completed_by_id,
        "completed_at": task.completed_at.isoformat() if task.completed_at else None,
        "due_at": task.due_at.isoformat() if task.due_at else None,
        "created_at": task.created_at.isoformat(),
        "updated_at": task.updated_at.isoformat(),
    }
    if task.node_run is not None and task.node_run.run is not None:
        value["run"] = {
            "id": task.run_id,
            "inspection_number": task.node_run.run.inspection_number,
            "status": task.node_run.run.status,
            "workflow_id": task.node_run.run.workflow_id,
            "workflow_name": task.node_run.run.workflow.name,
        }
    return value


def _artifact_dict(artifact: ExecutionArtifact) -> dict[str, Any]:
    return {
        "id": artifact.id,
        "run_id": artifact.run_id,
        "node_run_id": artifact.node_run_id,
        "root_id": artifact.storage_root.root_id,
        "relative_path": artifact.relative_path,
        "filename": artifact.filename,
        "role": artifact.role,
        "media_type": artifact.media_type,
        "size_bytes": artifact.size_bytes,
        "content_sha256": artifact.content_sha256,
        "immutable": artifact.immutable,
        "metadata": artifact.metadata_json or {},
        "created_at": artifact.created_at.isoformat(),
    }


def _mutation_for_run_detail(
    db: Session,
    *,
    run_id: str,
    mutation_id: str,
) -> ExecutionFileMutation:
    mutation = (
        db.query(ExecutionFileMutation)
        .filter(
            ExecutionFileMutation.run_id == run_id,
            ExecutionFileMutation.mutation_id == mutation_id,
        )
        .one_or_none()
    )
    if mutation is None:
        raise not_found("文件变更", mutation_id)
    return mutation


def _mutation_publish_receipt_dict(
    db: Session,
    receipt: ExecutionPublishReceipt,
) -> dict[str, Any]:
    target_root = db.get(
        ExecutionStorageRoot,
        receipt.target_storage_root_id,
    )
    return {
        "id": receipt.id,
        "artifact_id": receipt.artifact_id,
        "target_root_id": target_root.root_id if target_root else None,
        "target_relative_path": receipt.target_relative_path,
        "content_sha256": receipt.content_sha256,
        "status": receipt.status,
        "published_by_id": receipt.published_by_id,
        "published_at": receipt.published_at.isoformat(),
        "details": receipt.details or {},
    }


def _mutation_dict(
    db: Session,
    mutation: ExecutionFileMutation,
) -> dict[str, Any]:
    source_artifact = db.get(ExecutionArtifact, mutation.source_artifact_id)
    working_artifact = (
        db.get(ExecutionArtifact, mutation.working_artifact_id)
        if mutation.working_artifact_id
        else None
    )
    receipts = (
        db.query(ExecutionPublishReceipt)
        .filter(ExecutionPublishReceipt.mutation_id == mutation.id)
        .order_by(ExecutionPublishReceipt.published_at.asc())
        .all()
    )
    return {
        "id": mutation.id,
        "mutation_id": mutation.mutation_id,
        "run_id": mutation.run_id,
        "node_run_id": mutation.node_run_id,
        "status": mutation.status,
        "source": (
            _artifact_dict(source_artifact)
            if source_artifact is not None
            else None
        ),
        "working_copy": (
            _artifact_dict(working_artifact)
            if working_artifact is not None
            else None
        ),
        "source_fingerprint": mutation.source_fingerprint,
        "change_plan": mutation.change_plan or {},
        "verification_result": mutation.verification_result or {},
        "error": (
            {
                "code": mutation.error_code,
                "message": mutation.error_message,
            }
            if mutation.error_code
            else None
        ),
        "publish_receipts": [
            _mutation_publish_receipt_dict(db, receipt)
            for receipt in receipts
        ],
        "created_at": mutation.created_at.isoformat(),
        "updated_at": mutation.updated_at.isoformat(),
    }


def _workflow_node_for_mutation(
    db: Session,
    run: ExecutionRun,
    *,
    node_type: str,
    requested_node_id: Optional[str],
    mutation_id: Optional[str],
    stage: str,
    completed_statuses: set[str],
) -> tuple[
    ExecutionRun,
    ExecutionNodeRun,
    dict[str, Any],
    dict[str, Any],
]:
    run = (
        db.query(ExecutionRun)
        .filter(ExecutionRun.id == run.id)
        .populate_existing()
        .with_for_update()
        .one()
    )
    if mutation_id is not None:
        _require_mutation_run_state(
            db,
            run=run,
            mutation_id=mutation_id,
            completed_statuses=completed_statuses,
        )
    definition_nodes = {
        str(item.get("id")): item
        for item in (run.definition_snapshot or {}).get("nodes", [])
        if isinstance(item, dict) and item.get("id")
    }
    candidate_query = db.query(ExecutionNodeRun).filter(
        ExecutionNodeRun.run_id == run.id,
        ExecutionNodeRun.node_type == node_type,
    )
    if requested_node_id is not None:
        candidate_query = candidate_query.filter(
            ExecutionNodeRun.node_id == requested_node_id
        )
    candidates = (
        candidate_query.order_by(ExecutionNodeRun.node_id.asc())
        .populate_existing()
        .with_for_update()
        .all()
    )
    if not candidates:
        raise ExecutionApiError(
            409,
            "mutation_capability_not_declared",
            "当前流程未声明该文件变更能力",
            details={"node_type": node_type},
        )
    if len(candidates) > 1:
        raise ExecutionApiError(
            409,
            "mutation_node_ambiguous",
            "流程包含多个同类型写入节点，请指定 node_id",
            details={
                "node_type": node_type,
                "node_ids": sorted(node.node_id for node in candidates),
            },
        )
    node_run = candidates[0]
    node_definition = definition_nodes.get(node_run.node_id)
    if node_definition is None:
        raise ExecutionApiError(
            409,
            "mutation_node_definition_missing",
            "运行快照缺少文件变更节点定义",
        )
    # Never join a stage already owned by a Worker/API request. In particular,
    # taking the run lock and then waiting for that owner's mutation lock would
    # invert the worker completion order (mutation -> run) on PostgreSQL.
    if node_run.status == "running":
        raise ExecutionApiError(
            409,
            (
                "mutation_publish_in_progress"
                if stage == "publish"
                else "mutation_stage_in_progress"
            ),
            "当前文件变更阶段正在由 Worker 或另一个请求执行",
            details={
                "node_id": node_run.node_id,
                "stage": stage,
            },
        )

    mutation = None
    if mutation_id is not None:
        mutation = (
            db.query(ExecutionFileMutation)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .populate_existing()
            .with_for_update()
            .one_or_none()
        )
    bound_node_run_id = (
        ((mutation.change_plan or {}).get("node_bindings") or {}).get(stage)
        if mutation is not None
        else None
    )
    if (
        stage == "publish"
        and mutation is not None
        and mutation.node_run_id is not None
        and mutation.status in {"publishing", "published", "verified"}
    ):
        bound_node_run_id = mutation.node_run_id
    reached = node_run.status == "ready"
    replay = (
        node_run.status == "succeeded"
        and mutation is not None
        and mutation.status in completed_statuses
        and bound_node_run_id == node_run.id
    )
    if not reached and not replay:
        raise ExecutionApiError(
            409,
            "mutation_node_not_reached",
            "当前文件变更节点尚未进入实际执行路径",
            details={
                "node_id": node_run.node_id,
                "node_status": node_run.status,
                "stage": stage,
            },
        )
    node_definition, input_data = resolve_node_input(
        db,
        run=run,
        node_run=node_run,
    )
    return run, node_run, node_definition, input_data


def _artifact_ref_from_value(
    value: Any,
    *,
    code: str = "mutation_node_input_invalid",
) -> ArtifactRef:
    if not isinstance(value, dict):
        raise ExecutionApiError(
            409,
            code,
            "运行节点没有生成可验证的文件引用",
        )
    try:
        return ArtifactRef(
            str(value.get("root_id") or ""),
            str(value.get("relative_path") or ""),
        )
    except (StorageError, ValueError) as exc:
        raise ExecutionApiError(
            409,
            code,
            "运行节点生成的文件引用不合法",
        ) from exc


def _require_mutation_id_matches_node(
    *,
    run: ExecutionRun,
    node_run: ExecutionNodeRun,
    input_data: dict[str, Any],
    mutation_id: str,
) -> None:
    expected = str(
        input_data.get("mutation_id")
        or f"{run.id}-{node_run.node_id}"
    )
    if mutation_id != expected:
        raise ExecutionApiError(
            409,
            "mutation_node_input_mismatch",
            "mutation_id 与当前执行节点解析出的输入不一致",
            details={"expected_mutation_id": expected},
        )


def _require_equal_node_input(
    *,
    supplied: Any,
    expected: Any,
    field: str,
) -> None:
    if supplied != expected:
        raise ExecutionApiError(
            409,
            "mutation_node_input_mismatch",
            "请求内容与当前执行节点解析出的输入不一致",
            details={"field": field},
        )


def _require_workflow_write_capability(run: ExecutionRun) -> None:
    if (run.capabilities_snapshot or {}).get("write") is not True:
        raise ExecutionApiError(
            403,
            "workflow_write_capability_required",
            "当前流程未启用受控写入能力",
        )


def _require_mutation_run_state(
    db: Session,
    *,
    run: ExecutionRun,
    mutation_id: str,
    completed_statuses: set[str],
) -> None:
    if run.status in {
        "queued",
        "running",
        "waiting_human",
        "waiting_external",
    }:
        return
    if run.status == "completed":
        mutation_status = (
            db.query(ExecutionFileMutation.status)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .scalar()
        )
        if mutation_status in completed_statuses:
            return
    raise ExecutionApiError(
        409,
        "mutation_run_not_active",
        "当前流程运行状态不允许执行新的文件变更操作",
        details={"run_status": run.status},
    )


def _require_root_slot(
    run: ExecutionRun,
    *,
    root_id: str,
    access: str,
) -> None:
    declared = {
        (str(slot.get("root_id")), str(slot.get("access") or "read"))
        for slot in (run.definition_snapshot or {}).get("root_slots", [])
        if isinstance(slot, dict) and slot.get("root_id")
    }
    if (root_id, access) not in declared:
        raise ExecutionApiError(
            422,
            "mutation_root_capability_not_declared",
            "文件根目录未以所需权限声明在运行快照中",
            details={"root_id": root_id, "required_access": access},
        )


def _require_copy_configuration(
    node_definition: dict[str, Any],
) -> Optional[str]:
    config = node_definition.get("config") or {}
    configured_root = config.get(
        "staging_root_id",
        "execution_staging",
    )
    if configured_root != "execution_staging":
        raise ExecutionApiError(
            422,
            "mutation_staging_root_mismatch",
            "工作副本根目录与流程节点配置不一致",
        )
    configured_path = config.get("working_relative_path")
    return str(configured_path) if configured_path else None


def _require_publish_configuration(
    node_definition: dict[str, Any],
    *,
    target_root_id: str,
) -> None:
    configured_root = (node_definition.get("config") or {}).get(
        "publish_root_id"
    )
    if configured_root != target_root_id:
        raise ExecutionApiError(
            422,
            "publish_root_mismatch",
            "发布目标根与流程节点配置不一致",
        )


def _workflow_definition_for_type(
    run: ExecutionRun,
    *,
    node_type: str,
) -> dict[str, Any]:
    candidates = [
        node
        for node in (run.definition_snapshot or {}).get("nodes", [])
        if isinstance(node, dict) and node.get("type") == node_type
    ]
    if len(candidates) != 1:
        raise ExecutionApiError(
            409,
            "mutation_node_ambiguous",
            "流程中的文件变更节点定义缺失或不唯一",
            details={"node_type": node_type},
        )
    return candidates[0]


def _record_mutation_api_action(
    db: Session,
    *,
    run: ExecutionRun,
    mutation_id: str,
    action: str,
    actor: ExecutionUser,
    status: str,
    details: Optional[dict[str, Any]] = None,
) -> None:
    safe_details = {"status": status, **(details or {})}
    append_audit_log(
        db,
        action=f"file_mutation.{action}",
        resource_type="execution_file_mutation",
        resource_id=mutation_id,
        actor_user_id=actor.id,
        details={"run_id": run.id, **safe_details},
    )
    append_run_event(
        db,
        run_id=run.id,
        event_type=f"file_mutation.{action}",
        actor_type="user",
        actor_id=actor.id,
        payload={"mutation_id": mutation_id, **safe_details},
    )


def _artifact_for_auth(
    db: Session,
    *,
    artifact_id: str,
    auth: AuthContext,
) -> ExecutionArtifact:
    artifact = db.get(ExecutionArtifact, artifact_id)
    if artifact is None:
        raise not_found("文件制品", artifact_id)
    if _may_access_all_runs(db, auth):
        return artifact
    run = db.get(ExecutionRun, artifact.run_id) if artifact.run_id else None
    if run is not None and run.created_by_id == auth.user.id:
        return artifact
    if run is not None:
        tasks = (
            db.query(ExecutionHumanTask)
            .filter(ExecutionHumanTask.run_id == run.id)
            .all()
        )
        if any(
            can_user_handle_human_task(db, task=task, user=auth.user)
            for task in tasks
        ):
            return artifact
    raise not_found("文件制品", artifact_id)


def _artifact_path(
    db: Session,
    artifact: ExecutionArtifact,
) -> Path:
    try:
        return build_file_gateway(db).resolve(
            ArtifactRef(
                artifact.storage_root.root_id,
                artifact.relative_path,
            ),
            expected_type="file",
        )
    except (StorageError, ValueError) as exc:
        raise ExecutionApiError(
            409,
            "artifact_unavailable",
            "文件制品已移动、变化或当前数据根不可用",
            details={"artifact_id": artifact.id},
        ) from exc


def _node_run_dict(node: ExecutionNodeRun) -> dict[str, Any]:
    return {
        "id": node.id,
        "node_id": node.node_id,
        "node_type": node.node_type,
        "node_type_version": node.node_type_version,
        "name": node.node_name,
        "status": node.status,
        "attempt_count": node.attempt_count,
        "input_data": node.input_data or {},
        "output_data": node.output_data or {},
        "error": (
            {"code": node.error_code, "message": node.error_message}
            if node.error_code
            else None
        ),
        "started_at": node.started_at.isoformat() if node.started_at else None,
        "finished_at": node.finished_at.isoformat() if node.finished_at else None,
    }


def _event_dict(event: ExecutionEvent) -> dict[str, Any]:
    return {
        "sequence": event.id,
        "event_id": event.event_id,
        "type": event.event_type,
        "actor_type": event.actor_type,
        "actor_id": event.actor_id,
        "payload": event.payload,
        "occurred_at": event.occurred_at.isoformat(),
    }


def _run_dict(
    run,
    *,
    include_definition: bool = True,
    events: Optional[list[ExecutionEvent]] = None,
    artifacts: Optional[list[ExecutionArtifact]] = None,
) -> dict[str, Any]:
    value = {
        "id": run.id,
        "workflow_id": run.workflow_id,
        "workflow_name": run.workflow.name,
        "workflow_version_id": run.workflow_version_id,
        "inspection_number": run.inspection_number,
        "mode": run.mode,
        "status": run.status,
        "capabilities": run.capabilities_snapshot or {},
        "contract_checksum": run.contract_checksum,
        "input_data": run.input_data or {},
        "global_data": run.global_data or {},
        "output_data": run.output_data or {},
        "error": (
            {"code": run.error_code, "message": run.error_message}
            if run.error_code
            else None
        ),
        "created_by_id": run.created_by_id,
        "created_by": (
            {
                "id": run.created_by.id,
                "username": run.created_by.username,
                "display_name": run.created_by.display_name,
            }
            if run.created_by is not None
            else None
        ),
        "created_at": run.created_at.isoformat(),
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "updated_at": run.updated_at.isoformat(),
        "nodes": [
            _node_run_dict(node)
            for node in sorted(run.node_runs, key=lambda item: item.created_at)
        ],
        "human_tasks": [
            _human_task_dict(node.human_task)
            for node in run.node_runs
            if node.human_task is not None
        ],
    }
    if include_definition:
        value["definition"] = run.definition_snapshot
        value["definition_checksum"] = run.definition_checksum
    if events is not None:
        value["events"] = [_event_dict(event) for event in events]
    if artifacts is not None:
        value["artifacts"] = [_artifact_dict(artifact) for artifact in artifacts]
    return value


def _run_summary_dict(run: ExecutionRun) -> dict[str, Any]:
    node_runs = list(run.node_runs)
    completed_node_statuses = {"completed", "succeeded", "skipped"}
    open_human_statuses = {"open", "pending", "claimed"}
    open_human_task_count = sum(
        1
        for node in node_runs
        if node.human_task is not None
        and node.human_task.status in open_human_statuses
    )
    return {
        "id": run.id,
        "workflow_id": run.workflow_id,
        "workflow_name": run.workflow.name,
        "workflow_version_id": run.workflow_version_id,
        "inspection_number": run.inspection_number,
        "mode": run.mode,
        "status": run.status,
        "error": (
            {"code": run.error_code, "message": run.error_message}
            if run.error_code
            else None
        ),
        "created_by_id": run.created_by_id,
        "created_by": (
            {
                "id": run.created_by.id,
                "username": run.created_by.username,
                "display_name": run.created_by.display_name,
            }
            if run.created_by is not None
            else None
        ),
        "node_progress": {
            "completed": sum(
                1 for node in node_runs if node.status in completed_node_statuses
            ),
            "total": len(node_runs),
        },
        "open_human_task_count": open_human_task_count,
        "created_at": run.created_at.isoformat(),
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "updated_at": run.updated_at.isoformat(),
    }


def _index_job_dict(job: ExecutionIndexJob, root: ExecutionStorageRoot) -> dict[str, Any]:
    return {
        "id": job.id,
        "root_id": root.root_id,
        "status": job.status,
        "max_depth": job.max_depth,
        "counts": {
            "added": job.added_count,
            "updated": job.updated_count,
            "removed": job.removed_count,
            "total": job.total_count,
        },
        "errors": job.errors or [],
        "created_at": job.created_at.isoformat(),
        "started_at": job.started_at.isoformat() if job.started_at else None,
        "finished_at": job.finished_at.isoformat() if job.finished_at else None,
    }


@router.post("/auth/login")
def login(
    payload: LoginRequest,
    request: Request,
    response: Response,
    db: Session = Depends(get_db),
):
    bootstrap_execution_system(db)
    ensure_storage_roots(db)
    db.commit()
    client_address = (
        request.client.host
        if request.client is not None
        else "direct-call"
    )
    throttle_key = login_throttle_key(payload.username, client_address)
    require_login_not_throttled(db, throttle_key)
    user = (
        db.query(ExecutionUser)
        .filter(ExecutionUser.username == payload.username.strip())
        .one_or_none()
    )
    password_valid = verify_password(
        payload.password,
        user.password_hash
        if user is not None and user.is_active
        else DUMMY_PASSWORD_HASH,
    )
    if user is None or not user.is_active or not password_valid:
        record_login_failure(db, throttle_key)
        append_audit_log(
            db,
            action="auth.login_failed",
            resource_type="execution_login_throttle",
            resource_id=throttle_key[:16],
            details={"client": client_address},
        )
        db.commit()
        raise ExecutionApiError(401, "invalid_credentials", "用户名或密码错误")
    clear_login_failures(db, throttle_key)
    session, session_token, csrf_token = create_session(db, user)
    append_audit_log(
        db,
        action="auth.login",
        resource_type="execution_session",
        resource_id=session.id,
        actor_user_id=user.id,
    )
    db.commit()
    secure = bool(getattr(settings, "EXECUTION_COOKIE_SECURE", False))
    max_age = int(getattr(settings, "EXECUTION_SESSION_TTL_HOURS", 12)) * 3600
    response.set_cookie(
        SESSION_COOKIE,
        session_token,
        max_age=max_age,
        httponly=True,
        secure=secure,
        samesite="strict",
        path="/api/execution",
    )
    return {"user": _user_dict(user, db=db), "csrf_token": csrf_token}


@router.post("/auth/logout")
def logout(
    response: Response,
    auth: AuthContext = Depends(permission("workflow.read", csrf=True)),
    db: Session = Depends(get_db),
):
    auth.session.revoked_at = utcnow()
    append_audit_log(
        db,
        action="auth.logout",
        resource_type="execution_session",
        resource_id=auth.session.id,
        actor_user_id=auth.user.id,
    )
    db.commit()
    secure = bool(getattr(settings, "EXECUTION_COOKIE_SECURE", False))
    response.delete_cookie(
        SESSION_COOKIE,
        path="/api/execution",
        httponly=True,
        secure=secure,
        samesite="strict",
    )
    return {"success": True}


@router.get("/auth/me")
def current_user(
    auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    return {"user": _user_dict(auth.user, db=db)}


@router.get("/auth/csrf")
def csrf_token(
    auth: AuthContext = Depends(permission("workflow.read")),
):
    return {"csrf_token": csrf_token_for_session(auth.session)}


@router.get("/users")
def list_users(
    _auth: AuthContext = Depends(permission("user.manage")),
    db: Session = Depends(get_db),
):
    return {
        "items": [
            _user_dict(user)
            for user in db.query(ExecutionUser)
            .order_by(ExecutionUser.created_at.asc())
            .all()
        ]
    }


@router.get("/roles")
def list_roles(
    _auth: AuthContext = Depends(permission("user.manage")),
    db: Session = Depends(get_db),
):
    return {
        "items": [
            {
                "id": role.id,
                "key": role.key,
                "name": role.name,
                "description": role.description,
                "is_system": role.is_system,
                "permissions": sorted(
                    binding.permission.key
                    for binding in role.permission_bindings
                ),
            }
            for role in db.query(ExecutionRole)
            .order_by(ExecutionRole.key.asc())
            .all()
        ],
        "permissions": [
            {
                "id": item.id,
                "key": item.key,
                "name": item.name,
                "description": item.description,
            }
            for item in db.query(ExecutionPermission)
            .order_by(ExecutionPermission.key.asc())
            .all()
        ],
    }


@router.post("/users", status_code=201)
def add_user(
    payload: UserCreate,
    auth: AuthContext = Depends(permission("user.manage", csrf=True)),
    db: Session = Depends(get_db),
):
    user = ExecutionUser(
        username=payload.username,
        display_name=payload.display_name,
        password_hash=hash_password(payload.password),
        role=payload.role,
    )
    db.add(user)
    try:
        db.flush()
        bind_user_role(db, user, payload.role, created_by_id=auth.user.id)
        append_audit_log(
            db,
            action="user.create",
            resource_type="execution_user",
            resource_id=user.id,
            actor_user_id=auth.user.id,
            details={"role": payload.role},
        )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise conflict("username_conflict", "用户名已存在") from exc
    return _user_dict(user)


@router.patch("/users/{user_id}")
def update_user(
    user_id: str,
    payload: UserUpdate,
    auth: AuthContext = Depends(permission("user.manage", csrf=True)),
    db: Session = Depends(get_db),
):
    user = db.get(ExecutionUser, user_id)
    if user is None:
        raise not_found("用户", user_id)
    changes = payload.model_dump(exclude_unset=True)
    if changes.get("is_active") is False or (
        changes.get("role") is not None and changes["role"] != "admin"
    ):
        if user.role == "admin":
            active_admins = (
                db.query(ExecutionUser)
                .filter(
                    ExecutionUser.role == "admin",
                    ExecutionUser.is_active.is_(True),
                )
                .order_by(ExecutionUser.id.asc())
                .with_for_update()
                .all()
            )
            if len(active_admins) <= 1:
                raise conflict(
                    "last_admin_protected",
                    "不能停用或降级最后一个管理员",
                )
    if "display_name" in changes:
        user.display_name = changes["display_name"]
    if "password" in changes:
        user.password_hash = hash_password(changes["password"])
    if "is_active" in changes:
        user.is_active = changes["is_active"]
    if "role" in changes:
        role = (
            db.query(ExecutionRole)
            .filter(ExecutionRole.key == changes["role"])
            .one()
        )
        (
            db.query(ExecutionUserRole)
            .filter(ExecutionUserRole.user_id == user.id)
            .delete(synchronize_session=False)
        )
        db.flush()
        db.add(
            ExecutionUserRole(
                user_id=user.id,
                role_id=role.id,
                created_by_id=auth.user.id,
            )
        )
        user.role = changes["role"]
    if {"password", "is_active", "role"} & changes.keys():
        (
            db.query(ExecutionSession)
            .filter(
                ExecutionSession.user_id == user.id,
                ExecutionSession.revoked_at.is_(None),
            )
            .update(
                {"revoked_at": utcnow()},
                synchronize_session=False,
            )
        )
    append_audit_log(
        db,
        action="user.update",
        resource_type="execution_user",
        resource_id=user.id,
        actor_user_id=auth.user.id,
        details={"fields": sorted(changes)},
    )
    db.commit()
    return _user_dict(user)


@router.get("/credentials")
def list_credentials(
    auth: AuthContext = Depends(permission("credential.manage")),
    db: Session = Depends(get_db),
):
    return {
        "items": [
            {
                "id": item.id,
                "system_key": item.system_key,
                "account_name": item.account_name,
                "configured": True,
                "secret_mask": "••••••••",
                "is_active": item.is_active,
                "updated_at": item.updated_at.isoformat(),
            }
            for item in db.query(ExecutionCredential)
            .filter(ExecutionCredential.user_id == auth.user.id)
            .order_by(ExecutionCredential.system_key.asc())
            .all()
        ]
    }


@router.put("/credentials/{system_key}")
def upsert_credential(
    system_key: str,
    payload: CredentialUpsert,
    auth: AuthContext = Depends(permission("credential.manage", csrf=True)),
    db: Session = Depends(get_db),
):
    if system_key not in {"legacy_inspection", "new_inspection"}:
        raise ExecutionApiError(422, "system_key_invalid", "未知的外部系统")
    record = (
        db.query(ExecutionCredential)
        .filter(
            ExecutionCredential.user_id == auth.user.id,
            ExecutionCredential.system_key == system_key,
        )
        .with_for_update()
        .one_or_none()
    )
    if record is None:
        record = ExecutionCredential(
            user_id=auth.user.id,
            system_key=system_key,
            encrypted_secret=encrypt_credential(payload.secret),
            revision=1,
        )
        db.add(record)
    else:
        record.encrypted_secret = encrypt_credential(payload.secret)
        record.is_active = True
        record.revision += 1
    record.account_name = payload.account_name
    db.flush()
    append_audit_log(
        db,
        action="credential.upsert",
        resource_type="execution_credential",
        resource_id=record.id,
        actor_user_id=auth.user.id,
        details={"system_key": system_key},
    )
    db.commit()
    return {
        "id": record.id,
        "system_key": record.system_key,
        "account_name": record.account_name,
        "configured": True,
        "secret_mask": "••••••••",
        "is_active": record.is_active,
        "updated_at": record.updated_at.isoformat(),
    }


@router.get("/categories")
def categories(
    _auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    return {
        "items": [
            _category_dict(category)
            for category in db.query(ExecutionCategory)
            .filter(ExecutionCategory.is_active.is_(True))
            .order_by(ExecutionCategory.sort_order.asc())
            .all()
        ]
    }


@router.get("/node-types")
def node_types(_auth: AuthContext = Depends(permission("workflow.design"))):
    return {"items": [item.public_dict() for item in node_registry.all()]}


@router.get("/catalog/recommendations")
def workflow_recommendations(
    inspection_number: str = Query(default="", max_length=200),
    preferred_categories: list[str] = Query(default=[]),
    auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    items, cache_updated = catalog_recommendations(
        db,
        inspection_number=inspection_number,
        preferred_categories=preferred_categories,
        include_hidden=has_permission(db, auth.user, "workflow.design"),
    )
    if cache_updated:
        db.commit()
    item_query_states = {
        item.get("query_state")
        for item in items
        if item.get("query_state")
    }
    if "too_many_matches" in item_query_states:
        query_state = "too_many_matches"
    elif "complete" in item_query_states or "validated" in item_query_states:
        query_state = "complete"
    elif "incomplete" in item_query_states:
        query_state = "incomplete"
    else:
        query_state = "empty"
    return {
        "inspection_number": inspection_number.strip(),
        "query_state": query_state,
        "preferred_categories": list(
            dict.fromkeys(
                category.strip()
                for category in preferred_categories
                if category.strip()
            )
        ),
        "items": items,
    }


@router.post("/workflows/import", status_code=201)
def import_workflow(
    payload: WorkflowImportRequest,
    auth: AuthContext = Depends(permission("workflow.design", csrf=True)),
    db: Session = Depends(get_db),
):
    document = payload.document
    if (
        document.get("format") != "textile-execution-workflow"
        or document.get("format_version") != "1.0"
    ):
        raise ExecutionApiError(
            422,
            "workflow_import_format_invalid",
            "不是受支持的执行系统流程文件",
        )
    definition = document.get("definition")
    metadata = document.get("workflow") or {}
    if not isinstance(definition, dict):
        raise ExecutionApiError(422, "workflow_import_invalid", "流程定义缺失")
    expected_checksum = document.get("checksum")
    actual_checksum = definition_checksum(definition)
    if expected_checksum and expected_checksum != actual_checksum:
        raise ExecutionApiError(
            422,
            "workflow_checksum_mismatch",
            "流程文件校验和不匹配",
        )
    capabilities = metadata.get("capabilities") or {}
    expected_contract_checksum = document.get("contract_checksum")
    actual_contract_checksum = workflow_contract_checksum(
        definition,
        capabilities,
    )
    if (
        expected_contract_checksum
        and expected_contract_checksum != actual_contract_checksum
    ):
        raise ExecutionApiError(
            422,
            "workflow_contract_checksum_mismatch",
            "流程文件的定义与能力声明校验和不匹配",
        )
    validation = validate_definition(definition)
    if not validation.valid:
        raise ExecutionApiError(
            422,
            "workflow_invalid",
            "导入的流程定义校验失败",
            details=validation.as_dict(),
        )
    if payload.overwrite_workflow_id:
        if payload.expected_revision is None:
            raise ExecutionApiError(
                422,
                "expected_revision_required",
                "覆盖草稿时必须提供 expected_revision",
            )
        workflow = update_workflow_draft(
            db,
            workflow_id=payload.overwrite_workflow_id,
            expected_revision=payload.expected_revision,
            definition=definition,
            actor=auth.user,
            name=metadata.get("name"),
            description=metadata.get("description"),
            capabilities=metadata.get("capabilities"),
        )
    else:
        category = (
            db.query(ExecutionCategory)
            .filter(
                ExecutionCategory.key == metadata.get("category_key")
            )
            .one_or_none()
        )
        if category is None:
            raise ExecutionApiError(
                422,
                "workflow_category_invalid",
                "流程分类不存在",
            )
        workflow = create_workflow(
            db,
            actor=auth.user,
            slug=metadata.get("slug") or f"imported-{actual_checksum[:12]}",
            category_id=category.id,
            name=metadata.get("name") or "导入流程",
            description=metadata.get("description"),
            definition=definition,
            capabilities=metadata.get("capabilities") or {},
            is_enabled=True,
        )
    db.commit()
    return _workflow_dict(workflow, include_definition=True)


@router.get("/workflows")
def workflows(
    categories: list[str] = Query(default=[]),
    query: Optional[str] = Query(default=None, max_length=200),
    auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    statement = db.query(ExecutionWorkflow).join(ExecutionCategory)
    if categories:
        statement = statement.filter(ExecutionCategory.key.in_(categories))
    if query:
        statement = statement.filter(
            (ExecutionWorkflow.name.ilike(f"%{query}%"))
            | (ExecutionWorkflow.description.ilike(f"%{query}%"))
        )
    include_draft = has_permission(db, auth.user, "workflow.design")
    items = []
    can_access_all_runs = _may_access_all_runs(db, auth)
    for item in statement.order_by(
        ExecutionCategory.sort_order.asc(),
        ExecutionWorkflow.updated_at.desc(),
    ).all():
        published_capabilities = _published_workflow_capabilities(item)
        if published_capabilities.get("system_deprecated"):
            continue
        if published_capabilities.get("hidden") and not include_draft:
            continue
        value = _workflow_dict(item, include_definition=include_draft)
        recent = (
            db.query(ExecutionRun)
            .filter(ExecutionRun.workflow_id == item.id)
            .order_by(ExecutionRun.created_at.desc())
        )
        if not can_access_all_runs:
            recent = recent.filter(
                ExecutionRun.created_by_id == auth.user.id
            )
        recent = recent.first()
        latest_run = (
            {
                "id": recent.id,
                "status": recent.status,
                "inspection_number": recent.inspection_number,
                "created_at": recent.created_at.isoformat(),
                "finished_at": (
                    recent.finished_at.isoformat()
                    if recent.finished_at
                    else None
                ),
            }
            if recent
            else None
        )
        value["latest_run"] = latest_run
        value["recent_run"] = latest_run
        items.append(value)
    return {"items": items}


@router.post("/workflows", status_code=201)
def add_workflow(
    payload: WorkflowCreate,
    auth: AuthContext = Depends(permission("workflow.design", csrf=True)),
    db: Session = Depends(get_db),
):
    workflow = create_workflow(
        db,
        actor=auth.user,
        slug=payload.slug,
        category_id=payload.category_id,
        name=payload.name,
        description=payload.description,
        definition=payload.definition,
        capabilities=payload.capabilities,
        is_enabled=payload.is_enabled,
    )
    db.commit()
    return _workflow_dict(workflow, include_definition=True)


@router.get("/workflows/{workflow_id}")
def workflow_detail(
    workflow_id: str,
    auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, workflow_id)
    _ensure_workflow_visible(db, workflow=workflow, auth=auth)
    return _workflow_dict(
        workflow,
        include_definition=has_permission(db, auth.user, "workflow.design"),
        include_published_definition=True,
    )


@router.put("/workflows/{workflow_id}/draft")
def save_workflow_draft(
    workflow_id: str,
    payload: WorkflowDraftUpdate,
    auth: AuthContext = Depends(permission("workflow.design", csrf=True)),
    db: Session = Depends(get_db),
):
    workflow = update_workflow_draft(
        db,
        workflow_id=workflow_id,
        expected_revision=payload.revision,
        definition=payload.definition,
        actor=auth.user,
        name=payload.name,
        description=payload.description,
        capabilities=payload.capabilities,
        is_enabled=payload.is_enabled,
    )
    db.commit()
    return _workflow_dict(workflow, include_definition=True)


@router.post("/workflows/{workflow_id}/validate")
def validate_workflow(
    workflow_id: str,
    definition: Optional[dict[str, Any]] = Body(default=None),
    _auth: AuthContext = Depends(permission("workflow.design", csrf=True)),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, workflow_id)
    return validate_definition(
        definition or workflow.draft_definition,
    ).as_dict()


@router.post("/workflows/{workflow_id}/publish")
def publish(
    workflow_id: str,
    payload: WorkflowPublishRequest,
    auth: AuthContext = Depends(permission("workflow.publish", csrf=True)),
    db: Session = Depends(get_db),
):
    version = publish_workflow(
        db,
        workflow_id=workflow_id,
        expected_revision=payload.revision,
        actor=auth.user,
        release_note=payload.release_note,
    )
    db.commit()
    return _version_dict(version)


@router.get("/workflows/{workflow_id}/versions")
def workflow_versions(
    workflow_id: str,
    _auth: AuthContext = Depends(permission("workflow.read")),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, workflow_id)
    return {"items": [_version_dict(version) for version in workflow.versions]}


@router.get("/workflows/{workflow_id}/versions/{version_number}")
def workflow_version_detail(
    workflow_id: str,
    version_number: int,
    _auth: AuthContext = Depends(permission("workflow.design")),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, workflow_id)
    version = (
        db.query(ExecutionWorkflowVersion)
        .filter(
            ExecutionWorkflowVersion.workflow_id == workflow.id,
            ExecutionWorkflowVersion.version_number == version_number,
        )
        .one_or_none()
    )
    if version is None:
        raise not_found("流程版本", str(version_number))
    return {**_version_dict(version), "definition": version.definition}


@router.get("/workflows/{workflow_id}/export")
def workflow_export(
    workflow_id: str,
    _auth: AuthContext = Depends(permission("workflow.design")),
    db: Session = Depends(get_db),
):
    return export_workflow(get_workflow(db, workflow_id))


@router.post("/workflows/{workflow_id}/test", status_code=201)
def test_workflow(
    workflow_id: str,
    payload: WorkflowTestRequest,
    auth: AuthContext = Depends(permission("workflow.design", csrf=True)),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, workflow_id)
    try:
        run, duplicate = create_run(
            db,
            workflow=workflow,
            actor=auth.user,
            inspection_number=payload.inspection_number,
            input_data=payload.input_data,
            global_data=payload.global_data,
            idempotency_key=payload.idempotency_key,
            mode="test",
            draft_definition=workflow.draft_definition,
            target_sample_number=payload.target_sample_number,
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        run = (
            db.query(ExecutionRun)
            .filter_by(
                created_by_id=auth.user.id,
                idempotency_key=payload.idempotency_key,
            )
            .one()
        )
        assert_idempotent_run_matches(
            run,
            workflow=workflow,
            inspection_number=payload.inspection_number.strip(),
            input_data=payload.input_data,
            global_data=payload.global_data,
            mode="test",
            definition=run.definition_snapshot,
            capabilities=run.capabilities_snapshot,
            target_sample_number=payload.target_sample_number,
        )
        duplicate = True
    return {"duplicate": duplicate, "run": _run_dict(run)}


@router.get("/runs")
def list_runs(
    status: Optional[str] = None,
    status_group: Optional[str] = Query(
        default=None,
        pattern=r"^(active|terminal)$",
    ),
    inspection_number: Optional[str] = None,
    workflow_id: Optional[str] = None,
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=200),
    auth: AuthContext = Depends(permission("workflow.run")),
    db: Session = Depends(get_db),
):
    statement = db.query(ExecutionRun)
    if not _may_access_all_runs(db, auth):
        statement = statement.filter(ExecutionRun.created_by_id == auth.user.id)
    if status:
        statement = statement.filter(ExecutionRun.status == status)
    if status_group == "active":
        statement = statement.filter(
            ExecutionRun.status.in_(
                [
                    "created",
                    "pending",
                    "queued",
                    "running",
                    "waiting_human",
                    "waiting_external",
                    "paused",
                    "cancel_pending",
                    "failure_pending",
                ]
            )
        )
    elif status_group == "terminal":
        statement = statement.filter(
            ExecutionRun.status.in_(
                ["completed", "succeeded", "failed", "cancelled"]
            )
        )
    if inspection_number:
        statement = statement.filter(
            ExecutionRun.inspection_number.ilike(f"%{inspection_number}%")
        )
    if workflow_id:
        statement = statement.filter(ExecutionRun.workflow_id == workflow_id)
    total = statement.count()
    return {
        "items": [
            _run_summary_dict(item)
            for item in statement.order_by(
                ExecutionRun.created_at.desc(),
                ExecutionRun.id.desc(),
            )
            .offset(offset)
            .limit(limit)
            .all()
        ],
        "total": total,
        "offset": offset,
        "limit": limit,
    }


@router.post("/runs", status_code=201)
def start_run(
    payload: RunCreate,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    workflow = get_workflow(db, payload.workflow_id)
    _ensure_workflow_visible(db, workflow=workflow, auth=auth)
    try:
        run, duplicate = create_run(
            db,
            workflow=workflow,
            actor=auth.user,
            inspection_number=payload.inspection_number,
            input_data=payload.input_data,
            global_data=payload.global_data,
            idempotency_key=payload.idempotency_key,
            target_sample_number=payload.target_sample_number,
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        run = (
            db.query(ExecutionRun)
            .filter_by(
                created_by_id=auth.user.id,
                idempotency_key=payload.idempotency_key,
            )
            .one()
        )
        assert_idempotent_run_matches(
            run,
            workflow=workflow,
            inspection_number=payload.inspection_number,
            input_data=payload.input_data,
            global_data=payload.global_data,
            mode="live",
            definition=run.definition_snapshot,
            capabilities=run.capabilities_snapshot,
            target_sample_number=payload.target_sample_number,
        )
        duplicate = True
    return {"duplicate": duplicate, "run": _run_dict(run)}


@router.get("/runs/{run_id}")
def run_detail(
    run_id: str,
    auth: AuthContext = Depends(permission("workflow.run")),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    event_rows = (
        db.query(ExecutionEvent)
        .filter(ExecutionEvent.run_id == run.id)
        .order_by(ExecutionEvent.id.desc())
        .limit(101)
        .all()
    )
    has_earlier_events = len(event_rows) > 100
    events = event_rows[:100]
    events.reverse()
    artifacts = (
        db.query(ExecutionArtifact)
        .filter(ExecutionArtifact.run_id == run.id)
        .order_by(ExecutionArtifact.created_at.asc())
        .all()
    )
    payload = _run_dict(run, events=events, artifacts=artifacts)
    payload["event_history"] = {
        "has_more": has_earlier_events,
        "cursors": {
            "before_id": events[0].id if events else None,
            "after_id": events[-1].id if events else None,
        },
    }
    return payload


@router.get("/runs/{run_id}/event-history")
def run_event_history(
    run_id: str,
    after_id: Optional[int] = Query(default=None, ge=0),
    before_id: Optional[int] = Query(default=None, ge=1),
    limit: int = Query(default=100, ge=1, le=200),
    auth: AuthContext = Depends(permission("workflow.run")),
    db: Session = Depends(get_db),
):
    """Return a stable, chronological page without changing the SSE contract."""

    run = _run_for_auth(db, run_id=run_id, auth=auth)
    statement = db.query(ExecutionEvent).filter(ExecutionEvent.run_id == run.id)
    if after_id is not None:
        statement = statement.filter(ExecutionEvent.id > after_id)
    if before_id is not None:
        statement = statement.filter(ExecutionEvent.id < before_id)

    if after_id is not None:
        rows = (
            statement.order_by(ExecutionEvent.id.asc())
            .limit(limit + 1)
            .all()
        )
        has_more = len(rows) > limit
        events = rows[:limit]
    else:
        rows = (
            statement.order_by(ExecutionEvent.id.desc())
            .limit(limit + 1)
            .all()
        )
        has_more = len(rows) > limit
        events = rows[:limit]
        events.reverse()

    return {
        "items": [_event_dict(event) for event in events],
        "has_more": has_more,
        "cursors": {
            "before_id": events[0].id if events else None,
            "after_id": events[-1].id if events else None,
        },
    }


@router.get("/runs/{run_id}/external-operations")
def list_run_external_operations(
    run_id: str,
    auth: AuthContext = Depends(permission("workflow.run")),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    return {
        "items": [
            public_external_operation(operation)
            for operation in (
                db.query(ExecutionExternalOperation)
                .filter(ExecutionExternalOperation.run_id == run.id)
                .order_by(
                    ExecutionExternalOperation.created_at.asc(),
                    ExecutionExternalOperation.id.asc(),
                )
                .all()
            )
        ]
    }


@router.get("/external-operations/{operation_id}")
def external_operation_detail(
    operation_id: str,
    auth: AuthContext = Depends(permission("workflow.run")),
    db: Session = Depends(get_db),
):
    operation, _run = _external_operation_for_auth(
        db,
        operation_id=operation_id,
        auth=auth,
    )
    return public_external_operation(operation)


@router.post("/external-operations/{operation_id}/approve")
def approve_external_operation(
    operation_id: str,
    payload: ExternalOperationApprovalRequest,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    operation, run, node_run = _external_operation_for_approval(
        db,
        operation_id=operation_id,
        auth=auth,
    )
    if node_run.status != "waiting_external":
        raise conflict(
            "external_operation_node_not_waiting",
            "外部操作对应节点已不再等待连接器处理",
            operation_id=operation.id,
            node_status=node_run.status,
        )
    operation, duplicate = approve_prepared_external_operation(
        db,
        operation=operation,
        run=run,
        actor=auth.user,
        payload_checksum=payload.payload_checksum,
        confirmed_sample_number=payload.confirmed_sample_number,
        note=payload.note,
    )
    db.commit()
    return {
        "duplicate": duplicate,
        "operation": public_external_operation(operation),
        "remote_write_performed": False,
    }


@router.get("/external-operations/{operation_id}/reconciliation")
def external_operation_reconciliation_context(
    operation_id: str,
    auth: AuthContext = Depends(
        permission("external_operation.reconcile")
    ),
    db: Session = Depends(get_db),
):
    if auth.user.role != "admin":
        raise ExecutionApiError(
            403,
            "external_reconciliation_admin_required",
            "只有管理员可以查看旧系统人工对账上下文",
        )
    operation, _run = _external_operation_for_auth(
        db,
        operation_id=operation_id,
        auth=auth,
    )
    return public_external_reconciliation_context(operation)


@router.post("/external-operations/{operation_id}/reconcile")
def reconcile_external_operation_result(
    operation_id: str,
    payload: ExternalOperationReconciliationRequest,
    auth: AuthContext = Depends(
        permission("external_operation.reconcile", csrf=True)
    ),
    db: Session = Depends(get_db),
):
    operation, duplicate = reconcile_external_operation(
        db,
        operation_id=operation_id,
        actor=auth.user,
        action=payload.action,
        attempt_id=payload.attempt_id,
        payload_checksum=payload.payload_checksum,
        confirmed_sample_number=payload.confirmed_sample_number,
        note=payload.note,
        evidence=payload.evidence.model_dump(mode="json"),
    )
    db.commit()
    view = public_external_operation(operation)
    return {
        "duplicate": duplicate,
        "operation": view,
        "remote_write_performed": view["remote_write_performed"],
    }


def _require_bridge_key(provided_key: Optional[str]) -> None:
    if not settings.EXECUTION_BRIDGE_ENABLED:
        raise ExecutionApiError(
            503,
            "execution_bridge_disabled",
            "执行系统 Bridge 总开关未启用",
        )
    configured = settings.EXECUTION_BRIDGE_TOKEN.strip()
    if not configured:
        raise ExecutionApiError(
            503,
            "execution_bridge_not_configured",
            "服务端未配置执行系统 Bridge 接入令牌",
        )
    if not provided_key or not hmac.compare_digest(
        provided_key.encode("utf-8"),
        configured.encode("utf-8"),
    ):
        raise ExecutionApiError(
            401,
            "execution_bridge_unauthorized",
            "Bridge 接入令牌无效",
        )


def _require_readonly_bridge_key(provided_key: Optional[str]) -> None:
    """Authenticate the read-only task probe without enabling write Bridge."""

    configured = settings.EXECUTION_BRIDGE_TOKEN.strip()
    if not configured:
        raise ExecutionApiError(
            503,
            "execution_task_snapshot_bridge_not_configured",
            "服务端未配置任务信息只读 Bridge 接入令牌",
        )
    if not provided_key or not hmac.compare_digest(
        provided_key.encode("utf-8"), configured.encode("utf-8")
    ):
        raise ExecutionApiError(
            401,
            "execution_bridge_unauthorized",
            "Bridge 接入令牌无效",
        )


@router.post("/task-snapshots/{inspection_number}/refresh", status_code=202)
def refresh_task_snapshot(
    inspection_number: str,
    force: bool = Query(default=True),
    _auth: AuthContext = Depends(permission("workflow.read", csrf=True)),
    db: Session = Depends(get_db),
):
    row, queued = request_task_snapshot_refresh(
        db, inspection_number=inspection_number, force=force
    )
    db.commit()
    return {
        "inspection_number": row.inspection_number,
        "queued": queued,
        "status": row.status,
        "revision": row.revision,
        "remote_write_performed": False,
    }


@router.post("/task-snapshot-bridge/claim")
def claim_task_snapshot_bridge(
    payload: TaskSnapshotBridgeClaimRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None, alias="X-Execution-Bridge-Key"
    ),
):
    _require_readonly_bridge_key(x_execution_bridge_key)
    row = claim_task_snapshot_refresh(db, bridge_id=payload.bridge_id)
    if row is None:
        return {"claimed": False, "remote_write_performed": False}
    response = {
        "claimed": True,
        "inspection_number": row.inspection_number,
        "claim_token": row.claim_token,
        "claim_expires_at": row.claim_expires_at.isoformat(),
        "remote_write_performed": False,
    }
    db.commit()
    return response


@router.post("/task-snapshot-bridge/{inspection_number}/complete")
def complete_task_snapshot_bridge(
    inspection_number: str,
    payload: TaskSnapshotBridgeCompleteRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None, alias="X-Execution-Bridge-Key"
    ),
):
    _require_readonly_bridge_key(x_execution_bridge_key)
    row = complete_task_snapshot_refresh(
        db,
        inspection_number=inspection_number,
        bridge_id=payload.bridge_id,
        claim_token=payload.claim_token,
        snapshot=payload.snapshot,
    )
    db.commit()
    return {
        "inspection_number": row.inspection_number,
        "status": row.status,
        "revision": row.revision,
        "expires_at": row.expires_at.isoformat(),
        "remote_write_performed": False,
    }


@router.post("/task-snapshot-bridge/{inspection_number}/fail")
def fail_task_snapshot_bridge(
    inspection_number: str,
    payload: TaskSnapshotBridgeFailRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None, alias="X-Execution-Bridge-Key"
    ),
):
    _require_readonly_bridge_key(x_execution_bridge_key)
    row = fail_task_snapshot_refresh(
        db,
        inspection_number=inspection_number,
        bridge_id=payload.bridge_id,
        claim_token=payload.claim_token,
        error_code=payload.error_code,
        message=payload.message,
    )
    db.commit()
    return {
        "inspection_number": row.inspection_number,
        "status": row.status,
        "revision": row.revision,
        "remote_write_performed": False,
    }


@router.post("/external-bridge/claim")
def claim_external_bridge_operation(
    payload: ExternalBridgeClaimRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None,
        alias="X-Execution-Bridge-Key",
    ),
):
    _require_bridge_key(x_execution_bridge_key)
    result = claim_approved_external_operation(
        db,
        bridge_id=payload.bridge_id,
        account_name=payload.account_name,
        supported_operation_types=set(payload.supported_operation_types),
    )
    if result is None:
        return {"claimed": False}
    operation, attempt, credential = result
    ensure_bridge_credential_account(
        credential,
        account_name=payload.account_name,
    )
    response = {
        "claimed": True,
        "attempt": public_external_attempt(attempt),
        "operation": bridge_external_operation(
            operation,
            credential=credential,
        ),
    }
    db.commit()
    return response


@router.post("/external-bridge/attempts/{attempt_id}/heartbeat")
def heartbeat_external_bridge_attempt(
    attempt_id: str,
    payload: ExternalBridgeHeartbeatRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None,
        alias="X-Execution-Bridge-Key",
    ),
):
    _require_bridge_key(x_execution_bridge_key)
    attempt, _operation, abort_requested = heartbeat_external_attempt(
        db,
        attempt_id=attempt_id,
        bridge_id=payload.bridge_id,
        stage=payload.stage,
        stdout_append=payload.stdout_append,
    )
    response = {
        "attempt": public_external_attempt(attempt),
        "abort_requested": abort_requested,
    }
    db.commit()
    return response


@router.post("/external-bridge/attempts/{attempt_id}/stage")
def record_external_bridge_attempt_stage(
    attempt_id: str,
    payload: ExternalBridgeStageRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None,
        alias="X-Execution-Bridge-Key",
    ),
):
    _require_bridge_key(x_execution_bridge_key)
    attempt, operation = record_external_attempt_stage(
        db,
        attempt_id=attempt_id,
        bridge_id=payload.bridge_id,
        stage=payload.stage,
        detail=payload.detail,
    )
    response = {
        "attempt": public_external_attempt(attempt),
        "abort_requested": operation.status == "cancel_pending",
    }
    db.commit()
    return response


@router.post("/external-bridge/attempts/{attempt_id}/complete")
def complete_external_bridge_attempt(
    attempt_id: str,
    payload: ExternalBridgeCompleteRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None,
        alias="X-Execution-Bridge-Key",
    ),
):
    _require_bridge_key(x_execution_bridge_key)
    operation, attempt = complete_external_attempt(
        db,
        attempt_id=attempt_id,
        bridge_id=payload.bridge_id,
        receipt=payload.receipt,
        stdout_summary=payload.stdout_summary,
    )
    response = {
        "attempt": public_external_attempt(attempt),
        "operation": public_external_operation(operation),
    }
    db.commit()
    return response


@router.post("/external-bridge/attempts/{attempt_id}/fail")
def fail_external_bridge_attempt(
    attempt_id: str,
    payload: ExternalBridgeFailRequest,
    db: Session = Depends(get_db),
    x_execution_bridge_key: Optional[str] = Header(
        default=None,
        alias="X-Execution-Bridge-Key",
    ),
):
    _require_bridge_key(x_execution_bridge_key)
    operation, attempt = fail_external_attempt(
        db,
        attempt_id=attempt_id,
        bridge_id=payload.bridge_id,
        stage=payload.stage,
        error_code=payload.error_code,
        message=payload.message,
    )
    response = {
        "attempt": public_external_attempt(attempt),
        "operation": public_external_operation(operation),
    }
    db.commit()
    return response


@router.get("/runs/{run_id}/mutations")
def list_run_mutations(
    run_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    return {
        "items": [
            _mutation_dict(db, mutation)
            for mutation in (
                db.query(ExecutionFileMutation)
                .filter(ExecutionFileMutation.run_id == run.id)
                .order_by(ExecutionFileMutation.created_at.asc())
                .all()
            )
        ]
    }


@router.get("/runs/{run_id}/mutations/{mutation_id}")
def run_mutation_detail(
    run_id: str,
    mutation_id: str = ApiPath(
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$",
        min_length=1,
        max_length=100,
    ),
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    mutation = _mutation_for_run_detail(
        db,
        run_id=run.id,
        mutation_id=mutation_id,
    )
    return {"mutation": _mutation_dict(db, mutation)}


@router.post("/runs/{run_id}/mutations/preflight", status_code=201)
def preflight_run_mutation(
    run_id: str,
    payload: MutationPreflightRequest,
    auth: AuthContext = Depends(permission("file.write", csrf=True)),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    _require_mutation_run_state(
        db,
        run=run,
        mutation_id=payload.mutation_id,
        completed_statuses={
            "planned",
            "prepared",
            "written",
            "verified",
            "published",
            "failed",
        },
    )
    _require_workflow_write_capability(run)
    run, node_run, node_definition, node_input = _workflow_node_for_mutation(
        db,
        run,
        node_type="workbook.copy",
        requested_node_id=payload.node_id,
        mutation_id=payload.mutation_id,
        stage="copy",
        completed_statuses={
            "planned",
            "prepared",
            "written",
            "verified",
            "published",
        },
    )
    _require_mutation_id_matches_node(
        run=run,
        node_run=node_run,
        input_data=node_input,
        mutation_id=payload.mutation_id,
    )
    expected_source = (
        node_input.get("source")
        or node_input.get("selected_file")
        or node_input
    )
    _require_equal_node_input(
        supplied=ArtifactRef(
            payload.source.root_id,
            payload.source.relative_path,
        ),
        expected=_artifact_ref_from_value(expected_source),
        field="source",
    )
    configured_working_path = _require_copy_configuration(node_definition)
    if (
        payload.working_relative_path is not None
        and payload.working_relative_path != configured_working_path
    ):
        raise ExecutionApiError(
            422,
            "mutation_working_path_not_declared",
            "工作副本路径必须由流程节点配置声明，不能由请求任意指定",
        )
    _require_root_slot(
        run,
        root_id=payload.source.root_id,
        access="read",
    )
    _require_root_slot(
        run,
        root_id="execution_staging",
        access="write",
    )
    try:
        result = plan_file_mutation(
            db,
            run=run,
            mutation_id=payload.mutation_id,
            source_ref=ArtifactRef(
                payload.source.root_id,
                payload.source.relative_path,
            ),
            node_run_id=node_run.id,
            working_relative_path=configured_working_path,
        )
        mutation = _mutation_for_run_detail(
            db,
            run_id=run.id,
            mutation_id=payload.mutation_id,
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation.mutation_id,
            action="preflight",
            actor=auth.user,
            status=mutation.status,
            details={
                "reused": bool(result.get("reused")),
                "source_root_id": payload.source.root_id,
            },
        )
        db.commit()
        return {
            "reused": bool(result.get("reused")),
            "preflight": result,
            "mutation": _mutation_dict(db, mutation),
        }
    except ExecutionApiError as exc:
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=payload.mutation_id,
            action="preflight_failed",
            actor=auth.user,
            status="failed",
            details={"code": exc.code},
        )
        db.commit()
        raise


@router.post("/runs/{run_id}/mutations/{mutation_id}/copy")
def copy_run_mutation(
    run_id: str,
    payload: MutationStageRequest,
    mutation_id: str = ApiPath(
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$",
        min_length=1,
        max_length=100,
    ),
    auth: AuthContext = Depends(permission("file.write", csrf=True)),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    _require_mutation_run_state(
        db,
        run=run,
        mutation_id=mutation_id,
        completed_statuses={
            "prepared",
            "written",
            "verified",
            "published",
        },
    )
    _require_workflow_write_capability(run)
    run, node_run, node_definition, node_input = _workflow_node_for_mutation(
        db,
        run,
        node_type="workbook.copy",
        requested_node_id=payload.node_id,
        mutation_id=mutation_id,
        stage="copy",
        completed_statuses={
            "prepared",
            "written",
            "verified",
            "published",
        },
    )
    _require_mutation_id_matches_node(
        run=run,
        node_run=node_run,
        input_data=node_input,
        mutation_id=mutation_id,
    )
    _require_copy_configuration(node_definition)
    _require_root_slot(
        run,
        root_id="execution_staging",
        access="write",
    )
    try:
        receipt = prepare_file_mutation(
            db,
            run=run,
            mutation_id=mutation_id,
            node_run_id=node_run.id,
        )
        mutation = _mutation_for_run_detail(
            db,
            run_id=run.id,
            mutation_id=mutation_id,
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="copy",
            actor=auth.user,
            status=mutation.status,
            details={
                "reused": bool(receipt.get("reused")),
                "working_content_sha256": (
                    (receipt.get("working_copy") or {})
                    .get("fingerprint", {})
                    .get("sha256")
                ),
            },
        )
        db.commit()
        return {
            "reused": bool(receipt.get("reused")),
            "receipt": receipt,
            "mutation": _mutation_dict(db, mutation),
        }
    except ExecutionApiError as exc:
        mutation = (
            db.query(ExecutionFileMutation)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .one_or_none()
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="copy_failed",
            actor=auth.user,
            status=mutation.status if mutation is not None else "failed",
            details={"code": exc.code},
        )
        db.commit()
        raise


@router.post("/runs/{run_id}/mutations/{mutation_id}/write")
def write_run_mutation(
    run_id: str,
    payload: MutationWriteRequest,
    mutation_id: str = ApiPath(
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$",
        min_length=1,
        max_length=100,
    ),
    auth: AuthContext = Depends(permission("file.write", csrf=True)),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    _require_mutation_run_state(
        db,
        run=run,
        mutation_id=mutation_id,
        completed_statuses={"written", "verified", "published"},
    )
    _require_workflow_write_capability(run)
    run, node_run, node_definition, node_input = _workflow_node_for_mutation(
        db,
        run,
        node_type="workbook.write_cells",
        requested_node_id=payload.node_id,
        mutation_id=mutation_id,
        stage="write",
        completed_statuses={"written", "verified", "published"},
    )
    _require_mutation_id_matches_node(
        run=run,
        node_run=node_run,
        input_data=node_input,
        mutation_id=mutation_id,
    )
    supplied_writes = [item.model_dump() for item in payload.writes]
    expected_writes = (
        node_input.get("writes")
        or (node_definition.get("config") or {}).get("writes")
    )
    _require_equal_node_input(
        supplied=supplied_writes,
        expected=expected_writes,
        field="writes",
    )
    _require_root_slot(
        run,
        root_id="execution_staging",
        access="write",
    )
    try:
        receipt = write_file_mutation(
            db,
            run=run,
            mutation_id=mutation_id,
            writes=supplied_writes,
            node_run_id=node_run.id,
        )
        mutation = _mutation_for_run_detail(
            db,
            run_id=run.id,
            mutation_id=mutation_id,
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="write",
            actor=auth.user,
            status=mutation.status,
            details={
                "reused": bool(receipt.get("reused")),
                "write_count": len(payload.writes),
                "working_content_sha256": (
                    (receipt.get("working_copy") or {})
                    .get("after_fingerprint", {})
                    .get("sha256")
                ),
            },
        )
        db.commit()
        return {
            "reused": bool(receipt.get("reused")),
            "receipt": receipt,
            "mutation": _mutation_dict(db, mutation),
        }
    except ExecutionApiError as exc:
        mutation = (
            db.query(ExecutionFileMutation)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .one_or_none()
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="write_failed",
            actor=auth.user,
            status=mutation.status if mutation is not None else "failed",
            details={"code": exc.code},
        )
        db.commit()
        raise


@router.post("/runs/{run_id}/mutations/{mutation_id}/verify")
def verify_run_mutation(
    run_id: str,
    payload: MutationVerifyRequest,
    mutation_id: str = ApiPath(
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$",
        min_length=1,
        max_length=100,
    ),
    auth: AuthContext = Depends(permission("file.write", csrf=True)),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    _require_mutation_run_state(
        db,
        run=run,
        mutation_id=mutation_id,
        completed_statuses={"verified", "published"},
    )
    _require_workflow_write_capability(run)
    run, node_run, node_definition, node_input = _workflow_node_for_mutation(
        db,
        run,
        node_type="workbook.verify",
        requested_node_id=payload.node_id,
        mutation_id=mutation_id,
        stage="verify",
        completed_statuses={"verified", "published"},
    )
    _require_mutation_id_matches_node(
        run=run,
        node_run=node_run,
        input_data=node_input,
        mutation_id=mutation_id,
    )
    expected_target = node_input.get("target")
    _require_equal_node_input(
        supplied=ArtifactRef(
            payload.target.root_id,
            payload.target.relative_path,
        ),
        expected=_artifact_ref_from_value(expected_target),
        field="target",
    )
    expected_writes = (
        node_input.get("writes")
        or (node_definition.get("config") or {}).get("writes")
    )
    supplied_writes = (
        [item.model_dump() for item in payload.writes]
        if payload.writes is not None
        else expected_writes
    )
    _require_equal_node_input(
        supplied=supplied_writes,
        expected=expected_writes,
        field="writes",
    )
    _require_root_slot(
        run,
        root_id="execution_staging",
        access="write",
    )
    _require_root_slot(
        run,
        root_id=payload.target.root_id,
        access="publish",
    )
    publish_definition = _workflow_definition_for_type(
        run,
        node_type="artifact.publish",
    )
    _require_publish_configuration(
        publish_definition,
        target_root_id=payload.target.root_id,
    )
    try:
        result = verify_file_mutation(
            db,
            run=run,
            mutation_id=mutation_id,
            target_ref=ArtifactRef(
                payload.target.root_id,
                payload.target.relative_path,
            ),
            writes=supplied_writes,
            node_run_id=node_run.id,
        )
        mutation = _mutation_for_run_detail(
            db,
            run_id=run.id,
            mutation_id=mutation_id,
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="verify",
            actor=auth.user,
            status=mutation.status,
            details={
                "verified": bool(result.get("verified")),
                "checked_cells": result.get("checked_cells"),
                "working_content_sha256": result.get(
                    "working_content_sha256"
                ),
                "target_root_id": payload.target.root_id,
            },
        )
        db.commit()
        return {
            "verification": result,
            "mutation": _mutation_dict(db, mutation),
        }
    except ExecutionApiError as exc:
        mutation = (
            db.query(ExecutionFileMutation)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .one_or_none()
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="verify_failed",
            actor=auth.user,
            status=mutation.status if mutation is not None else "failed",
            details={"code": exc.code},
        )
        db.commit()
        raise


@router.post("/runs/{run_id}/mutations/{mutation_id}/publish")
def publish_run_mutation(
    run_id: str,
    payload: MutationPublishRequest,
    mutation_id: str = ApiPath(
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$",
        min_length=1,
        max_length=100,
    ),
    auth: AuthContext = Depends(permission("file.publish", csrf=True)),
    db: Session = Depends(get_db),
):
    run = _run_for_auth(db, run_id=run_id, auth=auth)
    _require_mutation_run_state(
        db,
        run=run,
        mutation_id=mutation_id,
        completed_statuses=(
            {"verified", "published"}
            if run.mode == "test"
            else {"published"}
        ),
    )
    _require_workflow_write_capability(run)
    run, node_run, node_definition, node_input = _workflow_node_for_mutation(
        db,
        run,
        node_type="artifact.publish",
        requested_node_id=payload.node_id,
        mutation_id=mutation_id,
        stage="publish",
        completed_statuses={"verified", "publishing", "published"},
    )
    _require_mutation_id_matches_node(
        run=run,
        node_run=node_run,
        input_data=node_input,
        mutation_id=mutation_id,
    )
    target_ref = ArtifactRef(
        payload.target.root_id,
        payload.target.relative_path,
    )
    _require_equal_node_input(
        supplied=target_ref,
        expected=_artifact_ref_from_value(node_input.get("target")),
        field="target",
    )
    _require_root_slot(
        run,
        root_id="execution_staging",
        access="write",
    )
    _require_root_slot(
        run,
        root_id=payload.target.root_id,
        access="publish",
    )
    _require_publish_configuration(
        node_definition,
        target_root_id=payload.target.root_id,
    )
    approval_context = payload.approval_context.model_dump()
    lease_token: Optional[str] = None
    try:
        # Validate the completed task before claiming the publish node. Invalid
        # or stale confirmation data must not turn a ready node into a failed
        # run.
        require_publish_confirmation(
            db,
            run=run,
            publish_node=node_definition,
            mutation_id=mutation_id,
            target_ref=target_ref,
            approval_context=approval_context,
        )
        run, node_run, lease_token, replay = claim_publish_node_for_api(
            db,
            run_id=run.id,
            node_run_id=node_run.id,
            mutation_id=mutation_id,
            actor=auth.user,
        )
        if replay:
            receipt = existing_publish_result(
                db,
                run=run,
                mutation_id=mutation_id,
            )
        else:
            receipt = publish_file_mutation(
                db,
                run=run,
                mutation_id=mutation_id,
                target_ref=target_ref,
                approval_context=approval_context,
                node_run_id=node_run.id,
                lease_token=lease_token,
                publish_node=node_definition,
                expected_working_ref=_artifact_ref_from_value(
                    node_input.get("working_copy")
                    or node_input.get("working_ref")
                ),
            )
            complete_node(
                db,
                node_run_id=node_run.id,
                lease_token=lease_token,
                output_data=receipt,
            )
        mutation = _mutation_for_run_detail(
            db,
            run_id=run.id,
            mutation_id=mutation_id,
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="publish",
            actor=auth.user,
            status=mutation.status,
            details={
                "dry_run": bool(receipt.get("dry_run")),
                "reused": bool(receipt.get("reused")),
                "target_root_id": payload.target.root_id,
                "content_sha256": (
                    receipt.get("content_sha256")
                    or (receipt.get("published") or {})
                    .get("fingerprint", {})
                    .get("sha256")
                ),
            },
        )
        db.commit()
        return {
            "reused": bool(receipt.get("reused")),
            "receipt": receipt,
            "mutation": _mutation_dict(db, mutation),
        }
    except ExecutionApiError as exc:
        if lease_token is not None:
            try:
                fail_node(
                    db,
                    node_run_id=node_run.id,
                    lease_token=lease_token,
                    error_code=exc.code,
                    error_message=exc.message,
                )
            except ExecutionApiError:
                # A lease winner or cancellation reconciliation now owns the
                # node. The durable mutation/receipt remains authoritative.
                pass
        mutation = (
            db.query(ExecutionFileMutation)
            .filter(
                ExecutionFileMutation.run_id == run.id,
                ExecutionFileMutation.mutation_id == mutation_id,
            )
            .one_or_none()
        )
        _record_mutation_api_action(
            db,
            run=run,
            mutation_id=mutation_id,
            action="publish_failed",
            actor=auth.user,
            status=mutation.status if mutation is not None else "failed",
            details={"code": exc.code},
        )
        db.commit()
        raise


def _run_action(
    run_id: str,
    action: str,
    auth: AuthContext,
    db: Session,
):
    _run_for_auth(db, run_id=run_id, auth=auth)
    run = set_run_control_status(
        db,
        run_id=run_id,
        action=action,
        actor=auth.user,
    )
    db.commit()
    return _run_dict(run)


@router.post("/runs/{run_id}/pause")
def pause_run(
    run_id: str,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    return _run_action(run_id, "pause", auth, db)


@router.post("/runs/{run_id}/resume")
def resume_run(
    run_id: str,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    return _run_action(run_id, "resume", auth, db)


@router.post("/runs/{run_id}/cancel")
def cancel_run(
    run_id: str,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    return _run_action(run_id, "cancel", auth, db)


@router.post("/runs/{run_id}/nodes/{node_id}/retry")
def retry_node(
    run_id: str,
    node_id: str,
    payload: NodeRetryRequest,
    auth: AuthContext = Depends(permission("workflow.run", csrf=True)),
    db: Session = Depends(get_db),
):
    _run_for_auth(db, run_id=run_id, auth=auth)
    node = retry_failed_node(
        db,
        run_id=run_id,
        node_id=node_id,
        actor=auth.user,
        reason=payload.reason,
    )
    db.commit()
    return _node_run_dict(node)


@router.get("/runs/{run_id}/events")
async def run_events(
    run_id: str,
    request: Request,
    last_event_id: Optional[str] = Header(default=None, alias="Last-Event-ID"),
):
    raw_session_token = request.cookies.get(SESSION_COOKIE)
    with SessionLocal() as auth_db:
        session, user = resolve_session(auth_db, raw_session_token)
        require_permission(auth_db, user, "workflow.run")
        _run_for_auth(
            auth_db,
            run_id=run_id,
            auth=AuthContext(session=session, user=user),
        )
        auth_db.commit()
    try:
        after = max(int(last_event_id or 0), 0)
    except ValueError as exc:
        raise ExecutionApiError(400, "event_cursor_invalid", "事件游标不合法") from exc

    async def stream():
        cursor = after
        started_at = asyncio.get_running_loop().time()
        max_seconds = max(
            int(getattr(settings, "EXECUTION_SSE_MAX_SECONDS", 300)),
            30,
        )
        while True:
            if await request.is_disconnected():
                break
            if asyncio.get_running_loop().time() - started_at >= max_seconds:
                yield (
                    "event: stream.rotate\n"
                    'data: {"reason":"connection_lifetime_reached"}\n\n'
                )
                break
            with SessionLocal() as event_db:
                try:
                    live_session, live_user = resolve_session(
                        event_db,
                        raw_session_token,
                        touch=False,
                    )
                    require_permission(event_db, live_user, "workflow.run")
                    _run_for_auth(
                        event_db,
                        run_id=run_id,
                        auth=AuthContext(
                            session=live_session,
                            user=live_user,
                        ),
                    )
                except ExecutionApiError as exc:
                    yield (
                        "event: stream.closed\n"
                        f"data: {json.dumps({'code': exc.code}, ensure_ascii=False)}\n\n"
                    )
                    break
                events = (
                    event_db.query(ExecutionEvent)
                    .filter(
                        ExecutionEvent.run_id == run_id,
                        ExecutionEvent.id > cursor,
                    )
                    .order_by(ExecutionEvent.id.asc())
                    .limit(100)
                    .all()
                )
                terminal = False
                if events:
                    for event in events:
                        cursor = event.id
                        payload = {
                            "sequence": event.id,
                            "event_id": event.event_id,
                            "type": event.event_type,
                            "occurred_at": event.occurred_at.isoformat(),
                            "payload": event.payload,
                        }
                        yield (
                            f"id: {event.id}\n"
                            f"event: {event.event_type}\n"
                            f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"
                        )
                        if event.event_type in {
                            "run.completed",
                            "run.failed",
                            "run.cancelled",
                        }:
                            terminal = True
                    if terminal:
                        break
                else:
                    run_status = event_db.query(ExecutionRun.status).filter(
                        ExecutionRun.id == run_id
                    ).scalar()
                    if run_status in {"completed", "failed", "cancelled"}:
                        yield (
                            "event: stream.closed\n"
                            f"data: {json.dumps({'status': run_status})}\n\n"
                        )
                        break
                    yield ": keepalive\n\n"
            await asyncio.sleep(1.0)

    return StreamingResponse(
        stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/human-tasks")
def human_tasks(
    status: Optional[str] = None,
    mine: bool = False,
    auth: AuthContext = Depends(permission("human_task.handle")),
    db: Session = Depends(get_db),
):
    statement = db.query(ExecutionHumanTask)
    if not _may_access_all_runs(db, auth):
        role_keys = [
            binding.role.key
            for binding in auth.user.role_bindings
        ]
        statement = statement.join(
            ExecutionRun,
            ExecutionRun.id == ExecutionHumanTask.run_id,
        ).filter(
            or_(
                ExecutionHumanTask.assigned_user_id == auth.user.id,
                ExecutionHumanTask.claimed_by_id == auth.user.id,
                (
                    ExecutionHumanTask.candidate_role_key.in_(role_keys)
                    if role_keys
                    else False
                ),
                and_(
                    ExecutionHumanTask.assigned_user_id.is_(None),
                    ExecutionHumanTask.candidate_role_key.is_(None),
                    ExecutionRun.created_by_id == auth.user.id,
                ),
            )
        )
    if status:
        statement = statement.filter(ExecutionHumanTask.status == status)
    if mine:
        statement = statement.filter(
            ExecutionHumanTask.claimed_by_id == auth.user.id
        )
    return {
        "items": [
            _human_task_dict(task)
            for task in statement.order_by(ExecutionHumanTask.created_at.desc()).all()
        ]
    }


@router.get("/human-tasks/{task_id}")
def human_task_detail(
    task_id: str,
    auth: AuthContext = Depends(permission("human_task.handle")),
    db: Session = Depends(get_db),
):
    task = _human_task_for_auth(db, task_id=task_id, auth=auth)
    node_run = task.node_run
    run = node_run.run
    return {
        "task": _human_task_dict(task),
        "node_run": {
            "id": node_run.id,
            "node_id": node_run.node_id,
            "node_type": node_run.node_type,
            "name": node_run.node_name,
            "status": node_run.status,
            # This is the server-persisted candidate/form context. The submit
            # endpoint still canonicalizes candidate IDs against this snapshot.
            "input_data": node_run.input_data or {},
        },
        "run": {
            "id": run.id,
            "inspection_number": run.inspection_number,
            "status": run.status,
            "mode": run.mode,
            "created_by_id": run.created_by_id,
            "created_at": run.created_at.isoformat(),
        },
        "workflow": {
            "id": run.workflow.id,
            "name": run.workflow.name,
            "category": _category_dict(run.workflow.category),
            "version_id": run.workflow_version_id,
        },
    }


@router.post("/human-tasks/{task_id}/claim")
def claim_task(
    task_id: str,
    payload: HumanTaskClaimRequest,
    auth: AuthContext = Depends(permission("human_task.handle", csrf=True)),
    db: Session = Depends(get_db),
):
    _human_task_for_auth(db, task_id=task_id, auth=auth)
    task = claim_human_task(
        db,
        task_id=task_id,
        expected_revision=payload.revision,
        actor=auth.user,
    )
    db.commit()
    return _human_task_dict(task)


@router.put("/human-tasks/{task_id}/draft")
def save_task_draft(
    task_id: str,
    payload: HumanTaskDraftRequest,
    auth: AuthContext = Depends(permission("human_task.handle", csrf=True)),
    db: Session = Depends(get_db),
):
    _human_task_for_auth(db, task_id=task_id, auth=auth)
    task = save_human_task_draft(
        db,
        task_id=task_id,
        expected_revision=payload.revision,
        data=payload.data,
        actor=auth.user,
    )
    db.commit()
    return _human_task_dict(task)


@router.post("/human-tasks/{task_id}/submit")
def submit_task(
    task_id: str,
    payload: HumanTaskSubmitRequest,
    auth: AuthContext = Depends(permission("human_task.handle", csrf=True)),
    db: Session = Depends(get_db),
):
    _human_task_for_auth(db, task_id=task_id, auth=auth)
    task = submit_human_task(
        db,
        task_id=task_id,
        expected_revision=payload.revision,
        data=payload.data,
        actor=auth.user,
    )
    db.commit()
    return _human_task_dict(task)


@router.post("/human-tasks/{task_id}/reject")
def reject_task(
    task_id: str,
    payload: HumanTaskRejectRequest,
    auth: AuthContext = Depends(permission("human_task.handle", csrf=True)),
    db: Session = Depends(get_db),
):
    _human_task_for_auth(db, task_id=task_id, auth=auth)
    task = reject_human_task(
        db,
        task_id=task_id,
        expected_revision=payload.revision,
        reason=payload.reason,
        actor=auth.user,
    )
    db.commit()
    return _human_task_dict(task)


@router.get("/artifacts/{artifact_id}")
def artifact_detail(
    artifact_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    artifact = _artifact_for_auth(db, artifact_id=artifact_id, auth=auth)
    parent_relations = (
        db.query(ExecutionArtifactRelation)
        .filter(ExecutionArtifactRelation.child_artifact_id == artifact.id)
        .all()
    )
    child_relations = (
        db.query(ExecutionArtifactRelation)
        .filter(ExecutionArtifactRelation.parent_artifact_id == artifact.id)
        .all()
    )
    return {
        "artifact": _artifact_dict(artifact),
        "lineage": {
            "parents": [
                {
                    "artifact_id": relation.parent_artifact_id,
                    "relation_type": relation.relation_type,
                }
                for relation in parent_relations
            ],
            "children": [
                {
                    "artifact_id": relation.child_artifact_id,
                    "relation_type": relation.relation_type,
                }
                for relation in child_relations
            ],
        },
    }


@router.get("/artifacts/{artifact_id}/download")
def download_artifact(
    artifact_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    artifact = _artifact_for_auth(db, artifact_id=artifact_id, auth=auth)
    path = _artifact_path(db, artifact)
    append_audit_log(
        db,
        action="artifact.download",
        resource_type="execution_artifact",
        resource_id=artifact.id,
        actor_user_id=auth.user.id,
    )
    db.commit()
    media_type = artifact.media_type or mimetypes.guess_type(artifact.filename)[0]
    return FileResponse(
        path,
        filename=artifact.filename,
        media_type=media_type or "application/octet-stream",
    )


def _preview_cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


_INLINE_RASTER_MEDIA_TYPES = {
    ".bmp": "image/bmp",
    ".gif": "image/gif",
    ".jpeg": "image/jpeg",
    ".jpg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".webp": "image/webp",
}
_WORKBOOK_PREVIEW_MAX_ENTRIES = 5_000
_WORKBOOK_PREVIEW_MAX_ENTRY_BYTES = 50 * 1024 * 1024
_WORKBOOK_PREVIEW_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_WORKBOOK_PREVIEW_MAX_COMPRESSION_RATIO = 200


def _inline_raster_media_type(extension: str) -> Optional[str]:
    """Only allow browser-safe raster formats to be rendered inline."""

    return _INLINE_RASTER_MEDIA_TYPES.get(extension.lower())


def _validate_workbook_preview_archive(path: Path) -> None:
    """Reject encrypted or suspiciously amplified OOXML archives."""

    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > _WORKBOOK_PREVIEW_MAX_ENTRIES:
                raise ExecutionApiError(
                    413,
                    "workbook_preview_archive_too_large",
                    "工作簿内部文件数量过多，不能在线预览",
                )

            total_uncompressed = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ExecutionApiError(
                        422,
                        "workbook_preview_encrypted",
                        "加密工作簿不能在线预览",
                    )
                if entry.file_size > _WORKBOOK_PREVIEW_MAX_ENTRY_BYTES:
                    raise ExecutionApiError(
                        413,
                        "workbook_preview_archive_too_large",
                        "工作簿内部单个文件过大，不能在线预览",
                    )
                total_uncompressed += entry.file_size
                if total_uncompressed > _WORKBOOK_PREVIEW_MAX_UNCOMPRESSED_BYTES:
                    raise ExecutionApiError(
                        413,
                        "workbook_preview_archive_too_large",
                        "工作簿解压后体积过大，不能在线预览",
                    )
                if (
                    entry.file_size > 10 * 1024 * 1024
                    and entry.file_size
                    > max(entry.compress_size, 1)
                    * _WORKBOOK_PREVIEW_MAX_COMPRESSION_RATIO
                ):
                    raise ExecutionApiError(
                        413,
                        "workbook_preview_archive_ratio_exceeded",
                        "工作簿压缩比异常，不能在线预览",
                    )
    except ExecutionApiError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise ExecutionApiError(
            422,
            "workbook_preview_invalid_archive",
            "工作簿文件结构无效，不能在线预览",
        ) from exc


@router.get("/artifacts/{artifact_id}/preview")
def preview_artifact(
    artifact_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    artifact = _artifact_for_auth(db, artifact_id=artifact_id, auth=auth)
    path = _artifact_path(db, artifact)
    if path.stat().st_size > 30 * 1024 * 1024:
        raise ExecutionApiError(
            413,
            "artifact_preview_too_large",
            "文件超过 30 MiB，不能在线预览，请下载后查看",
        )
    extension = path.suffix.lower()
    inline_media_type = _inline_raster_media_type(extension)
    if inline_media_type is not None:
        return FileResponse(
            path,
            media_type=inline_media_type,
            headers={
                "Content-Disposition": "inline",
                "Content-Security-Policy": "default-src 'none'; sandbox",
                "X-Content-Type-Options": "nosniff",
            },
        )
    if extension in {".txt", ".csv", ".json", ".log"}:
        raw = path.read_bytes()
        truncated = len(raw) > 256 * 1024
        return {
            "kind": "text",
            "content": raw[: 256 * 1024].decode("utf-8", errors="replace"),
            "truncated": truncated,
            "encoding": "utf-8",
        }
    if extension in {".xlsx", ".xlsm"}:
        _validate_workbook_preview_archive(path)
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
                keep_vba=False,
                keep_links=False,
            )
            try:
                sheets = []
                for sheet in workbook.worksheets[:10]:
                    rows = [
                        [_preview_cell_value(value) for value in row[:30]]
                        for row in sheet.iter_rows(
                            min_row=1,
                            max_row=min(sheet.max_row or 1, 100),
                            max_col=min(sheet.max_column or 1, 30),
                            values_only=True,
                        )
                    ]
                    sheets.append(
                        {
                            "name": sheet.title,
                            "rows": rows,
                            "truncated": (
                                (sheet.max_row or 0) > 100
                                or (sheet.max_column or 0) > 30
                            ),
                        }
                    )
            finally:
                workbook.close()
        except Exception as exc:
            raise ExecutionApiError(
                422,
                "workbook_preview_failed",
                "工作簿无法以只读模式安全预览",
            ) from exc
        return {"kind": "workbook", "sheets": sheets}
    raise ExecutionApiError(
        415,
        "artifact_preview_unsupported",
        "该文件类型暂不支持在线预览，请下载后查看",
        details={"extension": extension},
    )


@router.get("/publish-receipts/{receipt_id}")
def publish_receipt_detail(
    receipt_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    receipt = db.get(ExecutionPublishReceipt, receipt_id)
    if receipt is None:
        raise not_found("发布回执", receipt_id)
    artifact = _artifact_for_auth(
        db,
        artifact_id=receipt.artifact_id,
        auth=auth,
    )
    mutation = db.get(ExecutionFileMutation, receipt.mutation_id)
    target_root = db.get(
        ExecutionStorageRoot,
        receipt.target_storage_root_id,
    )
    return {
        "id": receipt.id,
        "mutation_id": mutation.mutation_id if mutation else None,
        "artifact": _artifact_dict(artifact),
        "target": {
            "root_id": target_root.root_id if target_root else None,
            "relative_path": receipt.target_relative_path,
        },
        "content_sha256": receipt.content_sha256,
        "status": receipt.status,
        "published_by_id": receipt.published_by_id,
        "published_at": receipt.published_at.isoformat(),
        "details": receipt.details or {},
    }


@router.get("/audit")
def audit_records(
    action: Optional[str] = Query(default=None, max_length=100),
    resource_type: Optional[str] = Query(default=None, max_length=100),
    resource_id: Optional[str] = Query(default=None, max_length=100),
    actor_user_id: Optional[str] = Query(default=None, max_length=36),
    before_id: Optional[int] = Query(default=None, ge=1),
    limit: int = Query(default=100, ge=1, le=200),
    _auth: AuthContext = Depends(permission("audit.read")),
    db: Session = Depends(get_db),
):
    statement = db.query(ExecutionAuditLog)
    if action:
        statement = statement.filter(ExecutionAuditLog.action == action)
    if resource_type:
        statement = statement.filter(
            ExecutionAuditLog.resource_type == resource_type
        )
    if resource_id:
        statement = statement.filter(
            ExecutionAuditLog.resource_id == resource_id
        )
    if actor_user_id:
        statement = statement.filter(
            ExecutionAuditLog.actor_user_id == actor_user_id
        )
    if before_id:
        statement = statement.filter(ExecutionAuditLog.id < before_id)
    records = statement.order_by(ExecutionAuditLog.id.desc()).limit(limit).all()
    return {
        "items": [
            {
                "id": record.id,
                "actor_user_id": record.actor_user_id,
                "action": record.action,
                "resource_type": record.resource_type,
                "resource_id": record.resource_id,
                "request_id": record.request_id,
                "details": record.details or {},
                "created_at": record.created_at.isoformat(),
            }
            for record in records
        ],
        "next_before_id": records[-1].id if len(records) == limit else None,
    }


@router.get("/files/roots")
def storage_roots(
    _auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    return {
        "items": [
            {
                "root_id": root.root_id,
                "name": root.name,
                "access_mode": root.access_mode,
                "category": root.category_key,
                "is_available": root.is_available,
                "availability_message": root.availability_message,
                "last_scan_finished_at": (
                    root.last_scan_finished_at.isoformat()
                    if root.last_scan_finished_at
                    else None
                ),
                "last_scan_error": root.last_scan_error,
            }
            for root in db.query(ExecutionStorageRoot)
            .filter(
                ExecutionStorageRoot.is_active.is_(True),
                ExecutionStorageRoot.access_mode == "read",
            )
            .order_by(ExecutionStorageRoot.name.asc())
            .all()
        ]
    }


@router.get("/files/search")
def search_files(
    inspection_number: str = Query(min_length=1, max_length=200),
    root_ids: list[str] = Query(default=[]),
    categories: list[str] = Query(default=[]),
    recent_days: Optional[int] = Query(default=7, ge=0, le=3650),
    limit: int = Query(default=6, ge=1, le=100),
    include_electron_groups: bool = False,
    _auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    items = search_index(
        db,
        inspection_number=inspection_number,
        root_ids=root_ids or None,
        category_keys=categories or None,
        recent_days=recent_days,
        limit=limit,
    )
    result: dict[str, Any] = {"items": items, "count": len(items)}
    if include_electron_groups:
        result["electron_groups"] = electron_groups_from_index(
            db,
            root_ids=root_ids or None,
            inspection_number=inspection_number,
            recent_days=recent_days,
            limit=limit,
        )
    return result


@router.get("/files/index/{entry_id}/preview")
def preview_indexed_electron_image(
    entry_id: str,
    _auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    row = (
        db.query(ExecutionFileIndexEntry, ExecutionStorageRoot)
        .join(
            ExecutionStorageRoot,
            ExecutionStorageRoot.id
            == ExecutionFileIndexEntry.storage_root_id,
        )
        .filter(ExecutionFileIndexEntry.id == entry_id)
        .one_or_none()
    )
    if row is None:
        raise not_found("索引图片", entry_id)
    entry, root = row
    if (
        root.root_id != ELECTRON_ROOT_ID
        or entry.missing_since is not None
        or entry.extension.casefold() not in ELECTRON_IMAGE_SUFFIXES
    ):
        raise ExecutionApiError(
            415,
            "indexed_image_preview_unsupported",
            "该索引文件不是可预览的电镜图片",
        )
    gateway = build_file_gateway(db)
    try:
        path = gateway.resolve(
            ArtifactRef(root.root_id, entry.relative_path),
            expected_type="file",
        )
        stat = path.stat()
    except (StorageError, OSError) as exc:
        raise ExecutionApiError(
            409,
            "indexed_image_stale",
            "图片已移动或当前无法读取，请刷新文件索引",
        ) from exc
    if f"{stat.st_size}:{stat.st_mtime_ns}" != entry.fingerprint:
        raise ExecutionApiError(
            409,
            "indexed_image_stale",
            "图片已发生变化，请刷新文件索引后重新选择",
        )
    if stat.st_size > 30 * 1024 * 1024:
        raise ExecutionApiError(
            413,
            "indexed_image_preview_too_large",
            "图片超过 30 MiB，不能在线预览",
        )
    media_type = _inline_raster_media_type(entry.extension)
    if media_type is None:
        raise ExecutionApiError(
            415,
            "indexed_image_preview_unsupported",
            "该图片格式不能在线预览",
        )
    return FileResponse(
        path,
        media_type=media_type,
        headers={
            "Content-Disposition": "inline",
            "Content-Security-Policy": "default-src 'none'; sandbox",
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "private, max-age=60",
        },
    )


@router.post("/files/refresh", status_code=202)
def refresh_files(
    payload: FileRefreshRequest,
    auth: AuthContext = Depends(permission("file.read", csrf=True)),
    db: Session = Depends(get_db),
):
    job, duplicate = queue_refresh(
        db,
        root_id=payload.root_id,
        actor=auth.user,
        max_depth=payload.max_depth,
    )
    root = db.get(ExecutionStorageRoot, job.storage_root_id)
    db.commit()
    return {"duplicate": duplicate, "job": _index_job_dict(job, root)}


@router.get("/files/refresh/{job_id}")
def refresh_status(
    job_id: str,
    auth: AuthContext = Depends(permission("file.read")),
    db: Session = Depends(get_db),
):
    job = db.get(ExecutionIndexJob, job_id)
    if job is None or (
        job.requested_by_id not in {None, auth.user.id}
        and not _may_access_all_runs(db, auth)
    ):
        raise not_found("索引任务", job_id)
    root = db.get(ExecutionStorageRoot, job.storage_root_id)
    return _index_job_dict(job, root)
    ExecutionPublishReceipt,
